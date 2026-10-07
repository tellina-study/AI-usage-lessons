#!/usr/bin/env python3
"""Служебные надписи в ВИДИМОМ СЛОЕ готовой деки — чего зал видеть не должен.

Круг 4, правила А3 («надзаголовки жанра убрать») и А4 («комментариев-подсказок
лектору на слайдах быть не должно»). Владелец дословно: «это презентация, а не
инструкция преподавателя; я сам прекрасно помню, что они должны сделать».

ЧИТАЕТ ГОТОВЫЙ `.pptx`, А НЕ ИСХОДНИКИ — и это главное в этой проверке. Текст
слайда приходит из трёх разных мест: из `## Visual` самого `.md`, из питоновских
генераторов схем (`make_figures*.py`) и из самого сборщика. Grep по `.md` видит
только первое, а на Семинаре 5 снятая плашка «разбор — на следующем слайде»
жила ровно в третьем месте — в `build_sem05.py`, и ни в одном слайде её текста
не было. Проверка по исходникам объявила бы деку чистой — тот же риск остаётся
и для Семинара 6, поэтому проверка перенесена без ослабления правил.

«Инструмент, через который смотрят, обязан быть тем же, который собирает
продукт» (`tools/presentation-build/proverki-i-pravila.md`, общий корень).

ЗАМЕТКИ ДОКЛАДЧИКА НЕ ПРОВЕРЯЮТСЯ, И ЭТО НЕ НЕДОСМОТР: по А4 место таких вещей
именно в заметках. Запрет — на видимый слой.

    python3 audit_sluzhebnoe.py [файл.pptx]      # по умолчанию sem-06.pptx

Код возврата 1, если что-то найдено, — чтобы проверка годилась для сборочной
цепочки, а не только для чтения глазами.
"""
import re
import sys
from pathlib import Path

import yaml
from pptx import Presentation

HERE = Path(__file__).parent
ROOT = HERE.parent

# Слот надзаголовка: `deck_kit.auto_header` ставил его ровно в (0,55″; 0,34″).
# Положение — точный признак, а текст сам по себе нет: «ХУК» капителями это и
# надзаголовок, и шапка колонки таблицы на n32. Первое запрещено, второе —
# содержание слайда.
KICKER_X, KICKER_Y, SLOT = 0.55, 0.34, 0.12

# Тайминг запрещён НЕ ВЕЗДЕ. На этой деке минуты могут быть содержанием кейса
# (субагент-кейс 1: «задача стала занимать час» и т.п.) — запрет (CLAUDE.md,
# «No Timing / No Methodology in Slides») адресован ровно местам, где минута
# может быть ТОЛЬКО хронометражом занятия: обложка, дивайдеры, карта занятия,
# ось. Семинар 6 (в отличие от Семинара 5) не завёл отдельных имён приёма для
# «повторного вопроса»/«закрытия» — `reflection_question` (вопрос без карточек)
# обслуживает и n04 (первое предъявление вопроса занятия), и n67 (тот же
# вопрос второй раз) одним именем, а `question_with_option_cards` — обычные
# вопросы кейсов; ни один из двух сюда НЕ входит: запрет остался бы либо
# слишком узким (не считал бы тайминг дефектом на n67), либо слишком широким
# (считал бы дефектом любое число на любом вопросе кейса). Уже эти четыре
# приёма однозначно methodological-only.
TIMING_PATTERNS = {"hero_cover", "section_divider_macro", "lecture_map", "recap_table"}

# Имена разделов этой деки (SVODKA-SLAIDOV.md): Мостик, MCP, Субагент, Сборка.
# Надзаголовок жанра в Семинаре 6 не печатается ни на одном слайде (круг 4,
# А3 — то же решение, что в Семинаре 5), но проверка остаётся на будущее по
# той же логике, что и `audit_figs`/`audit_figure_text`.
RAZDELY = ("МОСТИК", "MCP", "СУБАГЕНТ", "СБОРКА")

# Надзаголовок печатался как «РАЗДЕЛ · ЖАНР». Форма с разделителем ловится по
# тексту где угодно на слайде; голое имя раздела — только в слоте надзаголовка,
# иначе под правило попадает любая шапка колонки.
KICKER_PAIR = re.compile(r"^(?:%s)\s·\s\S" % "|".join(RAZDELY))
KICKER_BARE = re.compile(r"^(?:%s)$" % "|".join(RAZDELY))

PRAVILA = [
    # Названные владельцем дословно (А4).
    ("подсказка лектору (А4)", re.compile(
        r"разбор\s+—\s+на\s+следующем\s+слайде"
        r"|здесь\s+ничего\s+не\s+реша"
        r"|выбор\s+делает\s+зал"
        r"|спрашивается\s+(?:\*\*)?порядок"
        r"|вы\s+здесь\b"
        # «Лектору:» / «Преподавателю —» как ОБРАЩЕНИЕ в начале строки. Просто
        # слово «преподавателя» запрещать нельзя: заголовок n29 — «Случай
        # преподавателя курса», и это содержание кейса.
        r"|^\s*(?:лектору|преподавателю)\s*[:—-]", re.I | re.M)),

    ("методический комментарий", re.compile(
        r"(?:методическ|педагогическ)\w*"
        r"|на\s+этом\s+этапе\s+студент|здесь\s+студент\s+усваивает"
        r"|зачем\s+это\s+в\s+(?:лекции|семинаре)|этот\s+раздел\s+учит", re.I)),

    # Строительные леса: в видимый слой не попадают никогда.
    ("разметка производства", re.compile(
        r"\[VERIFY-DAY-OF\]|\[FACT-CHECK\]|\[VFY-|\bLO[1-9][a-z]?\b|§\s?\d|→\s?s\d+")),
]

# Применяется только к приёмам из TIMING_PATTERNS — см. комментарий выше.
TIMING_RX = re.compile(
    r"\b\d+\s*мин(?:ут[аыу]?)?\b|время\s+раздел|тайминг|длительность|⏱|⏰", re.I)


EMU = 914400.0


def visible_text(slide):
    """Весь видимый текст слайда с ЕГО ПОЛОЖЕНИЕМ: (текст, x, y) в дюймах.

    Положение нужно одному правилу — надзаголовку: отличить его от шапки
    колонки можно только по слоту, в котором он стоял. БЕЗ заметок: по А4 их
    место именно там."""
    out = []

    def walk(shapes, dx=0.0, dy=0.0):
        for sh in shapes:
            x = (sh.left or 0) / EMU + dx
            y = (sh.top or 0) / EMU + dy
            if sh.shape_type == 6:                      # группа
                walk(sh.shapes, x, y)
                continue
            if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    t = "".join(r.text for r in p.runs).strip()
                    if t:
                        out.append((t, x, y))
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    for cell in row.cells:
                        t = cell.text.strip()
                        if t:
                            out.append((t, x, y))

    walk(slide.shapes)
    return out


def main():
    src = HERE / (sys.argv[1] if len(sys.argv) > 1 else "sem-06.pptx")
    prs = Presentation(str(src))
    patterns = [(s_.get("visual") or {}).get("pattern", "")
                for s_ in yaml.safe_load(
                    (ROOT / "deck.yaml").read_text(encoding="utf-8"))["slides"]]

    nayden = []
    for i, slide in enumerate(prs.slides, 1):
        pattern = patterns[i - 1] if i <= len(patterns) else ""
        for line, x, y in visible_text(slide):
            for imya, rx in PRAVILA:
                if rx.search(line):
                    nayden.append((i, imya, line))
            if KICKER_PAIR.match(line) or (
                    KICKER_BARE.match(line)
                    and abs(x - KICKER_X) < SLOT and abs(y - KICKER_Y) < SLOT):
                nayden.append((i, "надзаголовок жанра (А3)", line))
            if pattern in TIMING_PATTERNS and TIMING_RX.search(line):
                nayden.append((i, f"тайминг на слайде ({pattern})", line))

    print(f"Служебные надписи в видимом слое — {src.name}, {len(prs.slides)} слайдов\n")
    if not nayden:
        print("  чисто: ни одного совпадения по пяти правилам")
        return 0
    for nomer, imya, line in nayden:
        print(f"  слайд {nomer:>3}  {imya:<28} «{line[:88]}»")
    print(f"\n  всего: {len(nayden)} на {len({n for n, _i, _l in nayden})} слайдах")
    return 1


if __name__ == "__main__":
    sys.exit(main())
