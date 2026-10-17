# -*- coding: utf-8 -*-
"""САМОПРОВЕРКА ЯДРА: каждая проверка — на нарочно сломанном входе.

Стандарт здесь один и он не обсуждается: ПРОВЕРКА, НЕ ПРОВЕРЕННАЯ НА
СЛОМАННОМ ВХОДЕ, — НЕ ПРОВЕРКА. Поэтому на каждое предупреждение два прогона:
чистый кадр (обязан молчать) и сломанный (обязан сработать).

Сверх файловых идут РАЗНОСТНЫЕ — где мало знать, что предупреждение
сработало: надо знать, ОТ ЧЕГО. Пробы на РАСКЛАДКУ ТЕЛА — все разностные, и
лежат в `selftest_layout.py`; зовутся отсюда же.

  [отбивка]       от отбивки абзаца, и ровно на её величину (7 пт на разрыв)
  [курсив]        от курсива, и так, что коэффициентом это не лечится
  [мера-06]       мера ядра совпадает с мерой Лекции 6 на однородном тексте
  [предел-столбца] столбец не уже слова — И без предела был бы уже
  [зебра]         полосы таблицы идут в том порядке, в каком шли
  [схема-ии]      «модель» — шаг модели, «моделирование» — НЕ шаг модели
  [тень]          `<p:style>` снимается, и флаг делает ровно то, что обещает
  [шрифт]         отсутствующее начертание выходит пробелом, а не оценкой
  [ссылка-живая]  гиперссылка лежит в пакете как `r:id`, а не как синий текст
  [отпечаток]     регрессия ловит сдвиг фигуры на ОДИН EMU

Пробы на ПОЛ КЕГЛЯ, лестницу заголовка, число-заголовок и две колонки —
`selftest_floor.py`; зовутся оттуда же, откуда пробы раскладки.

Запуск:  python3 -m deckkit.selftest
"""
import sys
import tempfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE

from . import (axis, ladder, markup, metrics, plates, prims, probes,
               refdeck, regress, render, report, scheme, selftest_contrast,
               selftest_layout, tables, tracks)
from .theme import SURFACE, WHITE

OK, BAD, SAID = [], [], []


def say(line):
    SAID.append(line)
    print(line)


def _assets(d):
    """Снимки под пробы §5.9: портрет (0,67 — даёт четверть площади),
    широкий (1,5 — даёт 38,3%, ниже порога) и крупный (1,4 — даёт 41%,
    выше порога). Рисуются здесь, а не берутся из деки: проба ядра не
    вправе зависеть от ассетов лекции.

    Третий заведён для ЧИСТОГО входа О18 (#229): иллюстрация без подписи
    обязана пройти молча, а на отношении 1,5 она спотыкалась бы о ДРУГУЮ
    проверку — порог площади, — и молчание осталось бы недоказанным.
    Отношение 1,4 взято из `probes_contrast._hero_assets`, где тот же
    порог уже подобран."""
    Image.new("RGB", (670, 1000), (200, 210, 220)).save(d / "probe-portrait.png")
    Image.new("RGB", (1500, 1000), (210, 215, 225)).save(d / "probe-wide.png")
    Image.new("RGB", (1400, 1000), (205, 212, 222)).save(d / "probe-hero-big.png")


def _run(files, tmp, seq, name="probe"):
    report.reset()
    tables.reset_caches()
    seq[0] += 1
    d = tmp / f"d{seq[0]}"
    d.mkdir()
    _assets(d)
    for nm, txt in files:
        (d / nm).write_text(txt, encoding="utf-8")
    refdeck.build(d, tmp / f"{name}.pptx", root=d)
    return list(report.WARNINGS)


# ════════════════════════════════════════════════════════════════════
# файловые пробы
# ════════════════════════════════════════════════════════════════════


def file_probes(tmp, seq):
    base = _run([("s01-clean.md", probes.CLEAN),
                 ("s02-clean-table.md", probes.CLEAN_TABLE),
                 ("s03-clean-figure.md", probes.CLEAN_FIGURE),
                 ("s04-clean-kinds.md", probes.CLEAN_KINDS),
                 ("s05-clean-ladder.md", probes.CLEAN_LADDER),
                 ("s06-clean-axis.md", probes.CLEAN_AXIS),
                 ("s07-clean-illustration.md", probes.CLEAN_ILLUSTRATION)],
                tmp, seq, name="clean")
    say("── САМОПРОВЕРКА ЯДРА ────────────────────────────────────────")
    say(f"чистые кадры (прогон · таблица · одно число · дорожки · "
        f"лестница · ось · иллюстрация без подписи): предупреждений"
        f" {len(base)}"
        + (" — " + "; ".join(base) if base else " (молчат, как и должны)"))
    (BAD if base else OK).append("[чистый кадр]")
    if base:
        BAD.append("чистый кадр выдал предупреждение — проверка шумит")

    for key, (sid, txt) in probes.BROKEN.items():
        tag, _, need = key.partition(" ")
        w = _run([(f"{sid}-broken.md", txt)], tmp, seq)
        hit = [x for x in w if x.startswith(tag) and (not need or need in x)]
        say(f"{key:<24} на сломанном входе: "
            + ("ЛОВИТ" if hit else "НЕ ЛОВИТ")
            + (f"  →  {hit[0][:88]}" if hit else ""))
        if key in probes.LABELS:
            say(f"{'':24} доказывает: {probes.LABELS[key]}")
        (OK if hit else BAD).append(key)

    # Два файла на один номер кадра: читатель молча берёт первый, и об этом
    # обязан сказать.
    w = _run([("s01-a.md", probes.CLEAN), ("s01-b.md", probes.CLEAN)],
             tmp, seq)
    hit = [x for x in w if x.startswith("[файлы]")]
    say(f"{'[файлы]':<24} на сломанном входе: "
        + ("ЛОВИТ" if hit else "НЕ ЛОВИТ")
        + (f"  →  {hit[0][:88]}" if hit else ""))
    (OK if hit else BAD).append("[файлы]")
    return base


# ════════════════════════════════════════════════════════════════════
# разностные пробы
# ════════════════════════════════════════════════════════════════════


def probe_space_before():
    """Отбивка абзацев входит в меру — и ровно на свою величину.

    Без этой половины достаточно было бы уменьшить отбивку вдвое:
    предупреждение осталось бы на месте, а мера снова начала бы врать,
    только тише."""
    probe = [{"text": f"Абзац {i}.", "size": 15.0,
              **({"newpara": True, "space_before": 7} if i > 1 else {})}
             for i in range(1, 9)]
    flat = [{k: v for k, v in r.items() if k != "space_before"} for r in probe]
    w = 12.23
    delta = metrics.runs_h(probe, w, 1.26) - metrics.runs_h(flat, w, 1.26)
    want = 7 * 7.0 / 72.0            # семь разрывов на восемь абзацев
    if abs(delta - want) > 0.005:
        BAD.append(f"[отбивка] мера разрыва {delta:.3f}\" против положенных "
                   f"{want:.3f}\" — отбивка считается не вся")
    else:
        OK.append("[отбивка]")
        say(f"{'[отбивка]':<24} в мере высоты: ЛОВИТ  →  восемь абзацев дают "
            f"{delta:.3f}\" разрывов, ровно 7 пт на каждый")


def probe_italic():
    """Курсив меряется СВОИМИ метриками, и коэффициентом это не лечится.

    Берётся та самая строка, на которой дефект был виден: оговорка плашки
    `s63` Лекции 7 во всю ширину плашечной колонки. По ней держатся три
    утверждения сразу.
      1. Правда — ТРИ строки, а мера по прямому начертанию обещает две.
      2. Коэффициент не спасает ни щедрый (1,031 — тот, что стоял), ни
         точный для этой строки: оба дают те же две строки.
      3. Единого верного коэффициента не существует, и видно это внутри
         ОДНОЙ строки: «дублированная» в курсиве УЖЕ прямого,
         «рецензентов» — ШИРЕ. Одно число не может быть одновременно меньше
         единицы и больше неё."""
    note = ("слепая дублированная оценка 22 рецензентов "
            "на 100 госпитализациях")
    st = plates.LEC07
    # Кегль ТОТ, НА КОТОРОМ ДЕФЕКТ БЫЛ ВИДЕН (оговорка Лекции 7 до поднятия
    # пола видимого слоя), а не нынешний `st.note_pt`: проба про МЕРУ
    # курсива, а не про сегодняшний набор. Привязанная к действующему числу,
    # она умерла молча в тот же круг, когда пол подняли до 12 пт.
    note_pt = 10.5
    iw = (12.23 - st.gap * 3) / 4 - 2 * st.pad
    true_n = metrics.n_lines(note, note_pt, iw, False, True)
    own = metrics._unit_w(note, False, True) / metrics._unit_w(note, False,
                                                               False)
    proxy = {c: metrics.n_lines(note, note_pt, iw / c, False, False)
             for c in (1.0, 1.0235, 1.031)}
    narrow = (metrics._unit_w("дублированная", False, True)
              / metrics._unit_w("дублированная", False, False))
    wide = (metrics._unit_w("рецензентов", False, True)
            / metrics._unit_w("рецензентов", False, False))
    if true_n != 3 or any(v != 2 for v in proxy.values()):
        BAD.append(f"[курсив] правда {true_n} строк, прокси по "
                   f"коэффициентам {proxy} — проба больше не доказывает, "
                   "что коэффициент не спасает")
    elif not narrow < 1.0 < wide:
        BAD.append(f"[курсив] внутри одной строки курсив больше не отходит "
                   f"от прямого в ДВЕ стороны ({narrow:.3f} и {wide:.3f}) — "
                   "похоже, мера снова идёт по одному начертанию")
    else:
        OK.append("[курсив]")
        say(f"{'[курсив]':<24} в мере числа строк: ЛОВИТ  →  правда "
            f"{true_n} строки, прокси по прямому даёт 2 при любом "
            f"коэффициенте (1,0 · {own:.4f} своё · 1,031 прежний); в одной "
            f"строке курсив и уже прямого ({narrow:.3f}), и шире "
            f"({wide:.3f})")


def probe_lec06_parity():
    """Мера ядра = мера Лекции 6 на ОДНОРОДНОМ тексте.

    Это то утверждение, на котором держится вся регрессия: если бы мера
    расходилась, общее ядро сдвигало бы геометрию доставленной деки,
    прочитанной залу. Здесь оно проверяется прямо — своей, независимой
    реализацией прежней меры, а не ссылкой на то, что регрессия прошла."""
    def old_n_lines(text, size, w_in, bold=False):
        maxw = w_in * 72.0
        words = text.split()
        if not words:
            return 1
        lines, cur = 1, words[0]
        for wd in words[1:]:
            t = cur + " " + wd
            if metrics.text_w(t, size, bold) <= maxw:
                cur = t
            else:
                lines, cur = lines + 1, wd
        return lines

    def old_est(text, size, w_in, ls, bold=False):
        rows = sum(old_n_lines(p, size, w_in, bold) for p in text.split("\n"))
        return rows * size * ls * metrics.LH / 72.0 * metrics.SAFE

    cases = [
        ("Короткая строка", 15.0, 3.0, 1.24, False),
        ("Классификация на тексте: на входе элемент спецификации, на выходе "
         "метка «требование / не требование»", 13.5, 4.45, 1.18, False),
        ("ЗАГОЛОВОК ПОЛУЖИРНЫМ НА УЗКОЙ КОЛОНКЕ", 11.0, 1.7, 1.08, True),
        ("Абзац\nс разрывом\nи ещё одним", 16.0, 6.0, 1.26, False),
        ("", 16.0, 6.0, 1.18, False),
        ("Неразрывноеоченьдлинноеслово", 17.0, 1.2, 1.22, True),
    ]
    diff = []
    for text, size, w, ls, bold in cases:
        a, b = metrics.n_lines(text, size, w, bold), old_n_lines(text, size,
                                                                 w, bold)
        c, d = metrics.est_h(text, size, w, ls, bold), old_est(text, size, w,
                                                               ls, bold)
        if a != b or abs(c - d) > 1e-12:
            diff.append((text[:28], a, b, c, d))
    if diff:
        BAD.append(f"[мера-06] мера ядра разошлась с прежней на "
                   f"{len(diff)} входах: {diff[:2]}")
    else:
        OK.append("[мера-06]")
        say(f"{'[мера-06]':<24} совместимость меры: ЛОВИТ  →  на {len(cases)} "
            "однородных входах число строк и высота совпадают с прежней "
            "мерой Лекции 6 до последнего знака")


def probe_col_min():
    """Предел ширины столбца работает — И без него столбец был бы уже слова.

    Проба разностная нарочно: «все столбцы шире своих слов» — утверждение,
    которое верно и у таблицы, где делить нечего. Ценность в том, что
    ПРОПОРЦИОНАЛЬНОЕ деление (как было до правки) эту же таблицу ломает."""
    table = [["Прослеживаемость",
              "Что именно делает режим и как это проверяют на практике "
              "инженером и врачом вместе", "Да"],
             ["Электроэнцефалографический",
              "отбирает спорные места в спецификации и передаёт их человеку "
              "на проверку целиком", "Нет"],
             ["Структурированность",
              "проверяет разметку человека и сверяет её с разметкой модели "
              "по всем срезам", "Да"]]
    rows = tables.rows_of(table)
    size, w = 11.0, 8.0
    widths = tables.table_widths(rows, size, w)
    mins = tables.col_min(rows, size, tables.TBL_TIGHT)
    want = tables.col_want(rows, size)
    tot = sum(want)
    naive = [w * x / tot for x in want]
    starved = [j for j in range(len(mins)) if naive[j] < mins[j] - 1e-9]
    broken = [j for j in range(len(mins)) if widths[j] < mins[j] - 1e-9]
    if not starved:
        BAD.append("[предел-столбца] пропорциональное деление этой таблицы "
                   "уже никого не обделяет — проба перестала доказывать "
                   "правку, её вход надо подобрать заново")
    elif broken:
        BAD.append(f"[предел-столбца] столбцы {broken} вышли уже своего "
                   "самого длинного слова — предел не держит")
    else:
        OK.append("[предел-столбца]")
        say(f"{'[предел-столбца]':<24} ширины столбцов: ЛОВИТ  →  "
            f"пропорциональное деление обделяло столбцы {starved} "
            f"({naive[starved[0]]:.2f}\" против нужных "
            f"{mins[starved[0]]:.2f}\"), с пределом ни один не уже своего "
            "слова")


def probe_zebra(tmp):
    """Порядок полос таблицы. Переставленная пара даёт ровную шахматку
    наоборот — на снимке почти не видно, в XML видно сразу. Регрессия на
    деке Лекции 6 этот порядок уже один раз поймала; здесь он закреплён
    пробой, чтобы впредь ловилось без деки."""
    report.reset()
    tables.reset_caches()
    p = prims.new_deck()
    s = prims.blank(p)
    table = [["Шапка", "Вторая"], ["первая", "строка"], ["вторая", "строка"]]
    tables.draw_table(s, 0.5, 0.5, table, 11.0, 6.0, "проба")
    fills = []
    for shp in s.shapes:
        try:
            fills.append(str(shp.fill.fore_color.rgb))
        except Exception:
            fills.append(None)
    # На каждую ячейку по ДВЕ фигуры: плашка и надпись. Шапка — ячейки 0-1,
    # то есть фигуры 0-3; первая строка — фигуры 4-7, вторая — 8-11.
    row1, row2 = fills[4:8:2], fills[8:12:2]
    # Первая строка после шапки — БЕЛАЯ, вторая на поверхности: `i % 2`,
    # где `i` — номер строки вместе с шапкой. Порядок выглядит обратным
    # интуиции и ровно поэтому закреплён пробой: регрессия на деке Лекции 6
    # поймала его переставленным, и на снимке это почти не видно.
    want1, want2 = str(WHITE), str(SURFACE)
    if row1 != [want1, want1] or row2 != [want2, want2]:
        BAD.append(f"[зебра] полосы таблицы пошли {row1} / {row2}, а должны "
                   f"{want1} / {want2}")
    else:
        OK.append("[зебра]")
        say(f"{'[зебра]':<24} порядок полос таблицы: ЛОВИТ  →  первая строка "
            f"после шапки белая ({want1}), вторая на поверхности ({want2})")


def probe_scheme_ai():
    """«модель» — шаг модели, «моделирование» — НЕ шаг модели.

    Поймано на настоящем `s02` Лекции 6: карта цикла изделия получила
    золотую обводку на стадии «3. Моделирование», которая к модели не
    относится. Утверждение разностное: один вход обязан сработать, другой
    обязан НЕ сработать."""
    yes = bool(scheme.AI_WORD.search("модель отбирает спорное"))
    no = bool(scheme.AI_WORD.search("3. Моделирование"))
    also = bool(scheme.AI_WORD.search("нейросеть размечает"))
    if not (yes and also) or no:
        BAD.append(f"[схема-ии] «модель» {yes}, «нейросеть» {also}, "
                   f"«Моделирование» {no} — различение сломано")
    else:
        OK.append("[схема-ии]")
        say(f"{'[схема-ии]':<24} шаг модели: ЛОВИТ  →  «модель» и "
            "«нейросеть» да, «Моделирование» нет")


def probe_scheme_ai_explicit():
    """ЯВНОЕ сильнее шаблона в ОБЕ стороны, и `ai_by_word=False` снимает
    шаблон. Без этой пробы деке, у которой шаблон врёт (Лекция 7: четыре
    ложных золота из пяти), осталось бы только мириться с ним."""
    class S:
        sid = "s01"
        meta = {"scheme": {"rows": [{"steps": [
            "модель отбирает спорное",            # шаблон: да
            {"text": "модель отбирает", "ai": False},   # явное НЕТ сильнее
            {"text": "вычитка инженером", "ai": True},  # явное ДА сильнее
        ]}]}}
    off = scheme.SchemeStyle(ai_by_word=False)
    by_word = [x["ai"] for x in scheme.scheme_of(S(), scheme.DEFAULT)["rows"][0]["steps"]]
    no_word = [x["ai"] for x in scheme.scheme_of(S(), off)["rows"][0]["steps"]]
    if by_word != [True, False, True] or no_word != [False, False, True]:
        BAD.append(f"[схема-явное] по шаблону {by_word}, без шаблона "
                   f"{no_word} — ожидалось [да,нет,да] и [нет,нет,да]")
    else:
        OK.append("[схема-явное]")
        say(f"{'[схема-явное]':<24} явное против шаблона: ЛОВИТ  →  "
            "`ai: false` гасит шаблон, `ai: true` ставит шаг без слова, "
            "`ai_by_word=False` снимает шаблон целиком")


def probe_scheme_arrows():
    """`arrows: false` — ряд перечень, а не конвейер: стрелок нет, и
    просвет между боксами свой. Ядро поле игнорировало, то есть рисовало на
    перечне участников три стрелки (`s70` Лекции 7) — порядок, которого в
    материале нет."""
    class S:
        sid = "s01"
        meta = {"scheme": {"arrows": False,
                           "rows": [{"steps": ["раз", "два", "три"]}]}}
    class T:
        sid = "s01"
        meta = {"scheme": {"rows": [{"steps": ["раз", "два", "три"]}]}}
    off = scheme.scheme_of(S(), scheme.DEFAULT)
    on = scheme.scheme_of(T(), scheme.DEFAULT)
    w_off = scheme._step_w(off, 3, 12.0, scheme.DEFAULT)
    w_on = scheme._step_w(on, 3, 12.0, scheme.DEFAULT)
    def arrows_drawn(sch):
        p = prims.new_deck()
        sl = prims.blank(p)
        scheme.draw_scheme(sl, "s01", sch, 0.5, 0.5, 12.0, scheme.DEFAULT)
        n = 0
        for sh in sl.shapes:
            try:
                n += 1 if sh.auto_shape_type == MSO_SHAPE.RIGHT_ARROW else 0
            except (AttributeError, ValueError):
                pass
        return n
    n_off, n_on = arrows_drawn(off), arrows_drawn(on)
    if n_off != 0 or n_on != 2 or not w_off > w_on:
        BAD.append(f"[схема-стрелки] без стрелок нарисовано {n_off}, со "
                   f"стрелками {n_on}, бокс {w_off:.2f} против {w_on:.2f} — "
                   "поле `arrows` снова ничего не меняет")
    else:
        OK.append("[схема-стрелки]")
        say(f"{'[схема-стрелки]':<24} `arrows: false`: ЛОВИТ  →  стрелок 0 "
            f"против 2, бокс шире ({w_off:.2f}\" против {w_on:.2f}\")")


def probe_scheme_newline():
    """Перевод строки внутри шага — РАЗРЫВ АБЗАЦА, и мера обязана считать
    его так же, как его рисует `text_box`. Прежде мера считала шаг одним
    потоком и недомеривала карточку участника на строку: бокс выходил ниже
    своего текста (`s70` Лекции 7, карточка «Разработчик»)."""
    flat = "Разработчик контролирует данные"
    para = "Разработчик\nконтролирует данные"
    iw = 3.00
    n_flat = scheme._rows(flat, 11.5, iw)
    n_para = scheme._rows(para, 11.5, iw)
    n_blind = metrics.n_lines(para, 11.5, iw, True)
    if not (n_para > n_blind and n_para == n_flat + 1):
        BAD.append(f"[схема-абзац] по абзацам {n_para}, одним потоком "
                   f"{n_blind}, без разрыва {n_flat} — мера снова считает "
                   "перевод строки пробелом")
    else:
        OK.append("[схема-абзац]")
        say(f"{'[схема-абзац]':<24} разрыв абзаца в шаге: ЛОВИТ  →  строк "
            f"{n_para} против {n_blind}, которые давала мера одним потоком")


def probe_row_arrows(tmp, seq):
    """ПОРЯДНЫЙ `arrows` СИЛЬНЕЕ СХЕМНОГО, и это проверяется на собранном
    кадре, а не на разборе.

    Цена отсутствия заплачена живой декой: на `s36` порядный ключ УЖЕ
    стоял, ядро его молча выбрасывало, и автор считал ряд починенным;
    на `s01a` схемный `arrows: false` снял стрелки у всех трёх рядов, хотя
    ветка выбора только во втором. Поэтому проба меряет ТРИ состояния на
    одной схеме: ряд со своим `true`, ряд со своим `false` и ряд без
    ключа, наследующий схемный."""
    txt = """---
id: s01
type: case
title: "Проба порядного ключа"
source: "—"
scheme:
  arrows: false
  rows:
    - {label: "Свой true", arrows: true, steps: ["раз", "два", "три"]}
    - {label: "Свой false", arrows: false, steps: ["раз", "два", "три"]}
    - {label: "Без ключа", steps: ["раз", "два", "три"]}
---

## Видимый текст слайда

Тело.

## Заметки лектора

проба
"""
    report.reset()
    _run([("s01-probe.md", txt)], tmp, seq, name="порядный")
    n = report.PRIM["стрелка"]
    # Один ряд со стрелками → ровно две стрелки на три шага; два
    # остальных наследуют/объявляют отказ и не дают ни одной.
    sch = {"lab_w": 1.7, "arrows": False,
           "rows": [{"label": "", "steps": [1, 2, 3], "arrows": True},
                    {"label": "", "steps": [1, 2, 3], "arrows": None}]}
    шир_со = scheme._step_w(sch, sch["rows"][0], 11.0, scheme.DEFAULT)
    шир_без = scheme._step_w(sch, sch["rows"][1], 11.0, scheme.DEFAULT)
    if not (n == 2 and шир_без > шир_со):
        BAD.append(f"[схема-порядные] стрелок {n} вместо 2; ширина со "
                   f"стрелками {шир_со:.2f}, без {шир_без:.2f} — порядный "
                   "ключ снова не читается")
    else:
        OK.append("[схема-порядные]")
        say(f"{'[схема-порядные]':<24} на собранном кадре: ЛОВИТ  →  при "
            f"схемном `false` ряд со своим `true` дал {n} стрелки, два "
            f"других ни одной; и бокс ряда со стрелками уже ({шир_со:.2f}\" "
            f"против {шир_без:.2f}\") — мера читает тот же ключ")


def probe_axis_proportional():
    """РАДИ ЧЕГО ОСЬ ЗАВЕДЕНА: расстояние между засечками пропорционально
    разнице величин. Это то самое «разница видна глазом, а не читается»,
    и доказывать это надо числом, иначе ось — просто третий способ
    расставить подписи в ряд.

    Вход — настоящий `s49c`: 0,816 · 0,820 · 0,712. Полоса плашек ставит
    их на равные расстояния; ось обязана поставить две близкие рядом, а
    третью — далеко, И В ТОЙ ЖЕ ПРОПОРЦИИ, в какой отличаются числа."""
    d = axis.axis_of("проба", {
        "unit": "площадь под кривой",
        "marks": [{"value": "0,816", "label": "Модель"},
                  {"value": "0,820", "label": "Хирург"},
                  {"value": "0,712", "label": "Пара"}]})
    w = 11.0
    xs = [axis._x_of(m["v"], d, 0.0, w, axis.DEFAULT) for m in d["marks"]]
    близкие = abs(xs[1] - xs[0])          # 0,820 − 0,816 = 0,004
    далёкая = abs(xs[0] - xs[2])          # 0,816 − 0,712 = 0,104
    # Полоса плашек: три равных шага, то есть отношение ровно 1.
    плашки = 1.0
    ожидание = 0.104 / 0.004              # 26
    вышло = далёкая / близкие if близкие else 0
    if not (abs(вышло - ожидание) < 0.01 and вышло > плашки * 20):
        BAD.append(f"[ось-пропорция] отношение расстояний {вышло:.2f} при "
                   f"ожидаемом {ожидание:.2f} — ось расставляет засечки не "
                   "по величине")
    else:
        OK.append("[ось-пропорция]")
        say(f"{'[ось-пропорция]':<24} разностно: ЛОВИТ  →  0,816 и 0,820 "
            f"стоят в {близкие:.3f}\", а 0,712 — в {далёкая:.2f}\" от них: "
            f"отношение {вышло:.1f} против 1,0 у полосы плашек")


def probe_tracks_grid(tmp, seq):
    """РАДИ ЧЕГО ЗАВЕДЕНЫ ДОРОЖКИ: клетка одной ширины во всех дорожках.

    Настоящий `s62` — «было» в два шага против «стало» в четыре. Схема-
    конвейер считает ширину бокса по числу шагов РЯДА, поэтому короткий ряд
    получает боксы вдвое шире длинного, и сравнение «было/стало» выходит
    сравнением двух разных масштабов. Дорожки ставят обе линии на сетку
    самой длинной.

    Меряется ПО СОБРАННОМУ ФАЙЛУ, а не по формуле ширины: формула — это
    намерение, а доказывать надо то, что уехало в XML. Разница между двумя
    формами видна тогда набором ширин коробок на кадре."""
    rows = """
  rows:
    - {label: "Было", steps: ["набор", "подпись"]}
    - {label: "Стало", steps: ["разговор", "черновик", "вычитка", "подпись"]}
"""
    head = """---
id: s01
type: case
title: "Проба общей сетки"
source: "—"
scheme:
"""
    tail = """---

## Видимый текст слайда

Тело.

## Заметки лектора

проба
"""

    СЛОВА = {"набор", "подпись", "разговор", "черновик", "вычитка"}

    def widths(txt, name):
        """Ширины НАДПИСЕЙ шагов на собранном кадре.

        Меряется надпись, а не коробка, по двум причинам: надпись — это
        ровно то, что читает глаз, и она одна у обеих форм, тогда как
        коробок у дорожек две (клетка и рельс), и отобрать из них клетку
        пришлось бы по заливке — то есть по признаку оформления."""
        _run([("s01-probe.md", txt)], tmp, seq, name=name)
        prs = Presentation(str(tmp / f"{name}.pptx"))
        out = []
        for sh in prs.slides[0].shapes:
            if not sh.has_text_frame:
                continue
            if sh.text_frame.text.strip() in СЛОВА:
                out.append(round(sh.width / 914400.0, 2))
        return out

    ряды = widths(head + "  arrows: true" + rows + tail, "сетка-ряды")
    дорожки = widths(head + "  kind: dorozhki" + rows + tail, "сетка-дорожки")
    кл_р, кл_д = sorted(set(ряды)), sorted(set(дорожки))
    if not (len(кл_р) >= 2 and len(кл_д) == 1 and max(кл_р) > min(кл_р) * 1.8):
        BAD.append(f"[дорожки-сетка] рядами ширины {кл_р}, дорожками {кл_д} "
                   "— общей сетки не вышло")
    else:
        OK.append("[дорожки-сетка]")
        say(f"{'[дорожки-сетка]':<24} на собранном кадре: ЛОВИТ  →  рядами "
            f"клетки {кл_р} (в {max(кл_р) / min(кл_р):.1f} раза шире там, "
            f"где шагов вдвое меньше), дорожками одна ширина {кл_д}")


def probe_ladder_bottom():
    """РАДИ ЧЕГО ЗАВЕДЕНА ЛЕСТНИЦА: крайняя ступень отличается ВИДОМ, и
    отличается ТА, которая крайняя по направлению.

    «Подписывает не глядя» — дно шкалы участия, и в нём весь смысл кадра.
    При `direction: вверх` крайней обязана стать первая ступень, а не
    остаться последняя: иначе направление — подпись, а не свойство
    рисунка."""
    rows = [{"label": "Ступени",
             "steps": [{"text": f"ступень {i}", "ai": False}
                       for i in range(1, 5)]}]
    вниз = ladder.ladder_of("проба", {"direction": "вниз"}, rows)
    вверх = ladder.ladder_of("проба", {"direction": "вверх"}, rows)
    w = 11.0
    y_вниз = [i * ladder.DEFAULT.rise for i in range(4)]
    rise = ladder.DEFAULT.rise * 3
    y_вверх = [rise - i * ladder.DEFAULT.rise for i in range(4)]
    h1, h2 = ladder.ladder_h(вниз, w), ladder.ladder_h(вверх, w)
    if not (вниз["down"] and not вверх["down"]
            and y_вниз[-1] > y_вниз[0] and y_вверх[-1] < y_вверх[0]
            and abs(h1 - h2) < 1e-9):
        BAD.append("[лестница-дно] направление не меняет ни порядок высот, "
                   "ни крайнюю ступень")
    else:
        OK.append("[лестница-дно]")
        say(f"{'[лестница-дно]':<24} разностно: ЛОВИТ  →  вниз ступени "
            f"идут {y_вниз[0]:.2f}→{y_вниз[-1]:.2f}\", вверх "
            f"{y_вверх[0]:.2f}→{y_вверх[-1]:.2f}\"; высота обеих "
            f"{h1:.2f}\" — крайняя ступень разная, место одно")


def probe_kind_fallback():
    """Неизвестный `kind` НЕ роняет кадр и НЕ рисуется молча рядами: кадр
    собирается рядами и об этом сказано. Проверяется разностью: с чужим
    `kind` схема обязана дать те же фигуры, что без него."""
    raw = {"rows": [{"label": "Ряд", "steps": ["раз", "два"]}],
           "arrows": True}

    class _S:
        sid = "проба"
        meta = {"scheme": dict(raw)}

    class _S2:
        sid = "проба"
        meta = {"scheme": dict(raw, kind="spiral")}

    report.reset()
    a = scheme.scheme_of(_S())
    n_a = len(report.WARNINGS)
    report.reset()
    b = scheme.scheme_of(_S2())
    n_b = len(report.WARNINGS)
    report.reset()
    same = (a and b and a["rows"] == b["rows"] and a["arrows"] == b["arrows"]
            and not b.get("kind"))
    if not (same and n_a == 0 and n_b == 1):
        BAD.append(f"[форма-откат] предупреждений {n_b} при чужом `kind` и "
                   f"{n_a} без него; схемы {'совпали' if same else 'разные'}")
    else:
        OK.append("[форма-откат]")
        say(f"{'[форма-откат]':<24} чужой `kind`: ЛОВИТ  →  схема собрана "
            f"рядами, как без него (0 расхождений разбора), и сказано одной "
            f"строкой ({n_b} против {n_a})")


def probe_thin_switch(tmp, seq):
    """ДВЕ половины одного утверждения о диагностике тонкого запаса.

    1. ПРОБА НА НЕЁ — проба именно на неё. Файловая проба `[запас]` выше
       показывает, что предупреждение выходит; сама по себе она не
       доказывает, что выходит оно ОТ ДИАГНОСТИКИ. Здесь тот же вход
       собирается с выключенной диагностикой (`THIN_LINES = 0`), и она
       обязана замолчать. Без этой половины достаточно было бы любого
       предупреждения с тем же тегом.

    2. ДИАГНОСТИКА НЕ ДВИГАЕТ НИ ОДНОГО EMU. Это её главное обещание:
       она смотрит и не правит. Проверяется не на слово, а отпечатком —
       тот же вход, собранный с диагностикой и без неё, обязан дать
       побитово одинаковую геометрию всех фигур. Отпечаток ловит сдвиг на
       один EMU (проба `[отпечаток]`), то есть «расхождений 0» здесь —
       утверждение, а не отсутствие проверки."""
    sid, txt = probes.BROKEN["[запас]"]
    files = [(f"{sid}-broken.md", txt)]
    on = _run(files, tmp, seq, name="thin-on")
    saved = metrics.THIN_LINES
    try:
        metrics.THIN_LINES = 0.0
        off = _run(files, tmp, seq, name="thin-off")
    finally:
        metrics.THIN_LINES = saved
    hit_on = [w for w in on if w.startswith("[запас]")]
    hit_off = [w for w in off if w.startswith("[запас]")]
    if not hit_on or hit_off:
        BAD.append(f"[запас-выключено] с диагностикой предупреждений "
                   f"{len(hit_on)}, без неё {len(hit_off)} — проба ловит "
                   "не диагностику")
    else:
        OK.append("[запас-выключено]")
        say(f"{'[запас-выключено]':<24} проба на саму проверку: ЛОВИТ  →  "
            f"с диагностикой {len(hit_on)} предупреждение, с выключенной — "
            f"{len(hit_off)}")
    diff = regress.compare(regress.fingerprint(tmp / "thin-on.pptx"),
                           regress.fingerprint(tmp / "thin-off.pptx"))
    if diff:
        BAD.append(f"[запас-геометрия] диагностика сдвинула геометрию: "
                   f"{len(diff)} расхождений, например {diff[0]}")
    else:
        OK.append("[запас-геометрия]")
        say(f"{'[запас-геометрия]':<24} диагностика ничего не двигает: "
            "ЛОВИТ  →  тот же вход с ней и без неё даёт 0 расхождений "
            "отпечатка (а отпечаток ловит сдвиг на 1 EMU)")


def probe_shadow():
    """Тень снимается удалением `<p:style>`, и флаг делает ровно то, что
    обещает. По умолчанию — снимается: это разногласие двух сборщиков,
    разобранное в пользу Лекции 6 (`notes/mcp-limitations.md` [#223-8])."""
    P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
    res = {}
    for flag in (True, False):
        prims.STRIP_THEME_STYLE = flag
        p = prims.new_deck()
        s = prims.blank(p)
        shp = prims.ocean_box(s, 1, 1, 2, 1)
        res[flag] = len(shp._element.findall(P + "style"))
    prims.STRIP_THEME_STYLE = True
    if res[True] != 0 or res[False] == 0:
        BAD.append(f"[тень] со снятием осталось стилей {res[True]}, без "
                   f"снятия {res[False]} — флаг не делает того, что обещает")
    else:
        OK.append("[тень]")
        say(f"{'[тень]':<24} снятие темы: ЛОВИТ  →  по умолчанию `<p:style>` "
            f"снят (0), а в режиме воспроизведения Лекции 7 остаётся "
            f"({res[False]}) — именно он и даёт тень в LibreOffice")


def probe_font_gap():
    """Отсутствующее начертание выходит ПРОБЕЛОМ `[шрифт]`, а не молчаливой
    оценкой по соседнему файлу."""
    report.reset()
    saved = dict(metrics._TTF)
    cache = dict(metrics._FCACHE)
    gaps = dict(metrics._FONT_GAP)
    try:
        metrics._TTF[(False, True)] = Path("/nonexistent/NoSuchFace.ttf")
        metrics._FCACHE.pop((False, True), None)
        metrics._FONT_GAP.clear()
        metrics.text_w("проба", 12.0, False, True)
        metrics.font_gaps()
        hit = [w for w in report.WARNINGS if w.startswith("[шрифт]")]
    finally:
        metrics._TTF.clear(); metrics._TTF.update(saved)
        metrics._FCACHE.clear(); metrics._FCACHE.update(cache)
        metrics._FONT_GAP.clear(); metrics._FONT_GAP.update(gaps)
    if not hit:
        BAD.append("[шрифт] отсутствующее начертание прошло молча")
    else:
        OK.append("[шрифт]")
        say(f"{'[шрифт]':<24} на сломанном входе: ЛОВИТ  →  {hit[0][:88]}")


def probe_live_links(tmp):
    """Гиперссылка лежит в пакете как `r:id` на внешнее отношение, а не как
    подчёркнутый синий текст. Самоотчёт не принимается — считается XML.

    Считает его `render.verify_links` — та самая, которой сборка отчитывается
    о ссылках. Своя копия того же обхода стояла здесь до круга #229, и это
    была проверка, проверяющая НЕ ТО, чем пользуются: разойдись две копии,
    проба осталась бы зелёной."""
    n, targets = render.verify_links(tmp / "clean.pptx")
    if n < 2 or targets < 2:
        BAD.append(f"[ссылка-живая] в XML чистого кадра ссылок {n}, "
                   f"отношений {targets} — ждали хотя бы по две")
    else:
        OK.append("[ссылка-живая]")
        say(f"{'[ссылка-живая]':<24} в пакете: ЛОВИТ  →  ссылок на прогонах "
            f"{n} · отношений {targets}")


def probe_fingerprint(tmp):
    """Проверка на саму проверку: отпечаток регрессии обязан поймать сдвиг
    фигуры на ОДИН EMU. Отпечаток, который «всегда совпадает», доказывал бы
    не совпадение дек, а собственную слепоту."""
    from pptx.util import Emu
    src = tmp / "clean.pptx"
    p = prims.new_deck()
    s = prims.blank(p)
    prims.ocean_box(s, 1, 1, 2, 1)
    a = tmp / "fp-a.pptx"
    b = tmp / "fp-b.pptx"
    p.save(str(a))
    s.shapes[0].width = Emu(s.shapes[0].width + 1)
    p.save(str(b))
    same = regress.compare(regress.fingerprint(a), regress.fingerprint(a))
    diff = regress.compare(regress.fingerprint(a), regress.fingerprint(b))
    if same or not diff:
        BAD.append(f"[отпечаток] сам с собой {len(same)} расхождений, со "
                   f"сдвигом на 1 EMU — {len(diff)}; отпечаток слеп")
    else:
        OK.append("[отпечаток]")
        say(f"{'[отпечаток]':<24} чувствительность: ЛОВИТ  →  сдвиг на 1 EMU "
            f"даёт {len(diff)} расхождение, файл сам с собой — 0")
    # Побайтово те же два файла неотличимы от любых других: метки времени
    # делают байтовое сравнение неработающим, а не строгим.
    say(f"{'':24} (байтовое сравнение здесь не годится: `python-pptx` "
        "пишет в пакет метки времени)")


# ════════════════════════════════════════════════════════════════════


def selftest():
    tmp = Path(tempfile.mkdtemp(prefix="deckkit-selftest-"))
    seq = [0]
    file_probes(tmp, seq)
    say("── разностные пробы ─────────────────────────────────────────")
    probe_space_before()
    probe_italic()
    probe_lec06_parity()
    probe_col_min()
    probe_zebra(tmp)
    probe_scheme_ai()
    probe_scheme_ai_explicit()
    probe_scheme_arrows()
    probe_scheme_newline()
    probe_row_arrows(tmp, seq)
    probe_axis_proportional()
    probe_tracks_grid(tmp, seq)
    probe_ladder_bottom()
    probe_kind_fallback()
    probe_thin_switch(tmp, seq)
    probe_shadow()
    probe_font_gap()
    probe_live_links(tmp)
    probe_fingerprint(tmp)
    probe_cell_runs()
    probe_render_path()
    selftest_layout.run_all(OK, BAD, say, _run, tmp, seq)
    # Пол КОНТРАСТА и hero по ПОРЯДКУ деки — обе проверки круга прожарки
    # 2026-10-14, обе заведены после дефекта, который прошёл мимо всех
    # прежних сторожей. Свой файл, как и у раскладки: эти пробы собирают
    # фигуры и миниатюрные деки, а не кирпич текста.
    selftest_contrast.run_all(OK, BAD, say, _run, tmp, seq)
    say("─────────────────────────────────────────────────────────────")
    total = len(OK) + len(BAD)
    say(f"проверок доказано: {len(OK)} из {total}")
    if BAD:
        say("НЕ ДОКАЗАНО:")
        for x in BAD:
            say("  " + x)
        return 1
    say("все проверки доказаны на сломанном входе")
    return 0


def probe_cell_runs():
    """Клетка карты двумя начертаниями: своё полужирным, пройденное светлым.

    Доказывается РАЗНОСТНО и в обе стороны. Прежнее правило
    (`bold = "**" in cell`) красило ячейку полужирной ЦЕЛИКОМ от одной пары
    звёздочек, и накопительная карта стадии была невыполнима — три сессии
    подряд обошли это скобками, каждая сочтя свойством жанра. Проверка
    обязана ловить и обратное: при ВЫКЛЮЧЕННОМ переключателе вид доставленной
    Лекции 6 не двигается, иначе правка меняет уже прочитанную деку."""
    cell = "**подсказка, которой верят** (отбор · диагноз)"
    table = [["задача", "данные"], [cell, "изображение"]]

    def runs_of(mode):
        tables.CELL_RUNS = mode
        p = prims.new_deck()
        s = prims.blank(p)
        tables.reset_caches()
        tables.draw_table(s, 0.5, 0.5, table, 11.0, 6.0, "проба")
        for sh in s.shapes:
            if not sh.has_text_frame:
                continue
            for para in sh.text_frame.paragraphs:
                rs = [(r.text, bool(r.font.bold)) for r in para.runs]
                if rs and "подсказка" in "".join(x for x, _ in rs):
                    return rs
        return []

    try:
        off, on = runs_of(False), runs_of(True)
    finally:
        tables.CELL_RUNS = False

    n_off, n_on = len(off), len(on)
    b_on = sum(1 for _, b in on if b)
    light = "".join(x for x, b in on if not b).strip()
    ok_off = n_off == 1 and off[0][1] is True
    ok_on = n_on > 1 and b_on == 1 and "подсказка" in next(
        (x for x, b in on if b), "") and "отбор" in light
    if not (ok_off and ok_on):
        BAD.append(
            f"[клетка-прогоны] выключено — прогонов {n_off} "
            f"(ждали 1, полужирный); включено — прогонов {n_on}, "
            f"полужирных {b_on} (ждали >1 и ровно 1)")
    else:
        OK.append("[клетка-прогоны]")
        say(f"{'[клетка-прогоны]':<24} разностно: ЛОВИТ  →  выключено — "
            f"прогон {n_off}, клетка полужирная целиком (вид Лекции 6 не "
            f"двинулся); включено — прогонов {n_on}, полужирный из них "
            f"{b_on}, светлым идёт «{light}»")


def probe_render_path():
    """Пути рендера абсолютны. Доказывается на СЛОМАННОМ входе: подаётся
    относительный путь — ровно тот случай, на котором рендер и ломался.

    Обе команды делают `cd` в папку деки, поэтому относительный путь после
    `cd` указывает уже не туда, куда указывал у вызывающего: PDF уезжал в
    `<папка>/<папка>/`, а `pdftoppm` падал «Couldn't open file». Проба не
    запускает LibreOffice — она смотрит строки команд, которые его зовут."""
    import os
    from . import render as R
    rel_pptx, rel_snaps = Path("deck/lec-07-karkas.pptx"), Path("look/snaps")
    bad_cmds = R.render_cmds(rel_pptx, rel_snaps, 110, "")
    good_cmds = R.render_cmds(rel_pptx.absolute(), rel_snaps.absolute(), 110, "")
    here = os.getcwd().rstrip("/")
    # сломанный вход: папка в `cd` относительная, и `--outdir` повторяет её же
    broke = ('cd "deck"' in bad_cmds[0]
             and '--outdir "deck"' in bad_cmds[0]
             and f'"{here}' not in bad_cmds[1])
    fixed = (f'cd "{here}/deck"' in good_cmds[0]
             and f'--outdir "{here}/deck"' in good_cmds[0]
             and f'"{here}/look/snaps/k"' in good_cmds[1])
    if not (broke and fixed):
        BAD.append(
            f"[рендер-путь] проба не различает: на относительном входе "
            f"сломанность {broke}, на абсолютном исправность {fixed}")
    else:
        OK.append("[рендер-путь]")
        say(f"{'[рендер-путь]':<24} на сломанном входе: ЛОВИТ  →  "
            f"относительный путь даёт `cd \"deck\"` и `--outdir \"deck\"` "
            f"после того же `cd`, то есть PDF в `deck/deck/`, а снимки мимо; "
            f"абсолютный кладёт и PDF, и снимки туда, куда просили")


if __name__ == "__main__":
    sys.exit(selftest())
