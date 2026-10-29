#!/usr/bin/env python3
"""Адрес текстом без настоящей гиперссылки — находка, которую не мог найти
ни один сторож этой деки, потому что ни один из них не смотрел на ссылки.

Зачем. CLAUDE.md, «Форма подачи слайда»: «Источник на слайде и каждый URL в
заметках — настоящая гиперссылка (`run.hyperlink.address`)… Проверка —
`r:id` в XML сохранённого файла; самоотчёт агента не принимается.» Правило
старое, а сторожа на него не было ни одного (issue #299) — и молчало оно по
структурной причине: `tools/seminar-build/deck_kit.py::write_notes` до этой
правки вообще не умел разбирать `[текст](адрес)` — он печатал его буквально,
со скобками, и `a:hlinkClick`/`r:id` в XML не возникал нигде
(`notes/mcp-limitations.md` [#299-4]). Инструмент починен в том же круге,
этот файл — сторож, которого к нему не было.

Что здесь считается. Собранный `.pptx` читается ДВАЖДЫ, по двум разным
источникам правды:

* **адрес текстом** — `python-pptx`, прогон за прогоном (`run.text`), в
  теле каждого кадра и в заметках докладчика к нему. Если у прогона с
  адресом внутри нет `run.hyperlink.address` — адрес МЁРТВ: он на экране, но
  никуда не ведёт.
* **настоящая гиперссылка** — `r:id` у `<a:hlinkClick>` в XML СОХРАНЁННОГО
  файла, у которого в `.rels` того же листа лежит отношение
  `.../hyperlink`. Это и есть «проверка по `r:id`», а не самоотчёт о том,
  что вызывался `.hyperlink.address = url` где-то в коде сборки.

ПОЧЕМУ НЕ ПРОСТО `grep -c http` ПО ФАЙЛУ ЦЕЛИКОМ. У КАЖДОГО `slideN.xml` и
`notesSlideN.xml` в объявлении корня лежат три `xmlns:*="http://schemas.
openxmlformats.org/…"` — такой же `http`, как настоящий адрес в тексте.
На деке Семинара 9 (92 слайда с заметками) это даёт РОВНО 92×3 = 276 —
число, которое и привело к этому тикету как «276 адресов в заметках». Прямая
проверка (эта) нашла в тех же заметках ОДИН настоящий адрес, не 276: разница
— не опечатка, а разница между «посчитать `<a:t>`» и «посчитать байты
файла». Поэтому здесь текст читается через прогоны `python-pptx`, а не через
строку XML целиком — и это не перестраховка, это и есть находка, ради
которой переписан замер.

    python3 check_hyperlinks_pptx.py sem-09.pptx
    python3 check_hyperlinks_pptx.py --self-test

Код возврата 1, если найден хотя бы один мёртвый адрес.
"""
import re
import sys
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation

HERE = Path(__file__).resolve().parent

URL = re.compile(r"(?:https?://|www\.)[^\s<>\"«»'()\[\]]+", re.I)
_TAIL = ".,;:!?…»\"'"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def _bare(u):
    return u.rstrip(_TAIL)


def xml_link_counts(path):
    """Настоящие гиперссылки — `r:id` у `<a:hlinkClick>` в XML сохранённого
    файла, у которого в `.rels` ЭТОГО ЖЕ листа лежит отношение
    `.../hyperlink` (ссылка на другой слайд той же деки — другое отношение
    и в счёт не идёт). Раздельно тело/заметки. Разбор — тот же, что
    `tools/presentation-build/deckkit/render.py::verify_links` делает для
    лекций: правило одно на оба конвейера, и проверка должна быть одна же,
    а не второй автор той же регулярки."""
    body = notes = 0
    with zipfile.ZipFile(path) as z:
        rel_ok = {}
        for nm in z.namelist():
            if not nm.endswith(".rels"):
                continue
            root = etree.fromstring(z.read(nm))
            rel_ok[nm] = {e.get("Id") for e in root
                         if e.get("Type", "").endswith("/hyperlink")}
        for nm in z.namelist():
            if not nm.endswith(".xml") or "_rels" in nm:
                continue
            if nm.startswith("ppt/slides/slide"):
                bucket = "body"
            elif nm.startswith("ppt/notesSlides/notesSlide"):
                bucket = "notes"
            else:
                continue
            folder, fname = nm.rsplit("/", 1)
            ok = rel_ok.get(f"{folder}/_rels/{fname}.rels", set())
            root = etree.fromstring(z.read(nm))
            n = sum(1 for e in root.iter(f"{A}hlinkClick") if e.get(f"{R}id") in ok)
            if bucket == "body":
                body += n
            else:
                notes += n
    return body, notes


def find_dead(path):
    """([(место, адрес), …] мёртвых, адресов в теле, адресов в заметках).

    «Мёртвый» — прогон содержит адрес текстом, и у ЭТОГО ЖЕ прогона нет
    `run.hyperlink.address`. Якорь ссылки необязан совпадать с адресом
    (markdown-ссылка `[Название](url)` показывает название, не сам url) —
    поэтому дохлый адрес ищется там, где текст прогона И ЕСТЬ адрес, а не
    по совпадению счёта адресов со счётом ссылок."""
    prs = Presentation(str(path))
    dead, body_urls, notes_urls = [], 0, 0
    for i, sl in enumerate(prs.slides, start=1):
        for shp in sl.shapes:
            if not shp.has_text_frame:
                continue
            for p in shp.text_frame.paragraphs:
                for r in p.runs:
                    for u in URL.findall(r.text):
                        body_urls += 1
                        if not r.hyperlink.address:
                            dead.append((f"тело s{i:02d}", _bare(u)))
        if sl.has_notes_slide:
            for p in sl.notes_slide.notes_text_frame.paragraphs:
                for r in p.runs:
                    for u in URL.findall(r.text):
                        notes_urls += 1
                        if not r.hyperlink.address:
                            dead.append((f"заметки s{i:02d}", _bare(u)))
    return dead, body_urls, notes_urls


def check(path, *, quiet=False):
    dead, body_urls, notes_urls = find_dead(path)
    link_body, link_notes = xml_link_counts(path)
    if not quiet:
        print(f"адресов текстом: тело {body_urls} · заметки {notes_urls}")
        print(f"гиперссылок (r:id в XML сохранённого файла): "
              f"тело {link_body} · заметки {link_notes}")
        if dead:
            print(f"⚠ МЁРТВЫХ АДРЕСОВ (текст без гиперссылки): {len(dead)}")
            for where, u in dead:
                print(f"  ⚠ {where}: {u}")
        else:
            print("· все адреса текстом привязаны к настоящей гиперссылке")
    return dead


# ── Самопроверка на нарочно сломанном и нарочно годном входе ───────────────
#
# Проба кормит сторож не фикстурой, подогнанной под регулярку, а ДВУМЯ
# настоящими `.pptx`, собранными по двум разным путям:
#
#   * `_probe_dead` пишет адрес текстом НАПРЯМУЮ в `text_frame`/
#     `notes_text_frame`, минуя `deck_kit` целиком — ровно то, что умел
#     прежний `write_notes` (и сырой `add_textbox`) до починки: адрес на
#     экране есть, `r.hyperlink.address` нет.
#   * `_probe_good` собирает тот же смысл через `K.text_box`/`K.write_notes`
#     ПОСЛЕ починки — markdown-ссылкой и голым `https://…`.
#
# Если починку `deck_kit` откатят, вторая проба перестанет быть «чистой» и
# самопроверка упадёт — она не только ловит старый баг, но и держит новый
# код, который его закрыл.


def _probe_dead(out_pptx):
    from pptx import Presentation as P
    from pptx.util import Inches
    sys.path.insert(0, str(HERE))
    import deck_kit as K
    prs = P()
    prs.slide_width, prs.slide_height = Inches(K.W_IN), Inches(K.H_IN)
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    K.set_bg(sl, K.WHITE)
    tb = sl.shapes.add_textbox(Inches(1), Inches(1), Inches(10), Inches(1))
    tb.text_frame.text = "Источник: https://example.org/dead-body-link"
    sl.notes_slide.notes_text_frame.text = (
        "источник — Thinking Machines Lab, https://example.org/dead-notes-link.")
    prs.save(out_pptx)


def _probe_good(out_pptx):
    from pptx import Presentation as P
    from pptx.util import Inches
    sys.path.insert(0, str(HERE))
    import deck_kit as K
    prs = P()
    prs.slide_width, prs.slide_height = Inches(K.W_IN), Inches(K.H_IN)
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    K.set_bg(sl, K.WHITE)
    K.text_box(sl, 1.0, 1.0, 10.0, 1.0,
               ["Источник: [Пример источника](https://example.org/good-body-link)"],
               size=14)
    K.write_notes(sl, "источник — [Пример](https://example.org/good-notes-link) "
                       "и голый https://example.org/good-notes-bare.")
    prs.save(out_pptx)


def self_test():
    import tempfile
    ok = True
    with tempfile.TemporaryDirectory() as d:
        d = Path(d)

        dead_pptx = d / "probe-dead.pptx"
        _probe_dead(dead_pptx)
        dead, bu, nu = find_dead(dead_pptx)
        hit = len(dead) == 2   # один в теле, один в заметках
        ok = ok and hit
        mark = "✓" if hit else "✗ ПРОВЕРКА МЁРТВАЯ"
        print(f"{mark} ловит адрес без ссылки (ожидали 2, нашли {len(dead)}): {dead}")

        good_pptx = d / "probe-good.pptx"
        _probe_good(good_pptx)
        dead2, bu2, nu2 = find_dead(good_pptx)
        clean = not dead2
        ok = ok and clean
        mark = "✓" if clean else "✗ ЛОЖНОЕ СРАБАТЫВАНИЕ"
        print(f"{mark} не ругается на настоящую ссылку "
              f"(адресов без ссылки: {len(dead2)}): {dead2}")

        link_body, link_notes = xml_link_counts(good_pptx)
        links_ok = link_body == 1 and link_notes == 2
        ok = ok and links_ok
        mark = "✓" if links_ok else "✗ ПРОВЕРКА МЁРТВАЯ"
        print(f"{mark} считает r:id в XML (тело 1, заметки 2) — "
              f"нашли тело {link_body}, заметки {link_notes}")
    return 0 if ok else 1


def main():
    args = sys.argv[1:]
    if not args or args[0] == "--self-test":
        return self_test()
    path = Path(args[0])
    if not path.exists():
        print(f"файла нет: {path}")
        return 2
    dead = check(path)
    print(f"\n— мёртвых адресов: {len(dead)}")
    return 1 if dead else 0


if __name__ == "__main__":
    sys.exit(main())
