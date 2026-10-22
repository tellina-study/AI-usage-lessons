"""Роль блока и его отрисовка: что за работу несёт этот кусок текста.

Одна работа слайда — одна форма (канон `seminar-deck-kit.md` §1). Здесь
решается, какая именно: вопрос залу (золотая коробка, редкая), формула-итог
(золотая планка без заливки), реплика докладчика (тихая тиловая линия),
оговорка (пунктир), факт (тиловая отметка). Решение принимается ПО ТЕКСТУ и
по приёму слайда, а не по номеру блока — поэтому один и тот же механизм
работает на любом семинаре, а занятие настраивает только таблицу
«приём → роль по умолчанию» в своём конфиге.

Сюда же собрана починка разметки на входе (`sanitize_blocks`): одиночный
курсив, `**жирный**`, рассечённый переносом строки исходника, и карточки-
варианты, записанные несколькими строками. Это не правка кита
(`deck_kit.py`/`slide_parts.py` остаются как есть) — это снятие того, что
вёрстка не умеет, ДО входа в вёрстку. Замерено на принятой деке Семинара 6:
ни один из трёх механизмов не меняет там ни одного блока, поэтому они
работают всегда, без флага в конфиге.
"""

import re

from pptx import Presentation
from pptx.util import Inches

import deck_config as D
import deck_kit as K
import metrics as M

FORM = {"question": K.form_question, "formula": K.form_formula,
        "speech": K.form_speech, "caveat": K.form_caveat, "fact": K.form_fact,
        "footnote": K.form_footnote}

ITALIC_LINE = re.compile(r"^\s*\*([^*].*?)\*\s*$")


def forced_role(force, sid, quote_i):
    """Роль, назначенную деке вручную, — для БЛОКА, а не для всего слайда.

    `quote_role_force` в конфиге занятия принимает две записи:

        quote_role_force: {n26: formula}          — всем цитатам слайда
        quote_role_force: {n20: [footnote, null]} — по блокам, по порядку

    Первая форма — прежняя, и она осталась дословно: слайд, у которого ВСЕ
    цитаты несут одну работу, не должен обрастать списком из одного элемента.
    Вторая появилась, когда на одном экране впервые встретились две разные
    работы, обе из которых общий механизм определял одинаково: сноска-глосс
    под таблицей и вывод-находка под ней (n20 Семинара 8). `null` в списке —
    «эту цитату решает общий механизм», и список короче числа цитат не
    ошибка: хвост просто решается как прежде."""
    r = force.get(sid)
    if isinstance(r, (list, tuple)):
        return r[quote_i] if 0 <= quote_i < len(r) else None
    return r


def opens_speech(body):
    return body.lstrip().startswith(D.SPEECH_OPEN)


def quote_role(lines, pattern):
    body = " ".join(K.plain(l) for l in lines).strip()
    if body.rstrip("»\"' ").endswith("?") or D.ASK.search(body):
        return "question"
    if D.CAVEAT_OPEN.match(K.plain(body).lstrip("*> ")) or D.CAVEAT.search(body):
        return "caveat"
    base = D.pattern_role().get(pattern)
    if base == "question":
        return "speech" if opens_speech(body) else "fact"
    if opens_speech(body):
        return "speech"
    return base or "formula"


def para_role(text, pattern):
    m = ITALIC_LINE.match(text)
    if m:
        return "speech", m.group(1).strip()
    return quote_role([text], pattern), text


def split_question(lines):
    """Отделить сцену от вопроса внутри одного блока-цитаты: зал должен
    увидеть вопрос отдельной золотой коробкой, а не последней фразой абзаца."""
    flat = " ".join(lines).strip()
    op, cl = "", ""
    for o, c in D.WRAP_PAIRS:
        if not (flat.startswith(o) and flat.endswith(c) and len(flat) > len(o) + len(c)):
            continue
        if o == c and o in flat[len(o):len(flat) - len(c)]:
            continue
        op, cl = o, c
        break
    core = flat[len(op):len(flat) - len(cl)].strip() if op else flat

    bounds, pos = [], 0
    for m in re.finditer(r"[.!?…]+[»\"\')\s]*", core):
        bounds.append((pos, m.end()))
        pos = m.end()
    if pos < len(core):
        bounds.append((pos, len(core)))

    def rewrap(text):
        text = text.strip()
        return f"{op}{text}{cl}" if op and text else text

    for a, b in bounds:
        sent = core[a:b]
        if "?" in sent or D.ASK.search(sent):
            if a == 0:
                return [], [rewrap(core)], True
            return [rewrap(core[:a])], [rewrap(core[a:])], True
    return [], lines, False


def unseen(sid, pattern, blocks, handled, why):
    for kind, _b in blocks:
        if kind not in handled:
            M._WARNINGS.append(f"БЛОК НЕ ПОКАЗАН [{sid} · {pattern}]: блок вида "
                               f"«{kind}» — {why}")


# ── Починка разметки на входе ───────────────────────────────────────────────
#
# Одиночный курсив `*…*` вёрстка не умеет: `metrics.fits()` честно печатает
# предупреждение («звёздочки выйдут на слайд как есть»), но только НАЗЫВАЕТ
# находку. Срезать звёздочки в ките нельзя — так прежняя версия теряла
# целевые ответы, помеченные `**жирным**`; поэтому снимается здесь, тем же
# шаблоном, которым кит находку обнаруживает.
_LONE_ITALIC = re.compile(r"(?<![*\w])\*(?!\*)([^*\n]{2,}?)\*(?![*\w])")


def _degloss_italic(s):
    if not isinstance(s, str) or "*" not in s:
        return s
    return _LONE_ITALIC.sub(lambda m: m.group(1), s)


# `slide_parts.blocks()` собирает блок-цитату построчно, и каждая ФИЗИЧЕСКАЯ
# строка исходника рисуется отдельным абзацем PowerPoint. Если автор перенёс
# строку посередине `**жирного**` фрагмента (обычный перенос по ширине
# редактора, не смысловая граница), открывающая `**` остаётся без пары на
# своей строке, закрывающая — на следующей, и обе пары звёздочек выходят на
# слайд буквальным текстом. Чётность считается НАКОПЛЕННО, а не впритык по
# паре строк: так находится и фрагмент, растянутый больше чем на один перенос.
_BOLD_MARK = re.compile(r"\*\*")


def _rejoin_split_bold(lines):
    out, i = [], 0
    while i < len(lines):
        cur = lines[i]
        while len(_BOLD_MARK.findall(cur)) % 2 == 1 and i + 1 < len(lines):
            i += 1
            cur = cur + " " + lines[i]
        out.append(cur)
        i += 1
    return out


# Ряд карточек `slide_parts` опознаёт по строке, у которой НА НЕЙ САМОЙ есть
# «·» и ≥4 бэктика. Автор, записавший варианты в несколько строк ради
# читаемости исходника, этот признак не зажигает — и голосование вслепую
# рисуется прозой. Тот же признак применяется здесь на двух уровнях: внутри
# одного `para`-блока и слиянием соседних `para`-фрагментов, пока не
# наберётся порог «минимум 3 варианта» (ниже порога это не голосование, а
# случайное совпадение разметки).
_CARD_SEG = re.compile(r"`([^`]+)`")


def _card_fragment(text):
    """Сегменты карточек в ОДНОМ `para`-блоке, если блок — ЧИСТЫЙ фрагмент
    списка (после вычитания `` `сегментов` `` остаются только «·» и
    пробелы), иначе `None`. Без строгого «ничего кроме» обычный абзац со
    случайным кодовым словом в обратных кавычках прошёл бы тоже."""
    if "`" not in text or "·" not in text:
        return None
    segs = _CARD_SEG.findall(text)
    rest = _CARD_SEG.sub("", text)
    if segs and re.fullmatch(r"[\s·]*", rest):
        return segs
    return None


def _merge_card_fragments(blocks):
    out, i = [], 0
    while i < len(blocks):
        kind, b = blocks[i]
        frag = _card_fragment(b) if kind == "para" else None
        if frag is not None:
            segs, j = list(frag), i + 1
            while j < len(blocks) and blocks[j][0] == "para":
                nxt = _card_fragment(blocks[j][1])
                if nxt is None:
                    break
                segs.extend(nxt)
                j += 1
            if len(segs) >= 3:
                out.append(("cards", segs))
                i = j
                continue
        out.append((kind, b))
        i += 1
    return out


def sanitize_blocks(blocks):
    """Снять одиночный курсив из видимого слоя, слить строки блока-цитаты с
    незакрытым `**жирным**` и собрать рассечённые карточки-варианты.
    `code`-блоки НЕ трогаются: внутри листинга `*` часто настоящий знак
    (`*args`, глоб-шаблон), и срезать его значило бы показать код, который
    автор не писал."""
    out = []
    for kind, b in blocks:
        if kind == "quote":
            out.append((kind, [_degloss_italic(l) for l in _rejoin_split_bold(b)]))
        elif kind == "para":
            out.append((kind, _degloss_italic(b)))
        elif kind == "bullets":
            items, numbered = (b, False) if isinstance(b, list) else b
            out.append((kind, ([_degloss_italic(i) for i in items], numbered)))
        elif kind == "table":
            headers, rows = b
            out.append((kind, ([_degloss_italic(h) for h in headers],
                               [[_degloss_italic(c) for c in r] for r in rows])))
        else:
            out.append((kind, b))
    return _merge_card_fragments(out)


# ── Отрисовка одного блока ──────────────────────────────────────────────────

def _is_variant_table(headers, rows):
    """Таблица-вариант без второй колонки (`| Вариант | |`): автор записал
    список вариантов ответа markdown-таблицей из двух колонок, вторая пустая
    у всех строк кроме, может быть, одной (пометка верного ответа). По сути
    это те же карточки-варианты, что `option_row` рисует ряду «ВЫБЕРИ» —
    разбирать её как таблицу значило бы рисовать пустые слоты во всю высоту
    строки под каждой карточкой. Здесь распознавание формы, не новый приём."""
    if len(headers) != 2:
        return False
    return K.plain(headers[0]).strip().lower().startswith("вариант")


def _table_style(pattern):
    """Стиль таблицы для приёма: цветной акцент рамки и, если нужно,
    отключённая автоподсветка или отключённые слоты пустых ячеек. Осознанные
    исключения из стандартной рамки `#1C7293` — решение ДЕКИ, а не кода,
    поэтому живут в её конфиге (`table_style`), а не в ветвлении по имени
    приёма здесь.

    `slots: false` — пустая ячейка остаётся пустой, без серой штриховой
    плашки. Нужно там, где пустота значит «писать нечего» (итоговая строка
    под колонкой примеров), а не «зал заполнит по ходу занятия»; подробнее —
    `K.table_card`."""
    st = D.table("table_style").get(pattern) or {}
    accent = getattr(K, st["accent"]) if st.get("accent") else None
    kw = {"accent": accent}
    if st.get("highlight") == "none":
        kw["highlight"] = {}
    if st.get("slots") is False:
        kw["slots"] = False
    return kw, st.get("label", "таблица")


def block_drawer(kind, b, sid, pattern, *, role=None):
    if kind == "quote":
        r = role or quote_role(b, pattern)
        fn = FORM[r]
        return lambda sl, y, mh: fn(sl, D.LEFT, y, D.WIDTH, b, max_h=mh,
                                    label=f"{sid} {r}")
    if kind == "table":
        headers, rows = b
        if _is_variant_table(headers, rows):
            options = [r[0] for r in rows if r and r[0].strip()]
            hi = next((i for i, r in enumerate(rows)
                       if len(r) > 1 and r[1].strip()), None)
            return lambda sl, y, mh: K.option_row(
                sl, D.LEFT, y, D.WIDTH, options, highlight_idx=hi,
                max_h=min(mh or 1.8, 1.8), label=f"{sid} варианты")
        kw, label = _table_style(pattern)
        return lambda sl, y, mh: K.table_card(sl, D.LEFT, y, D.WIDTH, headers, rows,
                                              max_h=mh, label=f"{sid} {label}", **kw)
    if kind == "code":
        lang, lines = b
        return lambda sl, y, mh: K.terminal_card(sl, D.LEFT, y, D.WIDTH, lines,
                                                 max_h=mh,
                                                 markup=lang in ("markdown", "md"),
                                                 label=f"{sid} код")
    if kind == "cards":
        if pattern in D.vote_patterns():
            return lambda sl, y, mh: K.option_row(sl, D.LEFT, y, D.WIDTH, b,
                                                  max_h=min(mh or 1.7, 1.7),
                                                  label=f"{sid} варианты")
        return lambda sl, y, mh: K.term_pills(sl, D.LEFT, y, D.WIDTH, b,
                                              label=f"{sid} термины")
    if kind == "bullets":
        items, numbered = (b, False) if isinstance(b, list) else b
        return lambda sl, y, mh: K.numbered_list(sl, D.LEFT, y, D.WIDTH, items,
                                                 max_h=mh, numbered=numbered,
                                                 label=f"{sid} список")
    if kind == "para":
        r, text = para_role(b, pattern)
        fn = FORM[role or r]
        return lambda sl, y, mh: fn(sl, D.LEFT, y, D.WIDTH, [text], max_h=mh,
                                    label=f"{sid} {role or r}")
    M._WARNINGS.append(f"БЛОК НЕ ПОКАЗАН [{sid} · {pattern}]: блок вида «{kind}» "
                       f"ни один приём не рисует — содержание пропало бы молча")
    return None


# ── Композиция: воздух как параметр, а не остаток ───────────────────────────

_scratch = Presentation()
_scratch.slide_width, _scratch.slide_height = Inches(K.W_IN), Inches(K.H_IN)


def measure(fn, *a, **kw):
    """Замер делает тот же код, который рисует: блок рисуется на черновом
    слайде, высота запоминается, слайд выбрасывается. Разойтись замер и
    отрисовка тогда не могут (канон §3)."""
    sl = _scratch.slides.add_slide(_scratch.slide_layouts[6])
    saved = list(M._WARNINGS)
    h = fn(sl, *a, **kw)
    M._WARNINGS[:] = saved
    return h


def with_footnote(prev, note, gap=0.06):
    """Склеить блок и сноску к нему в ОДИН блок композиции.

    Сноска — не равноправный сосед, а приписка к тому, что над ней, и
    клеится она не ради красоты. `compose` выдаёт каждому блоку долю
    `max(h*k, 0.5)`, и отсечка 0,5″ на тесном слайде стоит не зря (без неё
    мелкому блоку досталась бы доля, в которую не влезает ни одна строка).
    Но для сноски высотой 0,32″ эта отсечка означает, что из общего бюджета
    она вынимает 0,5″, отдаёт обратно 0,18″ — и ровно на эту разницу
    композиция уходит ниже рабочего поля, сколько ни сокращай содержание
    слайда: таблица и схема упруги и мгновенно съедают всё освободившееся.
    Замерено на n20 Семинара 8: отдельным блоком сноска давала «НЕ
    ПОМЕСТИЛОСЬ» на 0,09″ и не чинилась ни сокращением таблицы, ни
    сокращением находки, ни понижением потолка схемы — каждый раз перебор
    возвращался. Приклеенная, она просто входит в долю своего блока.

    Чинить это внутри `compose` (пересчётом доли в два захода) пробовали
    первым делом; пересчёт работает, но сдвигает вид уже принятой деки
    Семинара 6 — её таблицу n39 с 11,5 pt до 11. Поэтому починка стоит
    здесь, где касается только тех дек, которые сноску заказали."""
    def draw(sl, y, mh):
        nh = measure(note, y, None)
        ph = prev(sl, y, None if mh is None else max(mh - nh - gap, 0.5))
        return ph + gap + note(sl, y + ph + gap, None)
    return draw


def compose(sl, sid, y0, drawers, *, bottom=None, center=True):
    if bottom is None:
        bottom = D.BOTTOM
    if not drawers:
        return
    avail = bottom - y0
    nat = [measure(d, y0, None) for d in drawers]
    total = sum(nat) + D.GAP * (len(nat) - 1)

    if total <= avail:
        slack = avail - total
        gap = D.GAP + min(slack / max(len(nat), 1), 0.55)
        used = sum(nat) + gap * (len(nat) - 1)
        kc = 0.8 if center is True else (0.0 if center is False else float(center))
        y = y0 + (avail - used) / 2 * kc
        for d, h in zip(drawers, nat):
            d(sl, y, None)
            y += h + gap
        return

    # Урезание. Нижняя отсечка 0,5″ в доле блока стоит не зря: без неё мелкий
    # блок на тесном слайде получил бы долю, в которую не влезает ни одна
    # строка. У неё есть известная цена — блок, которому хватает меньше
    # отсечки, всё равно занимает в бюджете 0,5″, и на слайде из четырёх
    # блоков, где два упираются в собственный пол (таблица на дне кегля,
    # золотая планка при 11 pt), композиция уходит ниже рабочего поля на эту
    # разницу. Пересчёт в два захода (закрепить тех, кому хватает natural, и
    # переделить остаток) эту цену снимает и на Семинаре 8 закрывал ровно её —
    # но на принятой деке Семинара 6 он же ужимал таблицу n39 с 11,5 pt до 11
    # и усиливал переполнение её золотой планки. Платить сменой вида принятой
    # деки за чужой слайд нельзя, поэтому пересчёт не внесён: там, где
    # отсечка создаёт перебор, он называется предупреждением ниже, и слайд
    # сокращают содержанием.
    budget = avail - D.GAP * (len(drawers) - 1)
    k = budget / sum(nat)
    y = y0
    for d, h in zip(drawers, nat):
        drawn = d(sl, y, max(h * k, 0.5))
        y += drawn + D.GAP

    end = y - D.GAP
    over = end - bottom
    if over > 0.02:
        if end > K.H_IN - 0.15:
            M._WARNINGS.append(
                f"ЗА КРАЕМ ПОЛОТНА [{sid}]: композиция кончается на {end:.2f}″ при "
                f"высоте полотна {K.H_IN:.2f}″ — нижний блок в PowerPoint не виден вовсе")
        else:
            M._WARNINGS.append(
                f"НЕ ПОМЕСТИЛОСЬ [{sid}]: композиция на {over:.2f}″ ниже рабочего поля — "
                f"блок заходит в поле подписи, номер слайда ляжет поверх него")
