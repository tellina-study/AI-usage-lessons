# -*- coding: utf-8 -*-
"""ПОЛ КОНТРАСТА ВИДИМОГО СЛОЯ — читается ли текст, а не только крупен ли он.

ЧТО БЫЛО СЛОМАНО. В этом пакете есть пол КЕГЛЯ (`floor.py`) и нет пола
ЦВЕТА — и прожарка 2026-10-14 поймала ровно в этой щели дефект P0:
ключевое число деки Лекции 7 (значение «key»-плашки) набрано золотом
`#F0AB00` 23 пт на белой заливке плашки, контраст 1,99 : 1 при пороге 3,0 —
и так на 60 кадрах из 109. Это самый важный текст кадра, набранный самым
нечитаемым цветом деки.

Дефект прошёл ВСЕ сторожа, и не по недосмотру, а ПО ПОСТРОЕНИЮ: пол кегля
меряет кегль, мера высоты — высоту, регрессия — отпечаток XML, `[золото]`
проверяет, что золотой акцент на кадре ЕСТЬ. Ни одна проверка пакета не
знала слова «контраст», поэтому 60 кадров могли жить так сколько угодно
кругов. Прошлый круг это уже показал в другом носителе: оговорку под числом
подняли с 10,5 до 12 пт — и нечитаемость переехала с оговорки на само
число. Класс дефекта один: «с задних рядов этого не видно», и кегль
закрывает только его половину.

ПОЧЕМУ ЭТОТ МОДУЛЬ НАЗЫВАЕТ, А НЕ ЧИНИТ. Договор тот же, что у `floor.py`,
и по той же причине. У контраста ДВА способа починки — опустить цвет текста
(золото → `1C7293`/`028090`) или перевернуть плашку (заливка `21295C`,
число золотом, 6,85 : 1), — и они дают РАЗНЫЙ кадр: второй добавляет
иерархию заливкой, которую видно раньше, чем прочитан текст. Выбор между
ними — решение владельца, не сборки. Поэтому здесь только мера: кадр,
фигура, цвет, фон, отношение, вердикт.

ПОРОГ КУРСА РАВЕН ПОРОГУ WCAG; СТРОЖЕ НЕ СТАВИМ, ПОТОМУ ЧТО НА ПРОЕКЦИИ НЕ
МЕРЕНО. Равенство здесь — решение владельца (2026-10-14), а не недосмотр и не
значение по умолчанию: деку смотрят с задних рядов на проекторе с посаженной
лампой, то есть в условиях ХУЖЕ экранных, под которые писан WCAG, и строгий
порог был бы оправдан — но порог, которого на этом материале никто не мерил,
заменил бы одну догадку другой. Поднимать его можно после замера на настоящем
экране; до него числа ниже — измеренный стандарт, и ничего больше.

ПОЧЕМУ МЕРА ИДЁТ ПО СОБРАННОМУ `.pptx`, А НЕ ПО НАМЕРЕНИЮ СБОРЩИКА. Правило
`proverki-i-pravila.md`: «инструмент, через который смотрят, обязан быть тем
же, который собирает продукт». Цвет прогона и заливка фигуры в момент
вызова примитива — это намерение; в XML уезжает результат, и фон у надписи
вообще НЕ её свойство: `add_textbox` заливки не имеет, а видимый фон даёт
фигура под ней (коробка мотива, плашка, фон кадра). Поэтому фон считается
по z-порядку собранного файла — тем же способом, каким его видит
LibreOffice.
"""
import sys

from pptx import Presentation
from pptx.enum.dml import MSO_FILL
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

from . import report
from .theme import SLIDE_H, SLIDE_W

# ── пороги WCAG 2.x, SC 1.4.3 «Contrast (Minimum)», уровень AA ───────
# Откуда числа: W3C, Web Content Accessibility Guidelines 2.1, Success
# Criterion 1.4.3 — обычный текст не ниже 4,5 : 1, КРУПНЫЙ не ниже 3,0 : 1.
# «Крупный» там же определён как ≥18 пт, либо ≥14 пт полужирный (WCAG
# считает в CSS-пунктах; кегль в PPTX — те же пункты, 1/72 дюйма).
#
# Почему равен порогу WCAG, а не строже — в шапке модуля: решение владельца,
# а не значение по умолчанию.
NORMAL_RATIO = 4.5      # обычный текст
LARGE_RATIO = 3.0       # крупный текст
LARGE_PT = 18.0         # с этого кегля текст «крупный»
LARGE_BOLD_PT = 14.0    # полужирный — с этого

EPS = 1e-9              # ровно на пороге — ПРОХОДИТ, см. `passes`

WHITE_HEX = "FFFFFF"


def _run_word(n):
    """«прогон / прогона / прогонов» по числу. Мелочь, но отчёт сборки
    читают вслух на разборе, и «61 прогонов» в нём звенит."""
    n = abs(int(n))
    if n % 100 in range(11, 15):
        return "прогонов"
    return {1: "прогон", 2: "прогона", 3: "прогона", 4: "прогона"}.get(
        n % 10, "прогонов")


def _page_word(n):
    n = abs(int(n))
    if n % 100 in range(11, 15):
        return "кадрах"
    return "кадре" if n % 10 == 1 else "кадрах"


def lum(rgb):
    """Относительная яркость по WCAG 2.x. `rgb` — (r, g, b) 0…255."""
    out = 0.0
    for c, k in zip(rgb, (0.2126, 0.7152, 0.0722)):
        c /= 255.0
        out += k * (c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4)
    return out


def hex_rgb(h):
    h = str(h).strip().lstrip("#").upper()
    return int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)


def ratio(fg, bg):
    """Контраст пары: (L1 + 0,05) / (L2 + 0,05), где L1 — светлейшая.

    Порядок цветов в паре не имеет значения — отношение симметрично, и
    поэтому «текст на фоне» и «фон под текстом» дают одно число."""
    a, b = lum(hex_rgb(fg)), lum(hex_rgb(bg))
    hi, lo = max(a, b), min(a, b)
    return (hi + 0.05) / (lo + 0.05)


def is_large(size, bold):
    """Крупный ли текст по WCAG: ≥18 пт, либо ≥14 пт полужирный."""
    if size is None:
        return False
    return size >= (LARGE_BOLD_PT if bold else LARGE_PT) - EPS


def threshold(size, bold):
    return LARGE_RATIO if is_large(size, bold) else NORMAL_RATIO


def passes(r, thr):
    """Ровно на пороге — ПРОХОДИТ. WCAG пишет «не ниже», и это не придирка:
    без допуска пара, посчитанная в 3,0000000001 и в 2,9999999999, попадала
    бы в разные разряды от последнего бита двойной точности."""
    return r >= thr - EPS


# ── справка по палитре курса ─────────────────────────────────────────
# Пять пар, которыми живёт дека. Считаются КАЖДЫЙ прогон, а не вписаны
# числами: вписанное число не ломается, когда ломается формула, — и
# расхождение с независимым замером владельца тогда не находится ничем.
PALETTE_PAIRS = (
    ("золото F0AB00 на белом", "F0AB00", WHITE_HEX, 1.99),
    ("бирюза 028090 на белом", "028090", WHITE_HEX, 4.67),
    ("средний 1C7293 на белом", "1C7293", WHITE_HEX, 5.42),
    ("глубокий 21295C на белом", "21295C", WHITE_HEX, 13.67),
    ("золото F0AB00 на глубоком 21295C", "F0AB00", "21295C", 6.85),
)
PALETTE_TOL = 0.05      # допуск сверки с независимым замером владельца


def palette_lines():
    """Справка: пять пар палитры, мой счёт против независимого замера.

    Расхождение больше `PALETTE_TOL` — это НАХОДКА, а не повод подогнать
    формулу: тогда говорится, какая формула даёт какое число."""
    out = ["── палитра курса (`theme.py`) против независимого замера ─────",
           f"{'пара':<36} {'мой счёт':>9} {'замер':>7} {'расх.':>7}  порог"]
    worst = 0.0
    for name, fg, bg, want in PALETTE_PAIRS:
        r = ratio(fg, bg)
        d = abs(r - want)
        worst = max(worst, d)
        mark = "" if d <= PALETTE_TOL else "   ← РАСХОЖДЕНИЕ"
        thr = f"крупный {LARGE_RATIO:.1f} · обычный {NORMAL_RATIO:.1f}"
        out.append(f"{name:<36} {r:>9.4f} {want:>7.2f} {d:>7.4f}  {thr}{mark}")
    out.append(f"наибольшее расхождение {worst:.4f} при допуске "
               f"{PALETTE_TOL:.2f} — "
               + ("сходится" if worst <= PALETTE_TOL else "НЕ СХОДИТСЯ"))
    return out


# ── чтение собранного файла ──────────────────────────────────────────


def _rect(shp):
    """Рамка фигуры в дюймах или None, если геометрии нет."""
    try:
        if None in (shp.left, shp.top, shp.width, shp.height):
            return None
        return (Emu(shp.left).inches, Emu(shp.top).inches,
                Emu(shp.left + shp.width).inches,
                Emu(shp.top + shp.height).inches)
    except Exception:                                    # noqa: BLE001
        return None


def _solid(shp):
    """Сплошная заливка фигуры в hex или None (прозрачная, иная, ошибка)."""
    try:
        if shp.fill.type == MSO_FILL.SOLID:
            return str(shp.fill.fore_color.rgb)
    except Exception:                                    # noqa: BLE001
        return None
    return None


def _covers(rc, x, y):
    return rc is not None and rc[0] <= x <= rc[2] and rc[1] <= y <= rc[3]


def background(shapes, i):
    """ФАКТИЧЕСКАЯ заливка под фигурой `i`, как её видит рендер.

    Порядок, и каждый шаг — наблюдаемое свойство файла, а не догадка:

    1. СВОЯ сплошная заливка фигуры (чип, клетка таблицы с текстом внутри).
    2. Прозрачная своя — идём по z-порядку ВНИЗ до первой фигуры, накрывающей
       ЦЕНТР нашей. Сплошная заливка — она и есть фон (коробка мотива,
       плашка, фон кадра). Снимок — фон НЕ ИЗМЕРЕН и так и сказано: цвет
       пикселя под буквой из XML не читается, а выдать белый за фон снимка
       значило бы соврать тихо.
    3. Прозрачная фигура сквозь — пропускается: сквозь неё видно нижнюю.
    4. Ничего не нашлось — белый, и это помечено («по умолчанию»), чтобы
       предположение не читалось как замер.

    Центр, а не вся рамка, нарочно: надпись часто чуть шире коробки мотива
    (у `refdeck` коробка начинается на 0,14" левее текста, у плашек наоборот),
    и требование «накрывает рамку целиком» отправляло бы в «по умолчанию»
    ровно те случаи, где фон очевиден. Цена допущения названа: прогон,
    который НАЧИНАЕТСЯ на плашке и уезжает с неё, мерится по плашке.
    """
    own = _solid(shapes[i])
    if own:
        return own, "своя заливка"
    rc = _rect(shapes[i])
    if rc is None:
        return WHITE_HEX, "белый по умолчанию"
    cx, cy = (rc[0] + rc[2]) / 2.0, (rc[1] + rc[3]) / 2.0
    for j in range(i - 1, -1, -1):
        under = shapes[j]
        box = _rect(under)
        if not _covers(box, cx, cy):
            continue
        if under.shape_type == MSO_SHAPE_TYPE.PICTURE:
            return None, "снимок — не измерен"
        fill = _solid(under)
        if fill:
            whole = (box[0] <= 0.01 and box[1] <= 0.01
                     and box[2] >= SLIDE_W - 0.01 and box[3] >= SLIDE_H - 0.01)
            return fill, ("фон кадра" if whole else f"заливка #{j}")
    return WHITE_HEX, "белый по умолчанию"


def _run_color(r):
    try:
        if r.font.color.type is None:
            return None
        return str(r.font.color.rgb)
    except Exception:                                    # noqa: BLE001
        return None


def rows_of(prs):
    """Все текстовые прогоны деки с фоном, контрастом и вердиктом.

    Прогон без НАЗВАННОГО цвета или кегля вердикта не получает: цвет там
    наследуется от темы, а тема у деки этого курса не задана — то есть
    измерять было бы нечего, и «прошёл» в таком случае значил бы только
    «мы не посмотрели»."""
    out = []
    for n, sl in enumerate(prs.slides, 1):
        shapes = list(sl.shapes)
        for i, shp in enumerate(shapes):
            if not shp.has_text_frame:
                continue
            bg, src = background(shapes, i)
            for para in shp.text_frame.paragraphs:
                for r in para.runs:
                    if not r.text.strip():
                        continue
                    fg = _run_color(r)
                    size = r.font.size.pt if r.font.size else None
                    bold = bool(r.font.bold)
                    row = {"page": n, "shape": f"#{i} {shp.name}",
                           "text": r.text.strip()[:40], "size": size,
                           "bold": bold, "fg": fg, "bg": bg, "bg_src": src,
                           "ratio": None, "thr": None, "ok": None}
                    if fg and bg and size is not None:
                        row["ratio"] = ratio(fg, bg)
                        row["thr"] = threshold(size, bold)
                        row["ok"] = passes(row["ratio"], row["thr"])
                    out.append(row)
    return out


def _cls(row):
    if row["size"] is None:
        return "кегль не назван"
    return ("крупный" if is_large(row["size"], row["bold"]) else "обычный")


def table_lines(rows, *, only_bad=True, limit=None):
    """Таблица: кадр · фигура · первые 40 знаков · кегль · цвет · фон ·
    контраст · вердикт."""
    pick = [r for r in rows if r["ok"] is False] if only_bad \
        else list(rows)
    if limit:
        pick = pick[:limit]
    head = (f"{'кадр':>4} {'фигура':<22} {'текст (40 знаков)':<42} "
            f"{'кегль':>6} {'цвет':>8} {'фон':>8} {'контр.':>7} вердикт")
    out = [head, "─" * len(head)]
    for r in pick:
        sz = "—" if r["size"] is None else \
            f"{r['size']:.1f}{'ж' if r['bold'] else ''}"
        rt = "—" if r["ratio"] is None else f"{r['ratio']:.2f}"
        if r["ok"] is None:
            v = "не измерен (" + ("цвет не назван" if not r["fg"]
                                  else r["bg_src"]) + ")"
        elif r["ok"]:
            v = f"проходит (порог {r['thr']:.1f})"
        else:
            v = f"НИЖЕ ПОРОГА {r['thr']:.1f} · {_cls(r)}"
        out.append(f"{r['page']:>4} {r['shape'][:22]:<22} "
                   f"{r['text'][:42]:<42} {sz:>6} {r['fg'] or '—':>8} "
                   f"{r['bg'] or '—':>8} {rt:>7} {v}")
    return out


def pairs_of(rows):
    """Нарушения, свёрнутые в ПАРЫ «цвет на фоне, класс кегля».

    Правка у контраста идёт парой, а не прогоном: владелец меняет цвет
    ключевого числа один раз, и 60 прогонов лечатся одной правкой. Тот же
    договор, что у `floor.pt`, который называет КЛАСС текста, а не
    перечисляет четыре кегля одного кадра."""
    agg = {}
    for r in rows:
        if r["ok"] is not False:
            continue
        key = (r["fg"], r["bg"], _cls(r), r["thr"])
        a = agg.setdefault(key, {"runs": 0, "pages": set(),
                                 "ratio": r["ratio"], "sample": r["text"]})
        a["runs"] += 1
        a["pages"].add(r["page"])
    return agg


def summary_lines(rows, n_pages):
    bad = [r for r in rows if r["ok"] is False]
    unmeasured = [r for r in rows if r["ok"] is None]
    pages = sorted({r["page"] for r in bad})
    out = [f"прогонов с текстом: {len(rows)} на {n_pages} кадрах",
           f"НИЖЕ ПОРОГА: {len(bad)} {_run_word(len(bad))} на "
           f"{len(pages)} {_page_word(len(pages))}",
           f"не измерено (цвет/кегль/фон не названы): {len(unmeasured)}"]
    for (fg, bg, cls, thr), a in sorted(
            pairs_of(rows).items(), key=lambda kv: -kv[1]["runs"]):
        out.append(f"  {fg} на {bg}: {a['ratio']:.2f} : 1 при пороге "
                   f"{thr:.1f} ({cls}) — {a['runs']} {_run_word(a['runs'])} "
                   f"на {len(a['pages'])} {_page_word(len(a['pages']))}, "
                   f"например «{a['sample'][:30]}»")
    if pages:
        out.append("кадры: " + ", ".join(str(p) for p in pages))
    return out


def check(prs, *, tag="дека"):
    """Назвать нарушения контраста в отчёте сборки. Ничего не правит.

    Предупреждение на ПАРУ, а не на прогон: см. `pairs_of`. Тег
    `[контраст]` не блокирующий — по той же причине, что `[пол]`: починка
    здесь не одна (опустить цвет либо перевернуть плашку), и выбор между
    ними меняет кадр, то есть принадлежит владельцу."""
    rows = rows_of(prs)
    for (fg, bg, cls, thr), a in sorted(
            pairs_of(rows).items(), key=lambda kv: -kv[1]["runs"]):
        report.warn(
            f"[контраст] {tag}: {fg} на {bg} даёт {a['ratio']:.2f} : 1 при "
            f"пороге WCAG AA {thr:.1f} ({cls} текст) — {a['runs']} "
            f"{_run_word(a['runs'])} на {len(a['pages'])} "
            f"{_page_word(len(a['pages']))}, например «{a['sample'][:30]}». С "
            "задних рядов этого не видно. Починка не одна (опустить цвет "
            "текста либо перевернуть заливку), и выбор между ними меняет "
            "кадр — решение владельца, не сборки")
    return rows


def scan(path, *, only_bad=True):
    """Прогон по собранному файлу: справка палитры, таблица, итог."""
    prs = Presentation(str(path))
    rows = rows_of(prs)
    n = len(prs.slides._sldIdLst)
    lines = palette_lines() + [""]
    lines += [f"── прогоны деки {path} ──────────────────────────────"]
    lines += table_lines(rows, only_bad=only_bad)
    lines += [""] + summary_lines(rows, n)
    return rows, lines


def main(argv=None):
    argv = list(sys.argv[1:] if argv is None else argv)
    if "--selftest" in argv:
        from . import selftest_contrast
        return selftest_contrast.main()
    if "--palette" in argv:
        print("\n".join(palette_lines()))
        return 0
    only_bad = "--all" not in argv
    files = [a for a in argv if not a.startswith("-")]
    if not files:
        print(__doc__.strip().splitlines()[0])
        print("\nпрогон:  python3 -m deckkit.contrast <файл.pptx> [--all]"
              "\n         python3 -m deckkit.contrast --palette"
              "\n         python3 -m deckkit.contrast --selftest")
        return 2
    code = 0
    for f in files:
        rows, lines = scan(f, only_bad=only_bad)
        print("\n".join(lines))
        code = max(code, 1 if any(r["ok"] is False for r in rows) else 0)
    return code


if __name__ == "__main__":
    sys.exit(main())
