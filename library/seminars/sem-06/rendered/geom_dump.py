"""Снимок геометрии и текста собранной деки: по форме на строку.

Нужен затем, чтобы правку общей рисовалки (`deck_kit._col_shares`) доказывать
ЗАМЕРОМ по собранному файлу, а не доводом: что сдвинулось, на сколько, и не
изменился ли текст и кегль.
"""
import sys
from pptx import Presentation
from pptx.util import Emu

def walk(shapes, path, out):
    for i, sh in enumerate(shapes):
        p = f"{path}.{i}"
        try:
            geo = (Emu(sh.left).inches, Emu(sh.top).inches,
                   Emu(sh.width).inches, Emu(sh.height).inches)
        except Exception:
            geo = (0, 0, 0, 0)
        txt = ""
        if sh.has_text_frame:
            txt = " / ".join(r.text for pa in sh.text_frame.paragraphs for r in pa.runs)
            szs = ",".join(str(r.font.size and r.font.size.pt)
                           for pa in sh.text_frame.paragraphs for r in pa.runs)
        else:
            szs = ""
        out.append(f"{p}\t{geo[0]:.4f}\t{geo[1]:.4f}\t{geo[2]:.4f}\t{geo[3]:.4f}\t{szs}\t{txt}")
        if sh.shape_type == 6:
            walk(sh.shapes, p, out)

pr = Presentation(sys.argv[1])
out = []
for n, sl in enumerate(pr.slides, 1):
    walk(sl.shapes, f"s{n:02d}", out)
    nt = sl.notes_slide.notes_text_frame.text if sl.has_notes_slide else ""
    out.append(f"s{n:02d}.NOTES\t\t\t\t\t\t{len(nt.split())} слов")
print("\n".join(out))
