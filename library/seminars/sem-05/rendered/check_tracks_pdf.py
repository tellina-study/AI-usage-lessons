#!/usr/bin/env python3
"""Сверка СОБРАННОГО pptx с тем, что на самом деле нарисовал LibreOffice.

Зачем. Вёрстка ставит фигуру под текстовый блок на высоту, посчитанную
`metrics`. `metrics.line_h` считает межстрочный как долю КЕГЛЯ (18 pt × 1.30 =
23.4 pt), а `python-pptx` пишет в файл `line_spacing` вещественным числом, и
LibreOffice с PowerPoint понимают его как долю СОБСТВЕННОЙ высоты строки
шрифта. Замерено на этой деке: 28.0 pt против 23.4 — строка рисуется на 19.7%
выше мерки. Ошибка копится по строкам, и первой её ловит та фигура, которую
ставят вплотную под текст: ряд пилюль-терминов под базой в приёме
`base_and_edge`.

Почему не хватает того, что уже есть. `qa_preview.py` рисует картинку САМ, по
той же мерке, которой верстали, — расхождения с LibreOffice он не видит по
построению. Сборочные сторожа (`ПЕРЕПОЛНЕНИЕ`, `НА ДНЕ КЕГЛЯ`, `ШИРЕ РАМКИ`)
меряют тем же `metrics`. На `n46` все они молчали, а пилюли резали последнюю
строку базы на 4,8 pt.

Эта проверка не меряет текст — она СМОТРИТ в PDF, который выдал LibreOffice, и
сравнивает нарисованные строки с рамками фигур из pptx.

    python3 check_tracks_pdf.py sem-05.pptx /tmp/out/sem-05.pdf 47 48 49
    python3 check_tracks_pdf.py --self-test      # прогон на сломанном входе

`РЕЖЕТ` — фигура пересекает строку текста ЧАСТИЧНО. Плашка, внутри которой
текст лежит целиком, — штатная форма и не считается.
"""
import subprocess
import sys
import tempfile
from pathlib import Path

import pymupdf
from pptx import Presentation
from pptx.enum.dml import MSO_FILL

HERE = Path(__file__).resolve().parent
EMU_PT = 12700.0
CUT_TOL = 0.75          # pt — ниже этого пересечение считаем кромкой отрисовки
OWN_TOL = 4.0           # pt — на столько строка вправе вылезти за СВОЮ рамку
                        # и остаться своей. Из-за той же недомерки строка
                        # вылезает за собственную рамку: вправо до 1.3 pt
                        # (n46, правая дорожка) и вверх до 2.1 pt (n09, плашка).
                        # С допуском 0.75 владелец не опознавался, и каждая
                        # такая строка приходила ложным «РЕЖЕТ».


def shapes_of(pptx_path, want=None):
    prs = Presentation(pptx_path)
    out = {}
    for i, sl in enumerate(prs.slides):
        if want is not None and i not in want:
            continue
        rows = []
        for sh in sl.shapes:
            try:
                l, t = sh.left / EMU_PT, sh.top / EMU_PT
                w, h = sh.width / EMU_PT, sh.height / EMU_PT
            except TypeError:
                continue
            try:
                solid = sh.fill.type == MSO_FILL.SOLID
            except (AttributeError, TypeError, NotImplementedError):
                solid = False
            rows.append((l, t, l + w, t + h, solid))
        out[i] = rows
    return out


def lines_of(page):
    out = []
    for b in page.get_text("dict")["blocks"]:
        for ln in b.get("lines", []):
            txt = "".join(s["text"] for s in ln["spans"]).strip()
            if txt:
                out.append((ln["bbox"], txt, ln["spans"][0]["size"]))
    return out


def cuts(shape, line_bbox):
    """Насколько фигура ПЕРЕСЕКАЕТ строку, не будучи её собственной рамкой.

    Признак владения — фигура НАЧИНАЕТСЯ не ниже строки и накрывает её по
    ширине. Нижнюю границу в признак не берём намеренно: из-за недомерки
    последний ряд вылезает за дно своей рамки почти целиком (замерено: рамка
    кончается на 1.3 pt ниже верха последней строки), и по условию «рамка
    содержит строку» владелец переставал опознаваться. Фигуру, которую ставят
    ПОД текст, от собственной рамки отличает именно верх: пилюли на пробе
    начинаются на 11 pt ниже верха строки, которую режут.
    """
    sl, st, sr, sb = shape[:4]
    lx0, ly0, lx1, ly1 = line_bbox
    ov_x = min(sr, lx1) - max(sl, lx0)
    ov_y = min(sb, ly1) - max(st, ly0)
    if ov_x <= 0 or ov_y <= CUT_TOL:
        return 0.0
    owns = (sl <= lx0 + OWN_TOL and sr >= lx1 - OWN_TOL
            and st <= ly0 + OWN_TOL)
    return 0.0 if owns else ov_y


def spills(shape, line_bbox):
    """Строка ВЫВАЛИЛАСЬ за дно ЗАЛИТОЙ карточки, которая её содержит.

    Отдельное сообщение, а не разновидность `РЕЖЕТ`: здесь никто ни на кого не
    налезает — карточка просто кончается выше своего текста, и последняя строка
    печатается тёмным по белому за её краем. Для незалитой рамки это не дефект
    (её не видно), поэтому проверка ограничена заливкой.
    """
    sl, st, sr, sb, solid = shape
    lx0, ly0, lx1, ly1 = line_bbox
    if not solid:
        return 0.0
    # Содержит — значит строка НАЧИНАЕТСЯ внутри фигуры. Без условия `sb`
    # контейнером оказывалась и планка в 3.6 pt над дорожкой: она шире ярлыка
    # и стоит выше него, и ярлык «вываливался» из неё на всю свою высоту.
    starts_inside = st <= ly0 + OWN_TOL and sb >= ly0 + CUT_TOL
    owns = sl <= lx0 + OWN_TOL and sr >= lx1 - OWN_TOL and starts_inside
    return (ly1 - sb) if owns and ly1 - sb > CUT_TOL else 0.0


def check(pptx_path, pdf_path, pages=None):
    doc = pymupdf.open(pdf_path)
    want = set(pages) if pages is not None else None
    msgs = []
    for i, shapes in sorted(shapes_of(pptx_path, want).items()):
        if i >= doc.page_count:
            continue
        for bbox, txt, _size in lines_of(doc[i]):
            for s in shapes:
                ov = cuts(s, bbox)
                if ov:
                    msgs.append(f"РЕЖЕТ [стр. {i + 1}]: фигура (верх {s[1]:.1f} pt) "
                                f"заходит на строку «{txt[:34]}» на {ov:.1f} pt")
                out = spills(s, bbox)
                if out:
                    msgs.append(f"ВЫВАЛИЛСЯ [стр. {i + 1}]: строка «{txt[:34]}» "
                                f"на {out:.1f} pt ниже дна залитой карточки")
    doc.close()
    return sorted(set(msgs))


def measured_pitch(pdf_path, page_index, size_pt):
    """Реальный шаг строки для кегля `size_pt` — то, чем калибруют `metrics`."""
    doc = pymupdf.open(pdf_path)
    tops = sorted(b[0][1] for b in lines_of(doc[page_index])
                  if abs(b[2] - size_pt) < 0.1)
    doc.close()
    steps = [b - a for a, b in zip(tops, tops[1:]) if 0 < b - a < size_pt * 3]
    return min(steps) if steps else None


# ── Самопроверка на нарочно сломанном входе ─────────────────────────────────
#
# ЧТО ЗДЕСЬ ПРОВЕРЯЕТСЯ И ПОЧЕМУ ИМЕННО ТАК.
#
# Прежние пробы собирались приёмами вёрстки (`base_and_edge`, `terminal_card`)
# и ломались НЕДОМЕРКОЙ МЕЖСТРОЧНОГО: мерка считала шаг долей кегля, LibreOffice
# рисовал долей высоты строки шрифта, разница 19,7% копилась по строкам — и
# пилюли резали базу, а листинг уезжал за дно карточки. Недомерку починили
# (`deck_kit.text_box` пишет `Pt(size * spacing)`), и обе пробы перестали
# ломаться: карточка теперь честно вырастает под свой листинг, пилюли честно
# встают под текст. Проба, которая больше не ломается, НИЧЕГО НЕ ДОКАЗЫВАЕТ про
# проверку — она доказывает только, что починка работает.
#
# Поэтому пробы разделены по тому, что каждая сторожит:
#
#   * `cuts`/`spills` — предикаты геометрии: «фигура налезла на строку»,
#     «залитая карточка кончилась выше своего текста». Их и надо кормить
#     ГЕОМЕТРИЕЙ, выставленной руками, а не ждать, пока её случайно соберёт
#     приём вёрстки. Такая проба не зависит от арифметики ни одного приёма и не
#     умирает от её починки — ровно то свойство, которого не хватало прежним.
#   * Сама недомерка — отдельной пробой на ШАГ СТРОКИ: две одинаковые надписи,
#     нарисованный шаг против мерки. Эта проба поймала бы исходный дефект и
#     поймает возврат `p.line_spacing = spacing` вещественным числом. Прежде
#     расхождение лишь печаталось числом рядом с результатом и ни на что не влияло:
#     19,7% на экране не делали прогон красным.

PITCH_TOL = 0.02        # доля — на столько нарисованный шаг вправе разойтись
                        # с меркой. 2% — это округление кегля в PDF, а не
                        # недомерка: исходный дефект давал 19,7%.


def _probe_cut(out_pptx, *, overlap):
    """Четыре строки текста и фигура под ними. `overlap=True` — фигура
    поставлена НА последнюю строку (так вставал ряд пилюль при недомерке);
    `False` — под ней с зазором. Координаты заданы числами: проба сторожит
    предикат `cuts`, а не чью-то арифметику."""
    from pptx import Presentation as P
    from pptx.util import Inches
    import deck_kit as K
    prs = P()
    prs.slide_width, prs.slide_height = Inches(K.W_IN), Inches(K.H_IN)
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    K.set_bg(sl, K.WHITE)
    lines = [f"База, строка {i}." for i in (1, 2, 3, 4)]
    size, spacing, w = 18, K.TRACK_SPACING, 5.0
    # Ни один абзац пробы не вправе ПЕРЕНЕСТИСЬ: расчёт `last_top` ниже считает
    # абзац одной строкой, и перенос молча сдвинул бы последнюю строку вниз —
    # годная проба тогда получает фигуру на настоящем тексте и приходит ложным
    # «РЕЖЕТ». Так и вышло на первом прогоне: строки были длинные, каждая
    # разошлась на две, и зазор 0,15″ оказался отмерен не от той строки.
    assert all(K.M.nlines(ln, size, w) == 1 for ln in lines), \
        "абзац пробы переносится — зазор будет отмерен не от последней строки"
    y, step = 1.00, K.M.line_h(size, spacing)
    K.text_box(sl, 1.0, y, w, 4 * step, lines, size=size, spacing=spacing)
    last_top = y + 3 * step
    # сломанный вход: верх фигуры ВНУТРИ последней строки, на треть её высоты
    # ниже верха — так, чтобы фигура не опознавалась как собственная рамка
    # строки (признак владения в `cuts` — верх не ниже строки).
    box_y = last_top + step / 3 if overlap else last_top + step + 0.22
    K.ocean_box(sl, 1.0, box_y, w, 0.42)
    prs.save(out_pptx)


def _probe_spill(out_pptx, *, short_card):
    """Залитая карточка и текст в ней. `short_card=True` — дно карточки
    приходится на середину одной из строк, и всё, что ниже, печатается тёмным
    по белому за её краем; `False` — карточка накрывает текст целиком.

    Дно обязано попасть ВНУТРЬ строки, а не выше всех: `spills` считает
    вываливанием только строку, которая НАЧАЛАСЬ внутри карточки (иначе
    контейнером оказывается любая фигура выше текста). Поэтому строк восемь —
    какая-нибудь из них заведомо окажется на кромке при любом округлении."""
    from pptx import Presentation as P
    from pptx.util import Inches
    import deck_kit as K
    prs = P()
    prs.slide_width, prs.slide_height = Inches(K.W_IN), Inches(K.H_IN)
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    K.set_bg(sl, K.WHITE)
    lines = [f"строка вывода номер {i} в залитой карточке" for i in range(1, 9)]
    size, spacing, pad = 14, 1.30, 0.20
    step = K.M.line_h(size, spacing)
    x, y, w = 1.0, 1.0, 6.0
    full = len(lines) * step + 2 * pad
    K.ocean_box(sl, x, y, w, (5.4 * step + 2 * pad) if short_card else full,
                fill=K.CODE_BG, stroke=K.LIGHT)
    K.text_box(sl, x + pad, y + pad, w - 2 * pad, len(lines) * step, lines,
               size=size, spacing=spacing, color=K.CODE_FG, mono=True)
    prs.save(out_pptx)


def _probe_pitch(out_pptx):
    """Две одинаковые надписи кеглем 18 и межстрочным `TRACK_SPACING` — по
    расстоянию между ними меряется НАРИСОВАННЫЙ шаг строки."""
    from pptx import Presentation as P
    from pptx.util import Inches
    import deck_kit as K
    prs = P()
    prs.slide_width, prs.slide_height = Inches(K.W_IN), Inches(K.H_IN)
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    K.set_bg(sl, K.WHITE)
    K.text_box(sl, 1.0, 1.0, 9.0, 3.0,
               ["Мерка шага строки — надпись один.",
                "Мерка шага строки — надпись два.",
                "Мерка шага строки — надпись три."],
               size=18, spacing=K.TRACK_SPACING)
    prs.save(out_pptx)


def self_test():
    sys.path.insert(0, str(HERE))
    import deck_kit as K
    import metrics as M
    script = HERE.parents[3] / "tools" / "presentation-build" / "pptx_to_png.sh"
    ok = True

    def render(build, tag, d):
        pptx = d / f"probe-{tag}.pptx"
        build(str(pptx))
        subprocess.run([str(script), str(pptx), str(d), "150", "1", "1"],
                       check=True, capture_output=True)
        return pptx, pptx.with_suffix(".pdf")

    with tempfile.TemporaryDirectory() as d:
        d = Path(d)
        cases = (
            ("фигура на строке: верх внутри строки", "РЕЖЕТ", True,
             lambda o: _probe_cut(o, overlap=True)),
            ("фигура под строкой: зазор 0,22″", "РЕЖЕТ", False,
             lambda o: _probe_cut(o, overlap=False)),
            ("дно карточки в середине строки", "ВЫВАЛИЛСЯ", True,
             lambda o: _probe_spill(o, short_card=True)),
            ("карточка накрывает текст целиком", "ВЫВАЛИЛСЯ", False,
             lambda o: _probe_spill(o, short_card=False)),
        )
        for i, (name, kind, want_hit, build) in enumerate(cases):
            pptx, pdf = render(build, f"{kind}-{i}", d)
            msgs = [m for m in check(str(pptx), str(pdf), pages=[0])
                    if m.startswith(kind)]
            hit = bool(msgs)
            ok = ok and hit == want_hit
            mark = "✓" if hit == want_hit else "✗ ПРОВЕРКА МЁРТВАЯ"
            print(f"{mark} {name:38} → {msgs[0][:78] if msgs else 'чисто'}")

        # ── шаг строки: проба, которая поймала бы саму недомерку ────────────
        pptx, pdf = render(_probe_pitch, "pitch", d)
        drawn = measured_pitch(str(pdf), 0, 18.0)
        est = M.line_h(18, K.TRACK_SPACING) * 72
        if drawn is None:
            ok = False
            print("✗ ПРОВЕРКА МЁРТВАЯ шаг строки: в PDF не нашлось двух строк кеглем 18")
        else:
            off = drawn / est - 1
            good = abs(off) <= PITCH_TOL
            ok = ok and good
            mark = "✓" if good else f"✗ НЕДОМЕРКА ВЕРНУЛАСЬ (допуск {PITCH_TOL:.0%})"
            print(f"{mark} {'шаг строки совпадает с меркой':38} → мерка {est:.1f} pt, "
                  f"нарисовано {drawn:.1f} pt, расхождение {off * 100:+.1f}%")
    return 0 if ok else 1


def main():
    args = sys.argv[1:]
    if not args or args[0] == "--self-test":
        return self_test()
    pptx, pdf = args[0], args[1]
    pages = [int(x) - 1 for x in args[2:]] or None
    msgs = check(pptx, pdf, pages)
    for m in msgs:
        print("⚠", m)
    print(f"— фигур, режущих текст: {len(msgs)}")
    return 1 if msgs else 0


if __name__ == "__main__":
    sys.exit(main())
