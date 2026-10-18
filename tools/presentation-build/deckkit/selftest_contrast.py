# -*- coding: utf-8 -*-
"""РАЗНОСТНЫЕ ПРОБЫ: ПОЛ КОНТРАСТА и HERO ПО ПОРЯДКУ ДЕКИ.

Правило пакета: ПРОВЕРКА, НЕ ПРОВЕРЕННАЯ НА СЛОМАННОМ ВХОДЕ, — НЕ ПРОВЕРКА.
Для обеих проверок этого файла оно звучит жёстче обычного, потому что обе
заведены ПОСЛЕ дефекта, который прошёл мимо всех прежних сторожей: сторож,
сданный молчанием на здоровом материале, здесь не отличался бы от того,
которого не было вовсе — а именно его и не было.

  [контраст]            три посаженных нарушения по одному на класс:
                        крупный ниже 3,0 · обычный ниже 4,5 · текст на
                        цветной заливке. Плюс палитра курса: пять пар, мой
                        счёт против независимого замера владельца
  [контраст-приманка]   ровно на пороге — МОЛЧАНИЕ; прогон без названного
                        цвета — «не измерен», а не «прошёл»
  [контраст-кегль]      порогов ДВА: один и тот же цвет ловится на 15 пт и
                        молчит на 20 пт, и молчит на 14 пт полужирным
  [контраст-фон]        фон читается по z-порядку СОБРАННОГО файла: тот же
                        прогон на белом проходит 13,67, на фактической
                        заливке — 1,82
  [hero-порядок]        последний кадр ПОКАЗА без hero назван — и в двух
                        видах сразу: объявленный hero на 10% площади и
                        кадр вовсе без снимка
  [hero-объявление]     разностно: прежний сторож (по полю `media_hero`)
                        на кадре-ленте МОЛЧИТ, новый (по порядку деки)
                        говорит — то есть проба про порядок, а не про порог
  [hero-здоровая]       деке, у которой оба конца несут крупный снимок, —
                        молчание

Запуск:  python3 -m deckkit.contrast --selftest
"""
import sys
import tempfile
from pathlib import Path

from PIL import Image
from pptx.dml.color import RGBColor

from . import contrast, media, prims, probes_contrast as PC, refdeck, report
from .theme import SLIDE_H, SLIDE_W

OK, BAD, SAID = [], [], []


def say(line):
    SAID.append(line)
    print(line)


# ════════════════════════════════════════════════════════════════════
# контраст
# ════════════════════════════════════════════════════════════════════


def _deck_with(cases):
    """Собрать в памяти кадр с посаженными парами «цвет на заливке».

    Фигурами, а не файлом кадра, нарочно: фон у надписи даёт фигура ПОД
    ней, и проба обязана проверить именно это чтение — то есть поставить
    заливку и текст отдельными фигурами, как это делает любая раскладка."""
    p = prims.new_deck()
    s = prims.blank(p)
    prims.rect(s, 0, 0, SLIDE_W, SLIDE_H, RGBColor(0xFF, 0xFF, 0xFF),
               count=None)
    y = 0.3
    for name, fg, bg, size, bold, _must in cases:
        if bg.upper() != "FFFFFF":
            prims.rect(s, 0.4, y, 6.0, 0.6, RGBColor(*contrast.hex_rgb(bg)),
                       count=None)
        prims.text_box(s, 0.5, y + 0.05, 5.8, 0.5, name, size=size,
                       bold=bold, color=RGBColor(*contrast.hex_rgb(fg)),
                       quiet=True)
        y += 0.8
    return p


def probe_contrast_classes():
    """Три посаженных нарушения — по одному на класс порога."""
    report.reset()
    p = _deck_with(PC.CASES)
    rows = contrast.rows_of(p)
    found = {r["text"]: r for r in rows}
    missed, false_hit = [], []
    for name, fg, bg, size, bold, must in PC.CASES:
        r = found.get(name)
        if r is None or r["ok"] is None:
            missed.append(f"{name}: прогон не измерен вовсе")
            continue
        if must and r["ok"]:
            missed.append(f"{name}: {r['ratio']:.4f} при пороге "
                          f"{r['thr']:.1f} — НЕ названо")
        if not must and not r["ok"]:
            false_hit.append(f"{name}: {r['ratio']:.6f} при пороге "
                             f"{r['thr']:.1f} — названо ЛОЖНО")
    report.reset()
    if missed or false_hit:
        BAD.append("[контраст] " + "; ".join(missed + false_hit))
        return
    OK.append("[контраст]")
    for name, *_ in PC.CASES:
        r = found[name]
        say(f"{'[контраст]':<24} {name:<20} {r['fg']} на {r['bg']} → "
            f"{r['ratio']:.4f} при пороге {r['thr']:.1f} "
            + ("НАЗВАНО" if not r["ok"] else "молчит"))


def probe_threshold_edge():
    """Ровно на пороге — проходит. И «ровно» здесь ИЗМЕРЕНО.

    Две половины. Первая — на уровне самой функции: `passes(3.0, 3.0)`
    обязано быть истиной, иначе порог зависел бы от последнего бита двойной
    точности. Вторая — на настоящей паре цветов: из всех 65 536 серых пар
    ближайшая к ровно 3,0 даёт 3,000004, и проба печатает это число, а не
    слово «ровно», потому что 8 бит на канал точнее не умеют."""
    exact = (contrast.passes(contrast.LARGE_RATIO, contrast.LARGE_RATIO)
             and contrast.passes(contrast.NORMAL_RATIO,
                                 contrast.NORMAL_RATIO))
    under = not contrast.passes(contrast.LARGE_RATIO - 1e-4,
                                contrast.LARGE_RATIO)
    r = contrast.ratio(PC.EDGE_FG, PC.EDGE_BG)
    near = abs(r - contrast.LARGE_RATIO) < 1e-4 and r >= contrast.LARGE_RATIO
    if not (exact and under and near):
        BAD.append(f"[контраст-приманка] ровно на пороге проходит {exact}, "
                   f"чуть ниже отвергается {under}, пара приманки "
                   f"{r:.6f} годна {near}")
        return
    OK.append("[контраст-приманка]")
    say(f"{'[контраст-приманка]':<24} ровно на пороге: МОЛЧИТ  →  пара "
        f"{PC.EDGE_FG} на {PC.EDGE_BG} даёт {r:.6f} при пороге "
        f"{contrast.LARGE_RATIO:.1f} — проходит; та же проверка при "
        f"{contrast.LARGE_RATIO - 1e-4:.4f} уже называет")


def probe_no_color():
    """Прогон без НАЗВАННОГО цвета — «не измерен», а не «прошёл».

    Сломанный вход — само молчание: прогон, которому цвет не задан,
    наследует его от темы, темы у дек этого курса нет, и «прошёл» здесь
    значил бы «мы не посмотрели». Проба требует третьего разряда."""
    from pptx.util import Pt
    report.reset()
    p = prims.new_deck()
    s = prims.blank(p)
    tb = s.shapes.add_textbox(0, 0, 3000000, 500000)
    r = tb.text_frame.paragraphs[0].add_run()
    r.text = "прогон без названного цвета"
    r.font.size = Pt(15)
    rows = contrast.rows_of(p)
    report.reset()
    mine = [x for x in rows if "без названного" in x["text"]]
    if len(mine) != 1 or mine[0]["ok"] is not None or mine[0]["fg"]:
        BAD.append(f"[контраст-приманка цвет] прогонов {len(mine)}, "
                   f"вердикт {mine[0]['ok'] if mine else '—'} — разряда "
                   "«не измерен» нет")
        return
    OK.append("[контраст-приманка цвет]")
    say(f"{'[контраст-приманка цвет]':<24} без явного цвета: НЕ ИЗМЕРЕН  →  "
        f"вердикт {mine[0]['ok']}, цвет {mine[0]['fg']} — третий разряд "
        "есть, «прошёл» не выдаётся за «посмотрели»")


def probe_size_split():
    """Порогов ДВА: тот же цвет ловится обычным кеглем и молчит крупным."""
    report.reset()
    cases = [(f"кегль {sz}{'ж' if b else ''}", PC.GREY_BELOW_NORMAL,
              "FFFFFF", sz, b, must) for sz, b, must in PC.SIZE_SPLIT]
    p = _deck_with(cases)
    rows = {r["text"]: r for r in contrast.rows_of(p)}
    report.reset()
    wrong = []
    for name, _fg, _bg, sz, b, must in cases:
        r = rows.get(name)
        if r is None or r["ok"] is None or r["ok"] == must:
            wrong.append(f"{name}: вердикт {r and r['ok']}, ждали "
                         f"{'НАЗВАНО' if must else 'молчание'}")
    if wrong:
        BAD.append("[контраст-кегль] " + "; ".join(wrong))
        return
    OK.append("[контраст-кегль]")
    one = rows[cases[0][0]]
    say(f"{'[контраст-кегль]':<24} разностно: ЛОВИТ  →  один цвет "
        f"{one['fg']} на белом даёт {one['ratio']:.4f}: на 15 пт это ниже "
        f"{contrast.NORMAL_RATIO:.1f} и НАЗВАНО, на 20 пт и на 14 пт "
        f"полужирным — выше {contrast.LARGE_RATIO:.1f} и молчание")


def probe_background_source():
    """Фон читается по z-порядку СОБРАННОГО файла, а не берётся белым.

    Разностно и в одном цвете: глубокий `21295C` на белом даёт 13,67 и
    проходит с запасом, на фактической заливке `065A82` — 1,82 и обязан
    быть назван. Без этой половины проба доказывала бы только, что 1,82
    меньше 4,5; с ней она доказывает, что фон вообще измеряется."""
    report.reset()
    pair = [("на белом", "21295C", "FFFFFF", 16.0, False, False),
            ("на заливке", "21295C", "065A82", 16.0, False, True)]
    rows = {r["text"]: r for r in contrast.rows_of(_deck_with(pair))}
    report.reset()
    white, fill = rows.get("на белом"), rows.get("на заливке")
    if not white or not fill or white["bg"] != "FFFFFF" \
            or fill["bg"] != "065A82" or not white["ok"] or fill["ok"]:
        BAD.append(f"[контраст-фон] на белом {white and white['bg']}/"
                   f"{white and white['ok']}, на заливке "
                   f"{fill and fill['bg']}/{fill and fill['ok']} — фон не "
                   "читается")
        return
    OK.append("[контраст-фон]")
    say(f"{'[контраст-фон]':<24} разностно: ЛОВИТ  →  один цвет 21295C: "
        f"фон {white['bg']} ({white['bg_src']}) → {white['ratio']:.2f} "
        f"проходит; фон {fill['bg']} ({fill['bg_src']}) → "
        f"{fill['ratio']:.2f} НАЗВАНО")


def probe_palette():
    """Пять пар палитры курса: мой счёт против независимого замера.

    Расхождение больше допуска — находка, а не повод подогнать формулу."""
    worst, lines = 0.0, []
    for name, fg, bg, want in contrast.PALETTE_PAIRS:
        r = contrast.ratio(fg, bg)
        worst = max(worst, abs(r - want))
        lines.append(f"{name} {r:.4f} против {want:.2f}")
    if worst > contrast.PALETTE_TOL:
        BAD.append(f"[палитра] наибольшее расхождение {worst:.4f} при "
                   f"допуске {contrast.PALETTE_TOL:.2f}: "
                   + "; ".join(lines))
        return
    OK.append("[палитра]")
    say(f"{'[палитра]':<24} сверено с независимым замером: СХОДИТСЯ  →  "
        f"наибольшее расхождение {worst:.4f} при допуске "
        f"{contrast.PALETTE_TOL:.2f} на пяти парах")


# ════════════════════════════════════════════════════════════════════
# hero по порядку деки
# ════════════════════════════════════════════════════════════════════


def _hero_assets(d):
    """Снимки под пробы hero. Рисуются здесь: проба не вправе зависеть от
    ассетов лекции. Отношения подобраны под порог §5.9 и под его провал —
    1,4 даёт 41% площади, 5,77 даёт 10% (ровно случай `s106` Лекции 7)."""
    Image.new("RGB", (1400, 1000), (200, 210, 220)).save(d / "probe-hero-big.png")
    Image.new("RGB", (1500, 260), (210, 215, 225)).save(d / "probe-hero-thin.png")


def _build(files, tmp, seq, *, partial=False):
    """Собрать миниатюрную деку и прогнать по ней ОБА сторожа hero."""
    report.reset()
    seq[0] += 1
    d = tmp / f"h{seq[0]}"
    d.mkdir(parents=True)
    _hero_assets(d)
    for nm, txt in files:
        (d / nm).write_text(txt, encoding="utf-8")
    out = tmp / f"hero{seq[0]}.pptx"
    refdeck.build(d, out, root=d)
    declared = list(report.WARNINGS)        # сказал прежний сторож
    from pptx import Presentation
    shares = media.check_hero_order(Presentation(str(out)), partial=partial)
    order = [w for w in report.WARNINGS if w not in declared]
    return declared, order, shares


def probe_hero_order(tmp, seq):
    """Последний кадр показа без hero — назван, и в двух видах."""
    res = {}
    for name, files in (("объявлен, 10%", PC.DECK_SMALL_HERO),
                        ("снимка нет", PC.DECK_NO_HERO),
                        ("здоровая", PC.DECK_OK)):
        res[name] = _build(files, tmp, seq)
    report.reset()
    bad = []
    for name in ("объявлен, 10%", "снимка нет"):
        if not [w for w in res[name][1] if "порядка деки" in w]:
            bad.append(f"{name}: порядок деки молчит, доли {res[name][2]}")
    if res["здоровая"][1]:
        bad.append(f"здоровая дека: {len(res['здоровая'][1])} предупреждений "
                   f"— шумит, доли {res['здоровая'][2]}")
    if bad:
        BAD.append("[hero-порядок] " + "; ".join(bad))
        return
    OK.append("[hero-порядок]")
    for name in ("объявлен, 10%", "снимка нет", "здоровая"):
        sh = res[name][2]
        say(f"{'[hero-порядок]':<24} {name:<16} "
            + " · ".join(f"стр {k} — {100 * v:.1f}%"
                         for k, v in sorted(sh.items()))
            + ("  →  НАЗВАНО" if res[name][1] else "  →  молчание"))


def probe_hero_declaration(tmp, seq):
    """Разностно: прежний сторож на этом входе МОЛЧИТ, новый говорит.

    Сломанный вход — сам прежний сторож. `media.place_hero` меряет кадры,
    где автор поставил `media_hero: true`; на кадре-ленте (`media` без
    `media_hero` — ровно `s106` Лекции 7) ему мерить, по его мнению, нечего,
    и он молчит. Без этой половины проба доказывала бы только, что 10%
    меньше 40%, и новый сторож был бы неотличим от старого с другим
    порогом."""
    ribbon = (("s01-first.md", PC.DECK_OK[0][1]),
              ("s02-mid.md", PC.DECK_OK[1][1]),
              ("s03-last.md", PC.frame("s03", img="probe-hero-thin.png")))
    declared, order, shares = _build(ribbon, tmp, seq)
    report.reset()
    old = [w for w in declared if "hero" in w]
    new = [w for w in order if "порядка деки" in w]
    last = shares.get(max(shares)) if shares else None
    if old or not new or last is None or last >= media.HERO_MIN_AREA:
        BAD.append(f"[hero-объявление] прежний сторож сказал {len(old)}, "
                   f"новый {len(new)}, доля последнего {last} — проба не "
                   "про порядок")
        return
    OK.append("[hero-объявление]")
    say(f"{'[hero-объявление]':<24} разностно: ЛОВИТ  →  на кадре-ленте "
        f"(`media` без `media_hero`) прежний сторож по объявлению автора "
        f"сказал {len(old)} слов, новый по порядку деки — последний кадр "
        f"показа занят снимком на {100 * last:.1f}% при пороге "
        f"{100 * media.HERO_MIN_AREA:.0f}%")


def probe_hero_partial(tmp, seq):
    """Сборка ЧАСТИ деки не молчит, а говорит, что не мерила.

    Сломанный вход — `--only`: «первый и последний» там означает «из
    отобранных», то есть не то, что видит зал. Молчание в этом случае
    неотличимо от «порог взят», и ровно так сторож и соврал бы тише."""
    _d, order, shares = _build(PC.DECK_NO_HERO, tmp, seq, partial=True)
    report.reset()
    said = [w for w in order if "НЕ МЕРЕН" in w]
    if not said or shares:
        BAD.append(f"[hero-часть] сказано {len(said)}, доли {shares} — "
                   "частичная сборка молчит либо мерит")
        return
    OK.append("[hero-часть]")
    say(f"{'[hero-часть]':<24} на частичной сборке: ГОВОРИТ  →  "
        f"{said[0][:96]}")


# ════════════════════════════════════════════════════════════════════


def run_all(ok, bad, say_, run=None, tmp=None, seq=None):
    """Зовётся из `deckkit.selftest`, списками того прогона."""
    global OK, BAD, say
    OK, BAD, say = ok, bad, say_
    tmp = tmp or Path(tempfile.mkdtemp(prefix="deckkit-contrast-"))
    seq = seq or [0]
    probe_palette()
    probe_contrast_classes()
    probe_threshold_edge()
    probe_no_color()
    probe_size_split()
    probe_background_source()
    probe_hero_order(tmp, seq)
    probe_hero_declaration(tmp, seq)
    probe_hero_partial(tmp, seq)
    report.reset()


def main():
    """`python3 -m deckkit.contrast --selftest`. Код 1 — есть недоказанные
    пробы: прогон, не поймавший посаженного, ОБЯЗАН валить сборку."""
    OK[:] = []
    BAD[:] = []
    say("── САМОПРОВЕРКА ПОЛА КОНТРАСТА И HERO ПО ПОРЯДКУ ДЕКИ ───────")
    run_all(OK, BAD, say)
    say("─────────────────────────────────────────────────────────────")
    total = len(OK) + len(BAD)
    say(f"проверок доказано: {len(OK)} из {total}")
    if not total:
        say("НИ ОДНА проба не выполнилась — это НЕ успех")
        return 2
    if BAD:
        say("НЕ ДОКАЗАНО:")
        for x in BAD:
            say("  " + x)
        return 1
    say("все проверки доказаны на сломанном входе")
    return 0


if __name__ == "__main__":
    sys.exit(main())
