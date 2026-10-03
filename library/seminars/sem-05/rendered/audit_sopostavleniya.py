#!/usr/bin/env python3
"""Формат сопоставления «не X, а Y» в готовой деке — главный курсовой запрет.

`tools/editorial/README.md` §1: «не X, а Y» · «дело не в X, дело в Y» · «с одной
стороны… с другой». Владелец, круг 4, правило А8: «вот эти вот сравнения
„или-или“, „да, но нет“ и так далее выправить».

ЧИТАЕТ ГОТОВЫЙ `.pptx`, А НЕ ИСХОДНИКИ — по той же причине, что и
`audit_sluzhebnoe.py`: текст слайда приходит из `## Visual`, из генераторов схем
и из самого сборщика, и grep по `.md` видит только первое. Вдобавок здесь это
единственный способ отделить ЭКРАН от РЕЧИ: в исходнике заметка и видимый слой
лежат в одном файле, а нормы у них разные (§1 курсового слоя плюс
`AUTHOR-BRIEF` п. 10 про регистр).

ПРОИЗВОДСТВЕННАЯ РАЗМЕТКА НЕ СЧИТАЕТСЯ: `visual.backup` во фронтматтере — это
протокол решений круга, он залу не показывается и в хронометраж не входит. Он и
не попадает сюда по построению, раз мы читаем собранный файл, — но сказать это
нужно, потому что grep по `.md` его считал и завышал счёт.

## Граница слова — не придирка, а причина ложных находок

Шаблон из §1 канона был записан как `не [^,.;:]{2,45}, а ` — БЕЗ `\b`. Такой
шаблон ловит хвост любого слова на «-не»: «на сторо**не** платформы, а …»,
«написан в пла**не** ступени, а …». На этой деке разница замерена: 107 находок
без границы против 88 с границей, то есть 19 из 107 — мусор, и «не X, а Y» на
экране завышалось с 7 до 11. Сам канон поправлен тем же кругом; здесь граница
стоит с самого начала.

Отдельно считается ХВОСТОВАЯ форма «…, а не …» — зеркало той же фигуры и на
этой деке её основная масса (58 из 64 в речи). Прожарка круга 3 считала обе и
сводила в одно число (94); здесь они разведены, потому что правятся по-разному:
головная переписывается утверждением, хвостовая чаще снимается целиком.

    python3 audit_sopostavleniya.py [файл.pptx ...]    # по умолчанию sem-05.pptx
    python3 audit_sopostavleniya.py --spisok           # с цитатами по слайдам
    python3 audit_sopostavleniya.py --self-test

Код возврата 1, если найдено хоть одно, — число печатается всегда, потому что
смысл этой проверки в ВЕЛИЧИНЕ, а не в наличии: ноль здесь недостижим и не нужен
(противопоставление, которое и есть содержание, §1 разрешает развернуть).
"""
import re
import sys
from pathlib import Path

from pptx import Presentation

HERE = Path(__file__).parent

# `\bне\b` обязательно. Хвостовая форма в границе слова не нуждается слева, но
# нуждается справа: «, а неделя» не противопоставление.
FIGURY = [
    ("хвостовое «, а не»", re.compile(r",\s+а\s+не\b")),
    ("головное «не X, а Y»", re.compile(r"\bне\b[^,.;:!?]{2,45},\s+а\s+")),
    ("«дело не в X»", re.compile(r"\bдело\s+не\s+в\b", re.I)),
    ("«с одной стороны»", re.compile(r"\bс\s+одной\s+стороны\b", re.I)),
]

EMU = 914400.0


def vidimyy(slide):
    """Видимый текст слайда одной строкой. БЕЗ заметок — они считаются отдельно."""
    out = []

    def walk(shapes):
        for sh in shapes:
            if sh.shape_type == 6:                      # группа
                walk(sh.shapes)
                continue
            if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    t = "".join(r.text for r in p.runs).strip()
                    if t:
                        out.append(t)
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            out.append(cell.text.strip())

    walk(slide.shapes)
    return "\n".join(out)


def zametka(slide):
    if not slide.has_notes_slide:
        return ""
    return slide.notes_slide.notes_text_frame.text


def zamer(put):
    """[(номер, слой, имя фигуры, цитата)] плюс объём слоёв в словах."""
    prs = Presentation(str(put))
    nayden, slov = [], {"экран": 0, "речь": 0}
    for i, slide in enumerate(prs.slides, 1):
        for sloy, tekst in (("экран", vidimyy(slide)), ("речь", zametka(slide))):
            slov[sloy] += len(tekst.split())
            for imya, rx in FIGURY:
                for m in rx.finditer(tekst):
                    a = max(0, m.start() - 34)
                    nayden.append((i, sloy, imya,
                                   tekst[a:m.end() + 34].replace("\n", " ")))
    return nayden, slov, len(prs.slides)


def otchet(put, spisok=False):
    nayden, slov, vsego = zamer(put)
    print(f"\n{Path(put).name} — {vsego} слайдов, "
          f"экран {slov['экран']} слов, речь {slov['речь']} слов")
    for sloy in ("экран", "речь"):
        podytog = 0
        for imya, _rx in FIGURY:
            k = sum(1 for n in nayden if n[1] == sloy and n[2] == imya)
            podytog += k
            if k:
                print(f"  {sloy:<6} {imya:<22} {k:>3}")
        plotnost = podytog / slov[sloy] * 1000 if slov[sloy] else 0
        print(f"  {sloy:<6} {'ИТОГО по слою':<22} {podytog:>3}"
              f"   ({plotnost:.1f} на 1000 слов)")
    if spisok:
        print()
        for nomer, sloy, imya, cit in nayden:
            print(f"  слайд {nomer:>3} {sloy:<6} {imya:<22} …{cit}…")
    print(f"  {'ВСЕГО':<30} {len(nayden):>3}")
    return len(nayden)


# --------------------------------------------------------------------------
# Проверка на нарочно сломанном входе. Без неё непонятно, умеет ли проверка
# вообще находить и называет ли величину.
SLOMANNOE = [
    ("головное — чистый случай",
     "Чужой скилл — не текст, а каталог решений.", "головное «не X, а Y»", 1),
    ("хвостовое — чистый случай",
     "Переносится решение, а не запись в журнале.", "хвостовое «, а не»", 1),
    ("ЛОВУШКА КАНОНА: хвост слова на «-не»",
     "Порог написан в плане ступени, а проверка стоит рядом.",
     "головное «не X, а Y»", 0),
    ("ЛОВУШКА КАНОНА: «на стороне платформы, а …»",
     "Защита стоит на стороне платформы, а барьер — у агента.",
     "головное «не X, а Y»", 0),
    ("хвостовое не путается с «а неделя»",
     "Срок мерили в днях, а неделя прошла впустую.", "хвостовое «, а не»", 0),
    ("«дело не в X»",
     "Дело не в длине файла — читается первая строка.", "«дело не в X»", 1),
    # Пара «с одной стороны… с другой» считается ОДНИМ вхождением: канон
    # называет фигуру целиком, а метка у неё одна. Одинокое «с одной стороны»
    # без второй половины — та же фигура, поэтому метка и ловится сама по себе.
    ("«с одной стороны» — пара считается один раз",
     "С одной стороны барьер, с другой — просьба.", "«с одной стороны»", 1),
    ("три фигуры в одной строке считаются все три",
     "Дело не в цифре, а в базе; с одной стороны это видно.",
     None, 3),
]


def self_test():
    """Печатает ДЕФЕКТ И ЕГО ВЕЛИЧИНУ на каждом сломанном входе."""
    plohih = 0
    print("Прогон на нарочно сломанном входе\n")
    for imya, tekst, figura, zhdem in SLOMANNOE:
        if figura is None:
            bylo = sum(len(rx.findall(tekst)) for _i, rx in FIGURY)
        else:
            rx = dict((i, r) for i, r in FIGURY)[figura]
            bylo = len(rx.findall(tekst))
        ok = bylo == zhdem
        plohih += not ok
        print(f"  {'ок ' if ok else 'ПЛОХО'} ждём {zhdem}, нашли {bylo:<3} "
              f"{imya}\n        «{tekst}»")
    print(f"\n  сломанных случаев: {len(SLOMANNOE)}, "
          f"проверка промолчала на {plohih}")
    return 1 if plohih else 0


def main():
    argv = [a for a in sys.argv[1:]]
    if "--self-test" in argv:
        return self_test()
    spisok = "--spisok" in argv
    argv = [a for a in argv if not a.startswith("--")]
    puti = argv or [str(HERE / "sem-05.pptx")]
    vsego = 0
    for p in puti:
        vsego += otchet(p, spisok)
    return 1 if vsego else 0


if __name__ == "__main__":
    sys.exit(main())
