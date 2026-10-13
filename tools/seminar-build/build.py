#!/usr/bin/env python3
"""Сборка деки семинара — одна на все занятия курса.

    python3 tools/seminar-build/build.py --seminar 07 --all

До октября 2026 этот сборщик жил копией в папке каждого занятия. Копии
расходились, и расходились НЕЗАМЕТНО: `deck_kit.py`, `metrics.py`,
`slide_parts.py` в папках Семинаров 6 и 7 лежали побайтово одинаковые, а
`build_semNN.py` отличались 1247 строками из 1400 — при том что почти всё
отличие было таблицами (границы блоков, словарь приёмов, список слайдов с
золотой цитатой), а не вёрсткой. Починка в одной копии не доезжала до
остальных: `**жирный**`, рассечённый переносом строки, был найден и починен
на Семинаре 7, и до Семинара 6 эта починка не дошла бы никогда.

Что общее (здесь): вёрстка (`deck_kit`), измерение (`metrics`), разбор
слайда (`slide_parts`), роли блоков (`roles`), библиотека жанровых приёмов
(`genres`), сборочные аудиты (`audits`), генераторы манифеста и карты.
Что своё у занятия: `library/seminars/sem-NN/rendered/deck-config.yaml` —
разметка, словарь `visual.pattern → жанр`, пути; и `make_figures_*.py` —
схемы каждого занятия свои.

Устройство и порядок чтения: `tools/seminar-build/README.md`.
"""

import sys
from pathlib import Path

import yaml
from pptx import Presentation
from pptx.util import Inches

import deck_config as D
import deck_kit as K
import metrics as M
import slide_parts as SP
import audits
import genres
from roles import quote_role, sanitize_blocks

USAGE = """build.py — сборка деки семинара (одна на все занятия курса).

  --seminar NN  НОМЕР ЗАНЯТИЯ. Обязателен, если команда запущена не из папки
                занятия (иначе номер берётся из пути).
  --block nNN   ПРЕДПРОСМОТР одного слайда или группы: строит из файлов на
                диске в отдельный `sem-NN-nMM.pptx`. Безопасно при
                параллельной работе: `deck.yaml` и `rendered/sem-NN.pptx`
                не трогает.
  --all         ПОЛНАЯ пересборка: пересобирает `deck.yaml` из слайдов и
                перезаписывает `rendered/sem-NN.pptx`. Общие файлы, то есть
                ход сведения, а не проверка своей зоны.
  --deck F      собрать по другому манифесту (в `sem-NN-F.pptx`).
  nNN nMM …     собрать полотно только этих слайдов, остальные страницы
                оставить пустыми (отладка вёрстки).
  --lang en     собирать АНГЛИЙСКУЮ деку: слайды `slides-en/`, манифест
                `deck.en.yaml`, выход `sem-NN-en.pptx`, и — главное — реплика
                опознаётся по `"`, а не по «ёлочке» (см. `LANG_SIGNALS`).
                Русскую деку и русский манифест при этом не трогает ничем.
                Занятие без английского трека в конфиге отказывается сразу.
  --quote-roles НИЧЕГО НЕ СОБИРАЕТ И НЕ ПИШЕТ: печатает, какую роль получит
                каждый блок-цитата деки. С `--lang en` — то же по английской.
                Нужно, чтобы сверить роли двух дек ДО рендера: роли обязаны
                совпасть слайд в слайд, иначе язык меняет не текст, а вёрстку.

Полная пересборка требует ЯВНОГО `--all` и без него не запускается. Так
решено после круга 2 замечаний владельца (issue 225): четыре сессии круга
независимо запустили полную пересборку, набрав `--help` или перепутав флаг,
и перезаписали общую деку — расплата за то, что «ничего не передано» значило
«сделай самое разрушительное».
"""

FLAGS = ("--seminar", "--block", "--deck", "--all", "--lang", "--quote-roles")


def seminar_from_cwd():
    """Номер занятия из текущего каталога — чтобы привычное «запустить из
    своей папки» продолжало работать без флага."""
    for p in [Path.cwd()] + list(Path.cwd().parents):
        if p.name.startswith("sem-") and p.parent.name == "seminars":
            return p.name[4:]
    return None


def refresh_manifest():
    """Пересобрать манифест из его слайдов перед сборкой — манифест
    производное, правят слайд, а не `deck.yaml` (канон §8.5)."""
    import difflib
    name = D.deck_file()
    try:
        import make_deck_yaml as G
        doc, slides = G.build("n", D.LANG)
        text = G.header_comment(slides, doc["deck"], D.LANG) + yaml.safe_dump(
            doc, allow_unicode=True, sort_keys=False, width=100,
            default_flow_style=False)
    except Exception as e:
        M._WARNINGS.append(f"МАНИФЕСТ: пересобрать не удалось — {e}. "
                           f"Собираю по тому, что лежит в {name}")
        return
    path = D.ROOT / name
    was = path.read_text(encoding="utf-8") if path.exists() else ""
    if was == text:
        return
    path.write_text(text, encoding="utf-8")
    D.invalidate()
    changed = [l for l in difflib.unified_diff(was.splitlines(), text.splitlines(), n=0)
               if l[:1] in "+-" and l[:3] not in ("+++", "---")]
    head = "; ".join(l.strip() for l in changed[:4])
    M._WARNINGS.append(
        f"МАНИФЕСТ ПЕРЕСОБРАН: {name} был старше слайдов, строк разошлось "
        f"{len(changed)} — {head}{'…' if len(changed) > 4 else ''}")


def print_quote_roles(only):
    """Какую роль получит каждый блок-цитата — БЕЗ сборки и без записи файлов.

    Зачем отдельный режим, а не «собери и посмотри глазами»: роль видна на
    картинке цветом полосы, то есть проверяется рендером, PDF и просмотром
    шестидесяти снимков. Между языками сверять надо не картинку, а решение,
    и решение печатается здесь одной строкой на блок.

    Контракт для EN-деки: набор ролей обязан СОВПАСТЬ с русским слайд в
    слайд. Расхождение значит, что язык поменял не текст, а вёрстку, — и это
    дефект независимо от того, какая из двух раскрасок красивее."""
    deck_path = D.ROOT / D.deck_file()
    if deck_path.exists():
        slides = yaml.safe_load(deck_path.read_text(encoding="utf-8"))["slides"]
    else:
        slides = D.deck_from_files("n")
    gold = D.table("gold_quote_sids")
    tally, total = {}, 0
    for sv in slides:
        sid = sv["id"]
        if only and sid not in set(only):
            continue
        pattern = (sv.get("visual") or {}).get("pattern", "")
        _t, _a, visual, _n = SP.sections((D.ROOT / sv["file"]).read_text(encoding="utf-8"))
        for kind, b in SP.blocks(visual):
            if kind != "quote":
                continue
            role = quote_role(b, pattern)
            if gold.get(sid) == "all":
                role += " (+gold по gold_quote_sids)"
            head = K.plain(" ".join(b))[:58]
            tally[role.split(" ")[0]] = tally.get(role.split(" ")[0], 0) + 1
            total += 1
            print(f"{sid}  {pattern:28} {role:30} {head}…")
    print(f"\nязык: {D.LANG}   знак реплики: {' / '.join(D.SPEECH_OPEN)}   "
          f"блоков-цитат: {total}")
    for r, v in sorted(tally.items()):
        print(f"  {r:10} {v:3}")
    return 0


def main():
    argv = list(sys.argv[1:])
    if {"--help", "-h", "help"} & set(argv):
        print(USAGE)
        return
    unknown = [a for a in argv if a.startswith("-") and a not in FLAGS]
    if unknown:
        print(f"неизвестный флаг: {' '.join(unknown)}\n")
        print(USAGE)
        raise SystemExit(2)

    seminar = None
    if "--seminar" in argv:
        i = argv.index("--seminar")
        seminar = argv[i + 1] if i + 1 < len(argv) else ""
        del argv[i:i + 2]
    seminar = seminar or seminar_from_cwd()
    if not seminar:
        print("не сказано, какое это занятие: передайте `--seminar NN` "
              "или запускайте из папки семинара.\n")
        print(USAGE)
        raise SystemExit(2)
    D.configure(seminar)

    if "--lang" in argv:
        i = argv.index("--lang")
        D.set_lang(argv[i + 1] if i + 1 < len(argv) else "")
        del argv[i:i + 2]
    if "--quote-roles" in argv:
        argv.remove("--quote-roles")
        return print_quote_roles(argv)

    full = "--all" in argv
    if full:
        argv.remove("--all")
    block = None
    if "--block" in argv:
        i = argv.index("--block")
        block = argv[i + 1] if i + 1 < len(argv) else "n"
        del argv[i:i + 2]

    deck_file = D.deck_file()
    if "--deck" in argv:
        i = argv.index("--deck")
        deck_file = argv[i + 1]
        del argv[i:i + 2]

    stem = f"sem-{D.SEM}" if D.LANG == "ru" else f"sem-{D.SEM}-{D.LANG}"
    if block:
        slides = D.deck_from_files(block)
        out_name = f"{stem}-{block}.pptx"
        if not slides:
            print(f"слайдов по образцу «{block}*.md» не найдено")
            return
    elif not full and deck_file == D.deck_file():
        print("полная пересборка деки не запрошена.\n")
        print(USAGE)
        raise SystemExit(2)
    else:
        deck_path = D.ROOT / deck_file
        if deck_path.exists():
            slides = yaml.safe_load(deck_path.read_text(encoding="utf-8"))["slides"]
        elif deck_file == D.deck_file():
            # Самая первая сборка занятия: манифеста ещё не существует вовсе
            # (не «устарел», а именно отсутствует). `refresh_manifest()`
            # создаст его из слайдов; здесь достаточно пустого списка, чтобы
            # не упасть раньше срока.
            slides = []
        else:
            raise SystemExit(f"файла «{deck_file}» нет — и это не deck.yaml, "
                             f"чтобы пересобрать его автоматически")
        out_name = D.out_name() if deck_file == D.deck_file() else \
            f"sem-{D.SEM}-{Path(deck_file).stem}.pptx"

    deck = {"slides": slides}
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(K.W_IN), Inches(K.H_IN)
    blank = prs.slide_layouts[6]
    M.reset()
    if not block and deck_file == D.deck_file():
        refresh_manifest()
        # Перечитать ПОСЛЕ пересборки: на бутстрапе занятия чтение выше было
        # пустым, и цикл сборки прошёл бы по пустому списку, молча сохранив
        # пустую деку.
        slides = yaml.safe_load((D.ROOT / deck_file).read_text(encoding="utf-8"))["slides"]
        deck["slides"] = slides

    only = set(argv)
    built = 0
    for s in deck["slides"]:
        sid = s["id"]
        pattern = (s.get("visual") or {}).get("pattern", "")
        md = (D.ROOT / s["file"]).read_text(encoding="utf-8")
        title, _assertion, visual, notes = SP.sections(md)
        sl = prs.slides.add_slide(blank)
        blocks = sanitize_blocks(SP.blocks(visual))
        if not only or sid in only:
            fn = genres.genre_fn(pattern)
            kw = {"meta": s.get("visual")} if fn is genres.g_divider else {}
            fn(sl, sid, title or sid, blocks, pattern, _assertion, **kw)
        audits.audit_question_answer(sid, pattern, blocks)
        if notes:
            K.write_notes(sl, notes)
        built += 1

    audits.audit_numbering(deck["slides"])
    audits.audit_crossrefs(deck["slides"])
    audits.audit_figs(slides)
    audits.audit_figure_text()
    out = D.HERE / out_name
    prs.save(out)
    warns = M.report()
    for w in warns:
        print("•", w)
    print(f"\nслайдов: {built}   предупреждений: {len(warns)}   "
          f"файл: {out.name} ({out.stat().st_size // 1024} КБ)")


if __name__ == "__main__":
    main()
