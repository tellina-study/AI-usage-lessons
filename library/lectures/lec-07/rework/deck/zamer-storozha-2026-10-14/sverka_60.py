# -*- coding: utf-8 -*-
"""Сверка: 60 прогонов отчёта прожарки против счёта пробы контраста."""
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[6]
                       / "tools" / "presentation-build"))
from pptx import Presentation
from deckkit import contrast

P = "lec-07-karkas.pptx"

# — запрос прожарки, §1, дословно: цвет F0AB00 и кегль >= 16, без различия
#   полужирности и без различия фона
p = Presentation(P)
theirs = []
for n, sl in enumerate(p.slides, 1):
    for s in sl.shapes:
        if not s.has_text_frame:
            continue
        for par in s.text_frame.paragraphs:
            for r in par.runs:
                try:
                    c = str(r.font.color.rgb)
                except Exception:
                    continue
                if c == "F0AB00" and (r.font.size.pt if r.font.size else 0) >= 16:
                    theirs.append((n, round(r.font.size.pt, 1),
                                   bool(r.font.bold), r.text.strip()[:28]))
print(f"запрос отчёта (F0AB00, кегль >= 16):     {len(theirs)} прогонов, "
      f"{len({t[0] for t in theirs})} кадров")

rows = contrast.rows_of(p)
gold = [r for r in rows if r["fg"] == "F0AB00"]
bad = [r for r in rows if r["ok"] is False]
large = [r for r in bad if contrast.is_large(r["size"], r["bold"])]
small = [r for r in bad if not contrast.is_large(r["size"], r["bold"])]
print(f"проба: золото всего                      {len(gold)} прогонов, "
      f"{len({r['page'] for r in gold})} кадров")
print(f"проба: ниже порога ВСЕГО                 {len(bad)} прогонов, "
      f"{len({r['page'] for r in bad})} кадров")
print(f"  из них крупный класс (порог 3,0)       {len(large)} прогонов, "
      f"{len({r['page'] for r in large})} кадров")
print(f"  из них обычный класс (порог 4,5)       {len(small)} прогонов, "
      f"{len({r['page'] for r in small})} кадров")
print(f"кегли золота ниже порога: "
      + ", ".join(f"{s} пт×{[r['size'] for r in bad].count(s)}"
                  for s in sorted({r['size'] for r in bad})))

t_keys = sorted({(t[0], t[3]) for t in theirs})
m_keys = sorted({(r["page"], r["text"][:28]) for r in large})
print(f"\nв запросе отчёта, но не в крупном классе пробы: "
      f"{[k for k in t_keys if k not in m_keys]}")
print(f"в крупном классе пробы, но не в запросе отчёта: "
      f"{[k for k in m_keys if k not in t_keys]}")
print("\nобычный класс (кегль < 18 и не полужирный от 14) — "
      "кадры и кегли:")
for r in small:
    print(f"  стр {r['page']:>3} {r['size']:>5} пт "
          f"{'полужирный' if r['bold'] else 'светлый':<11} "
          f"фон {r['bg']} ({r['bg_src']}) «{r['text'][:34]}»")
