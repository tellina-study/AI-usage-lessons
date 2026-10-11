#!/usr/bin/env python3
"""Смоук-сборка трёх пробных слайдов `rendered/smoke/x0*.md` в
`rendered/smoke/sem-06-smoke.pptx` — проверка, что путь сборки Семинара 6
живой, ДО того как настоящие слайды занятия дописаны (три параллельные
сессии пишут `../../slides/n*.md` прямо сейчас).

Не трогает `../../slides/` и не трогает `../../deck.yaml` — собирает деку
ИЗ ТЕКСТА ЭТОГО КАТАЛОГА, минуя `deck_from_files`/`refresh_manifest`
(которые читают `ROOT/slides`, то есть каталог НАСТОЯЩИХ слайдов занятия).

Идентификаторы смоук-слайдов (`x01`, `x02`, `x03`) намеренно вне словаря
`nNN`: `build_sem06._deck_slides("x")` не находит для префикса `x` ни
`deck.yaml`, ни файлов в `slides/x*.md` и возвращает пустой список — это
значит, что `slide_number()`/`badge_for()` честно отдадут `None` (см.
`deck_kit.slide_number_mark`: `None` — не печатать номер, у слайда, которого
нет в деке, порядкового номера не существует), а `sections_for()` упадёт на
статический запасной `SECTIONS_BY_PREFIX["n"]`. Это ожидаемо и безопасно —
смоук-слайд не является частью деки и не должен притворяться, что является.

    python3 smoke_build.py
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

HERE = Path(__file__).resolve().parent
RENDERED = HERE.parent
sys.path.insert(0, str(RENDERED))

import build_sem06 as B   # noqa: E402  (путь добавлен выше)
import deck_kit as K      # noqa: E402
import metrics as M       # noqa: E402
import slide_parts as SP  # noqa: E402

FILES = ["x01-zaglushka-divider.md",
         "x02-zaglushka-razbor-tri-kolonki.md",
         "x03-zaglushka-code-artifact.md"]


def load(name):
    import re
    import yaml
    md = (HERE / name).read_text(encoding="utf-8")
    m = re.match(r"^---\n(.*?)\n---\n", md, re.S)
    fm = (yaml.safe_load(m.group(1)) if m else {}) or {}
    title, assertion, visual, notes = SP.sections(md)
    return fm, title, assertion, visual, notes


def main():
    M.reset()
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(K.W_IN), Inches(K.H_IN)
    blank = prs.slide_layouts[6]

    slides_for_audit = []
    for name in FILES:
        fm, title, assertion, visual, notes = load(name)
        sid = fm["id"]
        pattern = (fm.get("visual") or {}).get("pattern", "")
        sl = prs.slides.add_slide(blank)
        fn = B.GENRE_FN.get(pattern, B.g_content)
        kw = {"meta": fm.get("visual")} if fn is B.g_divider else {}
        fn(sl, sid, title or sid, SP.blocks(visual), pattern, assertion, **kw)
        B.audit_question_answer(sid, pattern, SP.blocks(visual))
        if notes:
            K.write_notes(sl, notes)
        slides_for_audit.append(fm)

    B.audit_numbering(slides_for_audit)
    B.audit_crossrefs(slides_for_audit)
    B.audit_figs(slides_for_audit)

    out = HERE / "sem-06-smoke.pptx"
    prs.save(out)
    warns = M.report()
    for w in warns:
        print("•", w)
    print(f"\nсмоук-слайдов: {len(FILES)}   предупреждений: {len(warns)}   "
          f"файл: {out.name} ({out.stat().st_size // 1024} КБ)")


if __name__ == "__main__":
    main()
