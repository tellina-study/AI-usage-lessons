"""Библиотека приёмов вёрстки деки семинара.

Перенесена с Семинара 4 (`sem-04/rendered/build_sem04.py`) — но перенесён
именно СЛОВАРЬ ПРИЁМОВ, а не 57 функций `build_sNN`, намертво привязанных к
содержанию того занятия. Каждый приём здесь принимает координаты и данные и
возвращает высоту, которую занял, — поэтому вёрстка идёт от курсора, а не от
жёстких координат, и следующий блок всегда знает, где кончился предыдущий.

Правила, которые этот модуль держит:

* Высота считается по ИЗМЕРЕННОМУ тексту (`metrics`), а не по числу элементов.
* Таблица — это скруглённая коробка с текстовыми блоками и тонкими линиями,
  а не сетка PowerPoint: высоты строк тогда считает вёрстка, а не PowerPoint,
  и строка не вырастает под текстом уже после сборки.
* Подсветка означает смысл и ничего другого: `GOLD_TINT` — целевой ответ,
  `TEAL_TINT` — нейтральная отметка, пунктирная золотая рамка — выбранный
  вариант.
* Цвета — только именованные константы. Литералов в вёрстке нет.
"""
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from lxml import etree

import metrics as M

# ── Палитра Ocean Gradient (зафиксирована курсом) ────────────────────────────
DEEP  = RGBColor(0x21, 0x29, 0x5C)   # тёмно-синий: фон дивайдеров, плашка заголовка
MID   = RGBColor(0x06, 0x5A, 0x82)   # средний: надзаголовки на светлом, шапки
LIGHT = RGBColor(0x1C, 0x72, 0x93)   # светлый: обводка коробок
TEAL  = RGBColor(0x02, 0x80, 0x90)   # вторичный: ярлык жанра, маркеры чек-листа,
                                     # полоса микро-дивайдера
GOLD  = RGBColor(0xF0, 0xAB, 0x00)   # акцент: вопрос, целевой ответ, полоса раздела

# Золото как ТЕКСТ на светлом фоне не проходит WCAG AA — известный дефект
# палитры. Для текста и значков на светлом берётся тёмно-золотой.
GOLD_DARK = RGBColor(0x8A, 0x62, 0x00)
GOLD_TINT = RGBColor(0xFE, 0xF5, 0xE0)   # подложка строки-целевого ответа
TEAL_TINT = RGBColor(0xE0, 0xF1, 0xF2)   # подложка нейтральной отметки

SURFACE   = RGBColor(0xF4, 0xF7, 0xFA)   # поверхность коробки на светлом
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
INK       = RGBColor(0x14, 0x1B, 0x2E)   # основной текст на светлом
SLATE     = RGBColor(0x6B, 0x76, 0x85)   # шапки таблиц, подписи, сноски
SOFT_GREY = RGBColor(0xE5, 0xEA, 0xF0)   # разделители строк, спокойная обводка
GREY_FILL = RGBColor(0xE4, 0xE9, 0xEF)   # незаполненный слот

# На тёмном фоне
ON_DARK      = RGBColor(0xD6, 0xE2, 0xEC)   # абзац на тёмном
ON_DARK_MUTE = RGBColor(0x9F, 0xAE, 0xC4)   # подпись, номер слайда, неактивная пилюля
DEEP_PANEL   = RGBColor(0x2A, 0x34, 0x70)   # коробка на тёмном фоне
DEEP_INK     = RGBColor(0x0B, 0x14, 0x3A)   # плашка ярлыка на дивайдере
PILL_OFF     = RGBColor(0x3E, 0x4C, 0x8A)   # пройденная/будущая ступень в полосе

CODE_BG = RGBColor(0x16, 0x1C, 0x30)
CODE_FG = RGBColor(0xE3, 0xE9, 0xF2)

W_IN, H_IN = 13.333, 7.5
FONT, MONO = "Arial", "Consolas"
NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


# ── Разметка внутри строки: **жирный** и `моноширинный` ──────────────────────

def plain(s):
    """Текст без разметки — им меряется ширина и высота."""
    import re
    s = re.sub(r"\*\*(.+?)\*\*", r"\1", s)
    s = re.sub(r"`(.+?)`", r"\1", s)
    s = re.sub(r"\[(.+?)\]\(.+?\)", r"\1", s)
    return s.strip()


def inline_runs(s):
    """Разбор строки на прогоны [(текст, жирный, моноширинный)].

    Прежний `strip_md()` ВЫРЕЗАЛ `**` вместо того, чтобы превратить его в
    жирный прогон, — и целевой ответ, помеченный в исходнике жирным, выходил
    на слайд обычной строкой среди таких же. Здесь разметка становится
    оформлением, а не теряется.
    """
    return M.inline_segments(s)


def is_bold(s):
    """Строка целиком помечена жирным — признак целевого ответа."""
    st = s.strip()
    return st.startswith("**") and st.endswith("**") and len(st) > 4


# ── Низкоуровневые примитивы ────────────────────────────────────────────────

def set_bg(sl, color):
    f = sl.background.fill
    f.solid()
    f.fore_color.rgb = color


def _no_shadow(shp):
    sppr = shp._element.spPr
    for el in sppr.findall(NS + "effectLst"):
        sppr.remove(el)
    etree.SubElement(sppr, NS + "effectLst")


def text_box(sl, x, y, w, h, lines, *, size=13, bold=False, italic=False,
             color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             mono=False, spacing=1.18, space_after=0, rich=True):
    """Текстовая рамка. `lines` — строка или список абзацев; **жирный** и
    `моноширинный` внутри становятся прогонами, а не вырезаются."""
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0)
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    ls = lines if isinstance(lines, (list, tuple)) else [lines]
    for i, ln in enumerate(ls):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(space_after)
        for seg, b, mo in (inline_runs(ln) if rich else [(ln, False, False)]):
            r = p.add_run()
            r.text = seg
            r.font.name = MONO if (mono or mo) else FONT
            r.font.size = Pt(size)
            r.font.bold = bold or b
            r.font.italic = italic
            r.font.color.rgb = color
    return tb


def ocean_box(sl, x, y, w, h, *, fill=SURFACE, stroke=LIGHT, stroke_pt=1.5, radius_pt=12.0):
    """Визуальный мотив курса: скруглённая коробка, радиус 12."""
    shp = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                              Inches(x), Inches(y), Inches(w), Inches(h))
    try:
        shp.adjustments[0] = max(0.04, min(0.25, (radius_pt / 72.0) / max(min(w, h) / 2.0, 0.5)))
    except Exception:
        pass
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if stroke is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = stroke
        shp.line.width = Pt(stroke_pt)
    _no_shadow(shp)
    return shp


def rect(sl, x, y, w, h, fill, *, stroke=None, stroke_pt=0.0, radius=False, radius_adj=0.16):
    shp = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                              Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if stroke is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = stroke
        shp.line.width = Pt(stroke_pt)
    if radius:
        try:
            shp.adjustments[0] = radius_adj
        except Exception:
            pass
    _no_shadow(shp)
    return shp


def dashed_box(sl, x, y, w, h, *, fill=SURFACE, stroke=GOLD_DARK, stroke_pt=1.6, radius_pt=12.0):
    """Пунктирная рамка. Два смысла и только два: выбранный вариант (золотом)
    и ПУСТОЙ СЛОТ — место, которое будет заполнено. Пустой слот обязан
    выглядеть как слот, а не как сломанная ячейка таблицы."""
    shp = ocean_box(sl, x, y, w, h, fill=fill, stroke=stroke, stroke_pt=stroke_pt, radius_pt=radius_pt)
    ln = shp.line._get_or_add_ln()
    etree.SubElement(ln, NS + "prstDash").set("val", "dash")
    return shp


def gradient_rect(sl, x, y, w, h, stops):
    """Градиентная заливка. DEEP→MID→LIGHT на дивайдерах — приём Семинара 4:
    плоский тёмный фон делает все дивайдеры одинаковыми."""
    shp = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.line.fill.background()
    _no_shadow(shp)
    sppr = shp._element.spPr
    for tag in ("noFill", "solidFill", "gradFill", "blipFill", "pattFill"):
        for el in sppr.findall(NS + tag):
            sppr.remove(el)
    grad = etree.SubElement(sppr, NS + "gradFill")
    grad.set("rotWithShape", "1")
    gs_lst = etree.SubElement(grad, NS + "gsLst")
    for pos, color in stops:
        gs = etree.SubElement(gs_lst, NS + "gs")
        gs.set("pos", str(int(pos)))
        etree.SubElement(gs, NS + "srgbClr").set("val", "%02X%02X%02X" % (color[0], color[1], color[2]))
    lin = etree.SubElement(grad, NS + "lin")
    lin.set("ang", "2700000")
    lin.set("scaled", "1")
    return shp


def hairline(sl, x1, y, x2, *, color=SOFT_GREY, pt=0.75):
    """Тонкая линия-разделитель строк. Именно она заменяет рамки сетки."""
    ln = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                 Inches(x1), Inches(y), Inches(x2), Inches(y))
    ln.line.color.rgb = color
    ln.line.width = Pt(pt)
    return ln


# ── Заголовок слайда ────────────────────────────────────────────────────────

def section_tag(sl, x, y, text, *, color=TEAL, size=11.5, w=12.2):
    """Надзаголовок-ярлык: одна строка тиловыми капителями — «где мы · что это
    за шаг». Самый дешёвый приём Семинара 4 и единственный, который отвечает
    сразу и на «вопрос это или утверждение», и на «мы ещё в том же кейсе».
    В прежней деке этой строки не было нигде."""
    text_box(sl, x, y, w, 0.30, text.upper(), size=size, bold=True, color=color, rich=False)
    return 0.30


HEAD_SIZES = (26, 23, 20.5, 18, 16, 14.5, 13)


def auto_header(sl, label, title, *, x=0.55, w=12.23, y=0.34, color=INK, tag_color=TEAL):
    """Заголовок-утверждение: тёмный на белом, кегль подобран ПО ДЛИНЕ, и
    функция возвращает y, с которого можно начинать содержимое.

    Прежняя вёрстка ставила заголовок жёстко в 24 pt в плашку 1,02″ — отсюда и
    обрезание длинного заголовка, и одинаковый отступ независимо от того,
    одну строку занял заголовок или три."""
    ty = y
    if label:
        section_tag(sl, x, ty, label, color=tag_color, w=w)
        ty += 0.34
    size = M.fit_size(plain(title), w, 1.55, HEAD_SIZES, spacing=1.14, bold=True)
    h = M.text_h(plain(title), size, w, spacing=1.14, bold=True) + 0.06
    M.fits(plain(title), size, w, h, label=f"заголовок «{plain(title)[:24]}»", spacing=1.14, bold=True)
    text_box(sl, x, ty, w, h, title, size=size, bold=True, color=color, spacing=1.14)
    return ty + h + 0.22


def footer_note(sl, text, *, x=0.55, y=None, w=12.23, size=11.5, color=SLATE,
                align=PP_ALIGN.LEFT):
    """Серая подпись-курсив. На слайде-вопросе это «разбор — на следующем
    слайде» — третий сигнал жанра после ярлыка и золотой коробки."""
    h = M.text_h(plain(text), size, w) + 0.06
    if y is None:
        y = H_IN - 0.48 - h
    text_box(sl, x, y, w, h, text, size=size, italic=True, color=color, align=align)
    return h


def slide_id_mark(sl, sid, *, on_dark=False):
    text_box(sl, W_IN - 1.35, H_IN - 0.46, 1.0, 0.3, sid, size=10.5,
             color=ON_DARK_MUTE if on_dark else SLATE, align=PP_ALIGN.RIGHT, rich=False)


# ── Коробка-выноска ─────────────────────────────────────────────────────────

def gold_callout(sl, x, y, w, text, *, size=14, bold=True, align=PP_ALIGN.LEFT,
                 min_h=0.0, max_h=None, pad=0.24, label="выноска"):
    """Золотая коробка. Означает ровно одно: ВОПРОС, который зал должен решить,
    или ФОРМУЛА, которую надо запомнить. Один раз на слайде — иначе акцент
    обнуляется частотой (в прежней деке эта коробка стояла на 33 слайдах из 56
    и держала там утверждения)."""
    lines = text if isinstance(text, (list, tuple)) else [text]
    inner = w - 2 * pad
    if max_h:
        while size > 10.5 and M.rich_block_h(lines, size, inner,
                                        spacing=1.22, space_after=3) + 2 * 0.16 > max_h:
            size -= 0.5
    th = M.rich_block_h(lines, size, inner, spacing=1.22, space_after=3)
    h = max(min_h, th + 0.32)
    if max_h:
        h = min(h, max_h)
    rect(sl, x, y, w, h, GOLD_TINT, stroke=GOLD, stroke_pt=1.5, radius=True, radius_adj=0.1)
    M.fits(" ".join(plain(l) for l in lines), size, inner, h - 0.24, label=label, spacing=1.22)
    text_box(sl, x + pad, y + 0.14, inner, h - 0.28, lines, size=size, bold=bold,
             color=INK, align=align, anchor=MSO_ANCHOR.MIDDLE, spacing=1.22, space_after=3)
    return h


def quote_box(sl, x, y, w, lines, *, size=13.5, max_h=None, on_dark=False, label="цитата"):
    """Спокойная коробка для реплики или пояснения — НЕ золотая. Золото
    остаётся за вопросом и формулой."""
    pad = 0.24
    inner = w - 2 * pad
    if max_h:
        while size > 10 and M.rich_block_h(lines, size, inner, space_after=3) + 0.3 > max_h:
            size -= 0.5
    h = M.rich_block_h(lines, size, inner, space_after=3) + 0.3
    if max_h:
        h = min(h, max_h)
    if on_dark:
        ocean_box(sl, x, y, w, h, fill=DEEP_PANEL, stroke=LIGHT, stroke_pt=1.2)
        col = WHITE
    else:
        ocean_box(sl, x, y, w, h, fill=SURFACE, stroke=SOFT_GREY, stroke_pt=1.2)
        col = INK
    M.fits(" ".join(plain(l) for l in lines), size, inner, h - 0.26, label=label)
    text_box(sl, x + pad, y + 0.15, inner, h - 0.3, lines, size=size, color=col,
             anchor=MSO_ANCHOR.MIDDLE, space_after=3)
    return h


# ── Таблица как коробка с текстом ───────────────────────────────────────────

CELL_SIZES = (12.5, 12, 11.5, 11, 10.5, 10, 9.5, 9)

# ── Подгонка таблицы: что уступать, когда 0,05″ не хватает ──────────────────
#
# Прежде переполнение таблицы только ОБЪЯВЛЯЛОСЬ. Кегль сбавлялся по лестнице
# `CELL_SIZES`, и если не помогал даже 9 pt, вёрстка печатала предупреждение и
# рисовала как есть — таблица вылезала на 0,05–0,20″ за отведённое. Между тем
# кегль — не единственная уступка: межстрочный, отступ строки, минимальная
# высота строки и поля коробки вместе дают ещё около 15% высоты, и отдавать их
# правильнее, чем ещё полступени кегля.
#
# Каждая ступень — (межстрочный, отступ строки, мин. высота строки, поле):
TIGHTEN = ((1.15, 0.18, 0.34, 0.20),      # штатно, ничего не уступлено
           (1.09, 0.14, 0.30, 0.17),      # ступень 1: ~7% высоты
           (1.04, 0.11, 0.26, 0.14))      # ступень 2: ещё ~6%
#
# Порядок перебора: СНАЧАЛА вся лестница кеглей при штатном сжатии, и только
# если не подошёл ни один кегль — следующая ступень сжатия, снова по всем
# кеглям. Сжатие тут последнее средство, а не равноправная уступка.
#
# Смешанный порядок (цена = номер кегля + номер ступени) пробовался и отвергнут
# по измерению: на этой деке он уводил на сжатие 10 таблиц из 28 — то есть
# делал сжатый вид штатным, ровно то, чего сжатие должно избегать. При
# лестничном порядке сжатие включается на двух таблицах, которые до этой правки
# просто вылезали за рамку.
def _fit_ladder(sizes):
    return [(si, ti) for ti in range(len(TIGHTEN)) for si in range(len(sizes))]


def _col_shares(headers, rows, ncol, inner_w, size):
    """Доли ширины колонок.

    Ширина пропорциональна МАССЕ текста в колонке — тогда все колонки требуют
    примерно одинакового числа строк, и таблица не растёт из-за одной узкой
    колонки с длинным текстом. Но у массы есть два ограничителя, без которых
    она врёт:

    * ПОЛ — колонке нужно не меньше, чем её самому длинному СЛОВУ, иначе слово
      вылезет за край ячейки, и не меньше 7% ширины.
    * ПОТОЛОК — колонке не нужно БОЛЬШЕ, чем её самой длинной ячейке на одну
      строку. Без потолка колонка из пяти коротких подписей («Хук», «Скилл»,
      «MCP»…) набирает суммарную массу больше, чем соседняя с одним длинным
      заголовком, и забирает половину таблицы, хотя каждой её ячейке хватает
      дюйма. Ровно это и происходило на слайде оси занятия.

    Остаток после потолков разливается по колонкам, которым ещё тесно.
    """
    mass, floor, ceil = [], [], []
    for j in range(ncol):
        cells = [r[j] for r in rows if j < len(r) and r[j]]
        if headers and j < len(headers):
            cells.append(headers[j])
        cells = [(c, plain(c)) for c in cells]
        widths = [M.text_w(t, size, mono="`" in raw) for raw, t in cells] or [0.35]
        mass.append(max(sum(widths), 0.35))
        longest = max((M.text_w(wd, size, mono="`" in raw)
                       for raw, t in cells for wd in t.split()), default=0.4)
        floor.append(max(longest + 0.2, inner_w * 0.07))
        ceil.append(max(max(widths) + 0.24, floor[-1]))

    if sum(floor) >= inner_w:                       # места нет даже на полы
        k = inner_w / sum(floor)
        return [f * k for f in floor]

    share = [0.0] * ncol
    free = set(range(ncol))
    pool = inner_w
    for _ in range(ncol + 2):
        if not free:
            break
        tot = sum(mass[j] for j in free) or 1.0
        clamped = False
        for j in sorted(free):
            want = pool * mass[j] / tot
            if want < floor[j]:
                share[j] = floor[j]; pool -= floor[j]; free.discard(j); clamped = True
            elif want > ceil[j]:
                share[j] = ceil[j]; pool -= ceil[j]; free.discard(j); clamped = True
        if not clamped:
            tot = sum(mass[j] for j in free) or 1.0
            for j in free:
                share[j] = pool * mass[j] / tot
            free = set()
            break
    if free:
        tot = sum(mass[j] for j in free) or 1.0
        for j in free:
            share[j] = max(pool * mass[j] / tot, floor[j])

    # Остаток разливается ТОЛЬКО по колонкам, которым ещё тесно (доля ниже
    # потолка), пропорционально нехватке — и каждой не выше её потолка. Отдавать
    # остаток по массе нельзя: колонка из пяти коротких подписей имеет большую
    # массу и снова забрала бы место, которое ей не нужно.
    for _ in range(ncol + 2):
        rest = inner_w - sum(share)
        if rest <= 0.01:
            break
        hungry = [j for j in range(ncol) if share[j] < ceil[j] - 0.01]
        if not hungry:
            for j in range(ncol):                   # содержимое у́же таблицы —
                share[j] += rest / ncol             # лишнее идёт в поля поровну
            break
        gapsum = sum(ceil[j] - share[j] for j in hungry)
        for j in hungry:
            share[j] = min(ceil[j], share[j] + rest * (ceil[j] - share[j]) / gapsum)
    k = inner_w / sum(share)
    return [x * k for x in share]


def table_card(sl, x, y, w, headers, rows, *, col_w=None, highlight=None,
               max_h=None, pad=0.2, gap=0.16, min_row_h=0.34, label="таблица",
               header_size=None, cell_size=None, align=None, accent=None):
    """Таблица, собранная руками: скруглённая коробка, отступ 0,2″, шапка
    СЕРЫМИ КАПИТЕЛЯМИ БЕЗ ЗАЛИВКИ, разделители строк — тонкие линии 0,75 pt,
    целевая строка подсвечена `GOLD_TINT`.

    Семинар 4 не вызвал `add_table()` ни разу — и правильно: сетка PowerPoint
    с рамками, синей шапкой и зеброй читается как вставленный из Excel
    скриншот, а высоту строк PowerPoint определяет уже после сборки, поэтому
    таблица вырастает вниз и наезжает на то, что под ней. Здесь высоты строк
    считает вёрстка, и итоговая высота известна до сохранения файла.

    Строка, помеченная в исходнике **жирным**, распознаётся как целевой ответ
    и подсвечивается автоматически — это и есть «подсветка = смысл».
    Возвращает занятую высоту.
    """
    ncol = max([len(r) for r in rows] + ([len(headers)] if headers else [1]))
    rows = [list(r) + [""] * (ncol - len(r)) for r in rows]
    if highlight is None:
        # Подсветка = ЦЕЛЕВОЙ ОТВЕТ, и только он. Поэтому строка берётся в
        # золото, лишь когда жирным помечены ВСЕ её заполненные ячейки: так
        # автор и размечает разбор вариантов. Одной жирной ячейки мало —
        # в таблице свидетельств жирным набрано название источника в первой
        # колонке, и по правилу «хотя бы одна» подсвечивались случайные строки.
        highlight = {}
        for i, r in enumerate(rows):
            filled = [c for c in r if c.strip()]
            if len(filled) > 1 and all(is_bold(c) for c in filled):
                highlight[i] = "gold"
    inner_w = w - 2 * pad

    sizes = CELL_SIZES if cell_size is None else (cell_size,)

    def geometry(size, tighten):
        """Ширины колонок, высота шапки и высоты строк при заданных кегле и
        ступени сжатия. Ровно тот же расчёт, по которому потом рисуют."""
        spacing, row_pad, mrh, box_pad = tighten
        # Минимальная высота строки сжимается В ДОЛЯХ от заданной вызывающим, а
        # не заменяется табличной: `min(mrh, min_row_h)` молча игнорировал бы
        # вызов, который просит строки ВЫШЕ штатных.
        mrh = min_row_h * (mrh / TIGHTEN[0][2])
        sh = col_w or _col_shares(headers, rows, ncol, w - 2 * box_pad, size)
        hsz = header_size if header_size is not None else min(size, 11.5)
        hh = (max(M.text_h(plain(h_), hsz, sh[j] - gap, bold=True)
                  for j, h_ in enumerate(headers)) + 0.12) if headers else 0.0
        rhs = [max(mrh, max(M.rich_h(c, size, sh[j] - gap, spacing=spacing)
                            for j, c in enumerate(r)) + row_pad) for r in rows]
        return sh, hsz, hh, rhs, box_pad, spacing, row_pad, box_pad * 2 + hh + sum(rhs)

    step = 0
    chosen = sizes[-1]
    geo = geometry(chosen, TIGHTEN[0])
    if max_h is None:
        chosen = sizes[0]
        geo = geometry(chosen, TIGHTEN[0])
    else:
        best = None
        for si, ti in _fit_ladder(sizes):
            g = geometry(sizes[si], TIGHTEN[ti])
            if best is None or g[-1] < best[2][-1]:
                best = (sizes[si], ti, g)          # запасной: самый низкий из всех
            if g[-1] <= max_h:
                chosen, step, geo = sizes[si], ti, g
                break
        else:
            chosen, step, geo = best[0], best[1], best[2]

    shares, hs, head_h, rh, pad, cell_spacing, row_pad, h = geo
    if step:
        # Подгонка совершилась молча — и это ровно то, о чём потом спорят
        # («почему эта таблица мельче соседней?»). Поэтому она остаётся в
        # логе сборки наравне с переполнением: уступка названа и измерена.
        M._WARNINGS.append(
            f"ПОДГОНКА [{label}]: ступень {step} — кегль {chosen} pt, межстрочный "
            f"{cell_spacing}, поле {pad:.2f}″; таблица {h:.2f}″ при бюджете {max_h:.2f}″")
    if max_h is not None and h > max_h + 0.02:
        M._WARNINGS.append(f"ПЕРЕПОЛНЕНИЕ [{label}]: таблица {h:.2f}″ при бюджете "
                           f"{max_h:.2f}″ даже после сжатия — сократить содержание")

    # `accent` различает две таблицы, стоящие на одном слайде подряд: без него
    # «здесь ещё рано» и «здесь не нужно вообще» — две одинаковые сетки друг
    # над другом, хотя это два РАЗНЫХ вопроса (про момент и про задачу).
    head_col = SLATE if accent is None else accent
    ocean_box(sl, x, y, w, h, fill=WHITE,
              stroke=LIGHT if accent is None else accent, stroke_pt=1.4)
    if accent is not None:
        rect(sl, x + 0.02, y + 0.1, 0.075, h - 0.2, accent, radius=True, radius_adj=0.5)
    cx = [x + pad]
    for s in shares[:-1]:
        cx.append(cx[-1] + s)

    if headers:
        for j, ht in enumerate(headers):
            text_box(sl, cx[j], y + pad, shares[j] - gap, head_h, ht.upper(),
                     size=hs, bold=True, color=head_col, spacing=1.1,
                     align=(align[j] if align else PP_ALIGN.LEFT))

    ry = y + pad + head_h
    for i, r in enumerate(rows):
        if i in (highlight or {}):
            tint = GOLD_TINT if highlight[i] == "gold" else TEAL_TINT
            rect(sl, x + 0.07, ry, w - 0.14, rh[i], tint, radius=True, radius_adj=0.08)
        for j, c in enumerate(r):
            if not c.strip():
                # пустая ячейка — это СЛОТ, а не забытая клетка
                dashed_box(sl, cx[j] + 0.02, ry + 0.08, shares[j] - gap - 0.04,
                           rh[i] - 0.16, fill=GREY_FILL, stroke=SOFT_GREY,
                           stroke_pt=1.0, radius_pt=6)
                continue
            # Просвет ячейки — строка минус её отступ: он меняется вместе со
            # ступенью сжатия, и зашитая под штатный отступ константа объявляла
            # переполнением каждую подогнанную ячейку.
            M.fits(plain(c), chosen, shares[j] - gap, rh[i] - row_pad + 0.04,
                   spacing=cell_spacing,
                   label=f"{label} строка {i} колонка {j}", mono="`" in c)
            text_box(sl, cx[j], ry, shares[j] - gap, rh[i], c, size=chosen, color=INK,
                     anchor=MSO_ANCHOR.MIDDLE, spacing=cell_spacing,
                     align=(align[j] if align else PP_ALIGN.LEFT))
        if i < len(rows) - 1:
            hairline(sl, x + pad, ry + rh[i], x + w - pad)
        ry += rh[i]
    return h


# ── Карточки-варианты ───────────────────────────────────────────────────────

def option_row(sl, x, y, w, options, *, highlight_idx=None, size=13.5, gap=0.16,
               min_h=1.15, max_h=1.7, label="карточки"):
    """Ряд карточек-вариантов. Выбранный вариант — пунктирная ЗОЛОТАЯ рамка и
    тёмно-золотой текст; остальные спокойные. На слайде-вопросе не подсвечен
    ни один: голосование идёт вслепую, разбор — на следующем слайде."""
    n = max(len(options), 1)
    cw = (w - gap * (n - 1)) / n
    inner = cw - 0.24
    fs = size
    # кегль уменьшается не только пока карточка выше отведённого, но и пока
    # самое длинное СЛОВО шире карточки: «описанием-триггером» не переносится
    # и в прежнем расчёте просто вылезало за край
    def _widest_word(size_):
        return max((M.text_w(wd, size_, bold=True)
                    for o in options for wd in plain(o).split()), default=0.0)
    while fs > 9.0 and (
            max(M.text_h(plain(o), fs, inner, spacing=1.16) for o in options) + 0.28 > max_h
            or _widest_word(fs) > inner):
        fs -= 0.5
    h = max(min_h, max(M.text_h(plain(o), fs, inner, spacing=1.16) for o in options) + 0.34)
    cx = x
    for i, opt in enumerate(options):
        hl = (highlight_idx is not None and i == highlight_idx)
        box = dashed_box if hl else ocean_box
        box(sl, cx, y, cw, h, fill=(GOLD_TINT if hl else SURFACE),
            stroke=(GOLD_DARK if hl else SOFT_GREY), stroke_pt=1.5 if hl else 1.2)
        M.fits(plain(opt), fs, inner, h - 0.2, label=f"{label} {i}", spacing=1.16)
        # `обратные кавычки` в исходнике помечают СТРОКУ как ряд карточек, а не
        # её содержимое как код: подпись варианта набирается обычным шрифтом
        text_box(sl, cx + 0.12, y, cw - 0.24, h, plain(opt), size=fs, bold=True, rich=False,
                 color=(GOLD_DARK if hl else INK), align=PP_ALIGN.CENTER,
                 anchor=MSO_ANCHOR.MIDDLE, spacing=1.16)
        cx += cw + gap
    return h


def numbered_list(sl, x, y, w, items, *, size=13.5, marker=TEAL, max_h=None,
                  numbered=False, label="список"):
    """Список в коробке. Маркер — тиловый кружок (или номер); высота строки
    считается по измеренному тексту, поэтому пункт в две строки не выдавливает
    последний пункт наружу (прежняя вёрстка считала `0.32 × число пунктов`)."""
    pad, mk = 0.24, 0.42
    inner = w - 2 * pad - mk
    fs = size
    if max_h:
        while fs > 9.5 and sum(M.text_h(plain(i), fs, inner) + 0.14 for i in items) + 2 * pad > max_h:
            fs -= 0.5
    ih = [M.text_h(plain(i), fs, inner) + 0.14 for i in items]
    h = sum(ih) + 2 * pad
    ocean_box(sl, x, y, w, h, fill=SURFACE, stroke=SOFT_GREY, stroke_pt=1.2)
    iy = y + pad
    for i, item in enumerate(items):
        if numbered:
            text_box(sl, x + pad, iy, mk - 0.06, ih[i], str(i + 1), size=fs + 2, bold=True,
                     color=marker, anchor=MSO_ANCHOR.MIDDLE, rich=False)
        else:
            rect(sl, x + pad + 0.04, iy + ih[i] / 2 - 0.055, 0.11, 0.11, marker,
                 radius=True, radius_adj=0.5)
        M.fits(plain(item), fs, inner, ih[i], label=f"{label} {i}")
        text_box(sl, x + pad + mk, iy, inner, ih[i], item, size=fs, color=INK,
                 anchor=MSO_ANCHOR.MIDDLE)
        iy += ih[i]
    return h


def criterion_plate(sl, x, y, w, title, items, *, size=11.5, accent=SLATE, max_h=None,
                    min_h=0.0, label="критерий"):
    """Плашка критерия: заголовок капителями + пункты через тонкие линии.
    Приём для слайдов «ещё рано / не нужно вообще» — там, где прежняя вёрстка
    ставила две одинаковые сетки друг над другом."""
    pad = 0.2
    inner = w - 2 * pad
    fs = size
    if max_h:
        while fs > 8.5 and 0.42 + sum(M.text_h(plain(i), fs, inner) + 0.16 for i in items) + pad > max_h:
            fs -= 0.5
    ih = [M.text_h(plain(i), fs, inner) + 0.16 for i in items]
    h = max(min_h, 0.42 + sum(ih) + pad)
    ocean_box(sl, x, y, w, h, fill=SURFACE, stroke=SOFT_GREY, stroke_pt=1.2)
    text_box(sl, x + pad, y + 0.12, inner, 0.3, title.upper(), size=11, bold=True,
             color=accent, rich=False)
    iy = y + 0.42
    for i, item in enumerate(items):
        M.fits(plain(item), fs, inner, ih[i], label=f"{label} {i}")
        text_box(sl, x + pad, iy, inner, ih[i], item, size=fs, color=INK,
                 anchor=MSO_ANCHOR.MIDDLE)
        if i < len(items) - 1:
            hairline(sl, x + pad, iy + ih[i], x + w - pad)
        iy += ih[i]
    return h


def slot_row(sl, x, y, w, slots, *, accent_idx=(), h=1.5, gap=0.16, label="слоты"):
    """Ряд слотов. Незаполненный слот рисуется ПУНКТИРОМ и читается как место,
    которое будет заполнено, — а не как пустая ячейка таблицы."""
    n = max(len(slots), 1)
    cw = (w - gap * (n - 1)) / n
    cx = x
    for i, (title, desc) in enumerate(slots):
        gold, empty = i in accent_idx, not desc
        box = dashed_box if empty else ocean_box
        box(sl, cx, y, cw, h, fill=(GOLD_TINT if gold else (GREY_FILL if empty else SURFACE)),
            stroke=(GOLD_DARK if gold else (SOFT_GREY if empty else LIGHT)),
            stroke_pt=1.5 if gold else 1.2)
        text_box(sl, cx + 0.08, y + 0.14, cw - 0.16, 0.5, title, size=12.5, bold=True,
                 color=(GOLD_DARK if gold else INK), align=PP_ALIGN.CENTER, spacing=1.08)
        if desc:
            M.fits(plain(desc), 9.8, cw - 0.2, h - 0.72, label=f"{label} {i}")
            text_box(sl, cx + 0.1, y + 0.66, cw - 0.2, h - 0.76, desc, size=9.8,
                     italic=True, color=SLATE, align=PP_ALIGN.CENTER, spacing=1.1)
        cx += cw + gap
    return h


def terminal_card(sl, x, y, w, lines, *, size=11.5, max_h=None, title=None, label="код"):
    """Тёмная моноширинная карточка для вывода команд и файлов. Высота — по
    числу РЕАЛЬНЫХ строк вывода; перенос моноширинного текста не делается,
    длинная строка отмечается предупреждением, а не молча обрезается."""
    pad = 0.24
    inner = w - 2 * pad
    fs = size
    if max_h:
        head = 0.32 if title else 0.0
        while fs > 7.5 and head + len(lines) * M.line_h(fs, 1.3) + 2 * pad > max_h:
            fs -= 0.5
    head = 0.32 if title else 0.0
    h = head + len(lines) * M.line_h(fs, 1.3) + 2 * pad
    if max_h:
        h = min(h, max_h)
    ocean_box(sl, x, y, w, h, fill=CODE_BG, stroke=LIGHT, stroke_pt=1.2, radius_pt=8)
    if title:
        text_box(sl, x + pad, y + 0.12, inner, 0.28, title, size=10.5, bold=True,
                 color=ON_DARK_MUTE, rich=False)
    for ln in lines:
        if M.text_w(ln, fs, mono=True) > inner:
            M._WARNINGS.append(f"ПЕРЕПОЛНЕНИЕ [{label}]: строка кода «{ln[:40]}…» шире карточки")
            break
    text_box(sl, x + pad, y + pad + head - 0.04, inner, h - 2 * pad - head + 0.08,
             lines, size=fs, color=CODE_FG, mono=True, spacing=1.3, rich=False)
    return h


def figure(sl, path, x, y, w, max_h, *, center=True):
    """Схема из `figures/`. Вписывается в отведённый прямоугольник целиком,
    пропорции сохраняются. 23 схемы сделаны отдельно и хорошо — задача
    вёрстки не переделать их, а дать им место."""
    from PIL import Image
    iw, ih = Image.open(path).size
    fw = w
    fh = w * ih / iw
    if fh > max_h:
        fh, fw = max_h, max_h * iw / ih
    fx = x + (w - fw) / 2 if center else x
    sl.shapes.add_picture(str(path), Inches(fx), Inches(y), width=Inches(fw), height=Inches(fh))
    return fh


# ── Дивайдеры ───────────────────────────────────────────────────────────────

def strip_pills(sl, x, y, w, n, active_idx, *, active=GOLD, inactive=PILL_OFF, pill_h=0.26,
                gap=0.12):
    """Полоса прогресса из сегментов. Ширина полосы и цвет сегмента — это и
    есть уровень дивайдера: широкая золотая = раздел, узкая тиловая = кейс."""
    pw = (w - gap * (n - 1)) / n
    cx = x
    for i in range(n):
        rect(sl, cx, y, pw, pill_h, active if i == active_idx else inactive,
             radius=True, radius_adj=0.4)
        cx += pw + gap
    return pill_h


def divider_bg(sl):
    """Фон дивайдера — ГРАДИЕНТ DEEP→MID→LIGHT, а не плоский тёмный. Плоская
    заливка делала все дивайдеры неотличимыми друг от друга."""
    gradient_rect(sl, 0, 0, W_IN, H_IN, [(0, DEEP), (55000, MID), (100000, LIGHT)])


def divider_badge(sl, number, *, cx=1.02, cy=1.62, r=0.38):
    """Золотой номерной значок — единственный элемент, который несёт ТОЛЬКО
    дивайдер уровня раздела."""
    circ = sl.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - r), Inches(cy - r),
                               Inches(2 * r), Inches(2 * r))
    circ.fill.solid()
    circ.fill.fore_color.rgb = GOLD
    circ.line.fill.background()
    _no_shadow(circ)
    text_box(sl, cx - r, cy - r, 2 * r, 2 * r, str(number), size=25, bold=True,
             color=DEEP, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, rich=False)


def tag_plate(sl, x, y, text, *, w=6.2, stroke=GOLD, color=GOLD, size=11.5):
    """Плашка-ярлык дивайдера («1 кейс · 2 слоя провала · 5 форм обхода»).

    Рисуется ОТ ПЕРЕДАННОГО y. Прежняя вёрстка печатала этот ярлык жёстко в
    `Inches(3.0)` — в той же точке, где уже стоял абзац или коробка цитаты, —
    и два текстовых блока ложились друг на друга буква в букву на s07 и s21.
    Это и есть «текст наложенный сам на себя», который владелец увидел за
    минуту. Здесь координата приходит снаружи, от курсора вёрстки."""
    h = 0.55
    rect(sl, x, y, w, h, DEEP_INK, stroke=stroke, stroke_pt=1.2, radius=True, radius_adj=0.28)
    text_box(sl, x + 0.26, y, w - 0.52, h, text.upper(), size=size, bold=True, color=color,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, rich=False)
    return h


def roadmap(sl, stages, cur_idx, *, x=0.55, y=6.55, w=12.23):
    """Дорожная карта занятия внизу дивайдера: где мы среди ступеней."""
    n = len(stages)
    gap = 0.2
    cw = (w - gap * (n - 1)) / n
    cx = x
    for i, name in enumerate(stages):
        on = (i == cur_idx)
        done = (i < cur_idx)
        rect(sl, cx, y, cw, 0.42, GOLD if on else (TEAL if done else PILL_OFF),
             radius=True, radius_adj=0.3)
        text_box(sl, cx, y, cw, 0.42, name, size=12.5, bold=on,
                 color=DEEP if on else (WHITE if done else ON_DARK_MUTE),
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, rich=False)
        cx += cw + gap
    return 0.42


# ── Грамматика деки: одна работа слайда — одна форма ────────────────────────
#
# Прожарка измерила: одна и та же кремовая плашка во всю ширину стояла на 31
# слайде из 56 и несла ШЕСТЬ разных работ — вопрос залу, реплику докладчика,
# тезис-итог, оговорку о пробеле, факт и технический блок. Если два разных по
# смыслу слайда выглядят одинаково, зал не различит их и в зале.
#
# Ниже — шесть форм, по одной на работу. Ни одна не переиспользуется под
# другую работу, и каждая отличается не оттенком, а УСТРОЙСТВОМ: заливка,
# рамка, планка слева, курсив, пунктир, моноширинный шрифт.
#
#   вопрос залу   → золотая заливка + золотая рамка, по центру   (form_question)
#   формула/итог  → толстая золотая планка слева, БЕЗ заливки    (form_formula)
#   реплика «…»   → тонкая тиловая линия слева, курсив, тихо     (form_speech)
#   оговорка      → пунктирная серая рамка                       (form_caveat)
#   факт / число  → тиловая заливка + тиловая планка             (form_fact)
#   технический   → тёмная моноширинная карточка                 (terminal_card)
#
# Золотая ЗАЛИВКА С РАМКОЙ существует ровно в одном месте — на вопросе.


def form_question(sl, x, y, w, lines, *, size=16.5, max_h=None, label="вопрос"):
    """Вопрос залу. Единственная форма с золотой заливкой И золотой рамкой —
    ни одна другая работа слайда её не получает."""
    return gold_callout(sl, x, y, w, lines, size=size, bold=True, align=PP_ALIGN.CENTER,
                        max_h=max_h, min_h=0.9, label=label)


def form_formula(sl, x, y, w, lines, *, size=15, max_h=None, label="формула"):
    """Формула или тезис-итог — то, что надо запомнить. Толстая золотая планка
    слева и жирный текст без заливки: вес есть, а золотой коробки нет, поэтому
    с вопросом не спутать."""
    bar, pad = 0.09, 0.3
    inner = w - bar - pad - 0.2
    if max_h:
        while size > 11 and M.rich_block_h(lines, size, inner,
                                    spacing=1.24, space_after=3) + 0.3 > max_h:
            size -= 0.5
    h = M.rich_block_h(lines, size, inner, spacing=1.24, space_after=3) + 0.3
    if max_h:
        h = min(h, max_h)
    rect(sl, x, y, bar, h, GOLD, radius=True, radius_adj=0.5)
    M.fits(" ".join(plain(l) for l in lines), size, inner, h - 0.26, label=label, spacing=1.24)
    text_box(sl, x + bar + pad, y + 0.14, inner, h - 0.28, lines, size=size, bold=True,
             color=INK, anchor=MSO_ANCHOR.MIDDLE, spacing=1.24, space_after=3)
    return h


def form_speech(sl, x, y, w, lines, *, size=13, max_h=None, label="реплика"):
    """Реплика докладчика в кавычках. Тихая форма: тонкая тиловая линия слева,
    курсив, приглушённый цвет, никакой заливки. Студент видит, что это слова
    преподавателя, а не тезис слайда."""
    bar, pad = 0.035, 0.26
    inner = w - bar - pad - 0.2
    if max_h:
        while size > 10 and M.rich_block_h(lines, size, inner,
                                    spacing=1.26, space_after=3) + 0.2 > max_h:
            size -= 0.5
    h = M.rich_block_h(lines, size, inner, spacing=1.26, space_after=3) + 0.2
    if max_h:
        h = min(h, max_h)
    rect(sl, x, y + 0.04, bar, h - 0.08, TEAL)
    M.fits(" ".join(plain(l) for l in lines), size, inner, h - 0.18, label=label, spacing=1.26)
    text_box(sl, x + bar + pad, y + 0.1, inner, h - 0.2, lines, size=size, italic=True,
             color=SLATE, anchor=MSO_ANCHOR.MIDDLE, spacing=1.26, space_after=3)
    return h


def form_caveat(sl, x, y, w, lines, *, size=12.5, max_h=None, label="оговорка"):
    """Честный пробел, оговорка, «этого измерить не удалось». Пунктирная серая
    рамка — та же пунктирная грамматика, что у пустого слота: «здесь чего-то
    нет, и мы об этом говорим вслух»."""
    pad = 0.26
    inner = w - 2 * pad
    if max_h:
        while size > 9.5 and M.rich_block_h(lines, size, inner, space_after=3) + 0.3 > max_h:
            size -= 0.5
    h = M.rich_block_h(lines, size, inner, space_after=3) + 0.3
    if max_h:
        h = min(h, max_h)
    dashed_box(sl, x, y, w, h, fill=GREY_FILL, stroke=SLATE, stroke_pt=1.2)
    M.fits(" ".join(plain(l) for l in lines), size, inner, h - 0.26, label=label)
    text_box(sl, x + pad, y + 0.15, inner, h - 0.3, lines, size=size, italic=True,
             color=SLATE, anchor=MSO_ANCHOR.MIDDLE, space_after=3)
    return h


def form_fact(sl, x, y, w, lines, *, size=13.5, max_h=None, label="факт"):
    """Факт, число, наблюдение. Тиловая заливка и тиловая планка — нейтральная
    отметка, которая не претендует ни на вопрос, ни на вывод.

    Заодно здесь наконец работает TEAL: вторичный цвет палитры был
    импортирован прежним рендерером и не использован ни разу."""
    bar, pad = 0.07, 0.26
    inner = w - bar - 2 * pad
    if max_h:
        while size > 10 and M.rich_block_h(lines, size, inner, space_after=3) + 0.3 > max_h:
            size -= 0.5
    h = M.rich_block_h(lines, size, inner, space_after=3) + 0.3
    if max_h:
        h = min(h, max_h)
    ocean_box(sl, x, y, w, h, fill=TEAL_TINT, stroke=TEAL, stroke_pt=1.1)
    rect(sl, x + 0.02, y + 0.06, bar, h - 0.12, TEAL, radius=True, radius_adj=0.5)
    M.fits(" ".join(plain(l) for l in lines), size, inner, h - 0.26, label=label)
    text_box(sl, x + bar + pad, y + 0.15, inner, h - 0.3, lines, size=size, color=INK,
             anchor=MSO_ANCHOR.MIDDLE, space_after=3)
    return h


PILL_PAD, PILL_GAP, PILL_H = 0.22, 0.14, 0.44


def pill_rows(items, w, size):
    """Упаковка терминов в строки заданной ширины — отдельно от отрисовки,
    потому что высоту ряда пилюль надо знать ДО того, как он нарисован:
    дорожки Такта Б считают свою высоту вместе с терминами под базой."""
    rows, cur, cw = [], [], 0.0
    for it in items:
        iw = min(M.text_w(plain(it), size, bold=True) + 2 * PILL_PAD, w)
        if cur and cw + iw + PILL_GAP > w:
            rows.append(cur); cur, cw = [], 0.0
        cur.append((it, iw)); cw += iw + PILL_GAP
    if cur:
        rows.append(cur)
    return rows


def pill_rows_h(items, w, size):
    n = len(pill_rows(items, w, size))
    return n * (PILL_H + PILL_GAP) - PILL_GAP if n else 0.0


def term_pills(sl, x, y, w, items, *, size=13, label="термины", center=True):
    """Ряд терминов — работа «ЗАПОМНИ». Компактные тиловые пилюли по ширине
    текста, в одну-две строки.

    Отличается от `option_row` («ВЫБЕРИ») устройством, а не оттенком: там
    карточки равной ширины во всю строку, здесь пилюли по размеру слова. Ряд
    карточек прежде означал «выбери» на четырёх слайдах и «запомни» на двух —
    одна форма на две разные работы.

    `center=False` — для узкой колонки (дорожка «база» Такта Б): ряд из двух
    пилюль, поставленный по центру полуширины, читается как случайно съехавший
    от края текста, под которым он стоит."""
    pad, gap, ph = PILL_PAD, PILL_GAP, PILL_H
    rows = pill_rows(items, w, size)
    ry = y
    for row in rows:
        total = sum(iw for _, iw in row) + gap * (len(row) - 1)
        cx = x + ((w - total) / 2 if center else 0.0)
        for it, iw in row:
            ocean_box(sl, cx, ry, iw, ph, fill=TEAL_TINT, stroke=TEAL, stroke_pt=1.2, radius_pt=20)
            text_box(sl, cx, ry, iw, ph, plain(it), size=size, bold=True, color=MID,
                     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, rich=False)
            cx += iw + gap
        ry += ph + 0.14
    return ry - y - 0.14


# ── Такт Б: две дорожки на одном экране ─────────────────────────────────────
#
# Седьмая форма грамматики — и ЕДИНСТВЕННАЯ БЕЗ КОРОБКИ. Шесть прежних форм
# все до одной суть плашка, коробка или карточка; здесь рабочее поле делится
# вертикальным ШВОМ на две равные дорожки, каждая под своей горизонтальной
# планкой сверху. Спутать не с чем:
#
#   вопрос      → золотая ЗАЛИВКА с золотой РАМКОЙ      — здесь заливки нет
#   формула     → толстая золотая планка СЛЕВА           — здесь планка СВЕРХУ
#   факт        → тиловая заливка и тиловая планка       — здесь заливки нет
#   оговорка    → пунктирная РАМКА                       — здесь рамки нет
#   технический → тёмная моноширинная КАРТОЧКА           — здесь фона нет
#   критерий    → скруглённая коробка, кегль 11,5        — здесь коробки нет,
#                 (и она тоже бывает парой бок о бок)      кегль 15,5–16,5,
#                                                          а между дорожками
#                                                          шов, а не зазор
#
# Устройство читается за полсекунды: два равных поля, тонкий шов между ними,
# золотой тик на шве. Планка над левой дорожкой — MID (спокойная), над правой
# — GOLD: кромка и есть то, за чем на этом слайде приходит сильный.

TRACK_SIZES = (18, 17, 16, 15, 14, 13, 12)
TRACK_HEAD = 0.56          # планка + ярлык капителями над дорожкой
TRACK_SPACING = 1.30       # межстрочный: 45 секунд — значит воздух, не плотность


def _track_h(lines, size, w, terms=None, term_size=12.5):
    h = M.rich_block_h(lines, size, w, spacing=TRACK_SPACING, space_after=6)
    if terms:
        h += 0.26 + pill_rows_h(terms, w, term_size)
    return h


def base_edge_tracks(sl, x, y, w, left_lines, right_lines, *, left_label="база",
                     right_label="кромка", terms=(), size=None, max_h=None,
                     seam=0.46, label="дорожки"):
    """Такт Б: база слева, кромка справа, на одном экране и в равных правах.

    Равноправие здесь — не пожелание, а конструкция: ширины дорожек равны с
    точностью до шва, кегль у них ОДИН И ТОТ ЖЕ, обе начинаются с одного y, и
    шов проходит на всю высоту более высокой из двух. Поэтому ни одна из них
    не может прочитаться как сноска к другой. Различает их не вес, а цвет
    планки сверху: MID у базы, GOLD у кромки.

    Термины (`terms`) встают пилюлями ПОД базой, а не отдельным блоком: они
    принадлежат левой дорожке — это ровно то, что «дальше звучит без
    объяснения». Кегль у них на ступень мельче текста дорожки, потому что
    пилюля — ярлык, а не фраза.
    """
    col = (w - seam) / 2
    term_size = 12.5

    ladder = TRACK_SIZES if size is None else (size,)
    chosen = ladder[-1]
    for s in ladder:
        chosen = s
        need = TRACK_HEAD + max(_track_h(left_lines, s, col, terms, term_size),
                                _track_h(right_lines, s, col))
        if max_h is None or need <= max_h:
            break

    hl = _track_h(left_lines, chosen, col, terms, term_size)
    hr = _track_h(right_lines, chosen, col)
    h = TRACK_HEAD + max(hl, hr)
    if max_h is not None and h > max_h + 0.02:
        M._WARNINGS.append(f"ПЕРЕПОЛНЕНИЕ [{label}]: дорожки {h:.2f}″ при бюджете "
                           f"{max_h:.2f}″ — текста на 45 секунд слишком много")

    xr = x + col + seam
    # планки сверху: равной длины и толщины, разного цвета
    rect(sl, x, y, col, 0.05, MID)
    rect(sl, xr, y, col, 0.05, GOLD)
    text_box(sl, x, y + 0.14, col, 0.28, left_label.upper(), size=11.5, bold=True,
             color=MID, rich=False)
    text_box(sl, xr, y + 0.14, col, 0.28, right_label.upper(), size=11.5, bold=True,
             color=GOLD_DARK, rich=False)

    # шов: одна тонкая линия на всю высоту поля и золотой тик на её вершине.
    # Именно шов, а не зазор: зазор у нас уже означает «две отдельные плашки»
    # (слайд критериев), а шов — «одно поле, поделённое на два».
    sx = x + col + seam / 2
    rect(sl, sx - 0.004, y, 0.008, h, SOFT_GREY)
    rect(sl, sx - 0.05, y, 0.10, 0.10, GOLD)

    by = y + TRACK_HEAD
    M.fits(" ".join(plain(l) for l in left_lines), chosen, col, hl,
           label=f"{label} база", spacing=TRACK_SPACING)
    M.fits(" ".join(plain(l) for l in right_lines), chosen, col, hr,
           label=f"{label} кромка", spacing=TRACK_SPACING)
    text_box(sl, x, by, col, hl, left_lines, size=chosen, color=INK,
             spacing=TRACK_SPACING, space_after=6)
    text_box(sl, xr, by, col, hr, right_lines, size=chosen, color=INK,
             spacing=TRACK_SPACING, space_after=6)
    if terms:
        ty = by + M.rich_block_h(left_lines, chosen, col,
                                 spacing=TRACK_SPACING, space_after=6) + 0.26
        term_pills(sl, x, ty, col, list(terms), size=term_size, center=False,
                   label=f"{label} термины")
    return h
