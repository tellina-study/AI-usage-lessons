# -*- coding: utf-8 -*-
"""КАДРЫ: нормализация файла, боковая колонка, лента, обязательный hero.

Собрано из обоих сборщиков, и здесь они дополняют друг друга почти без
пересечений — у Лекции 6 была колонка на два кадра, лента и подпись под
кадром; у Лекции 7 — кадр-свидетельство на 40% площади (§5.9) и коробка,
ужимающаяся до самого снимка.

ПОЧЕМУ НОРМАЛИЗАЦИЯ ОБЯЗАТЕЛЬНА (lec-06, `notes/mcp-limitations.md`
[#223-9]). `python-pptx` не знает формата MPO — мультикадрового JPEG из
камеры, — и падает на нём ValueError'ом. В Лекции 6 такой оказалась ровно
одна картинка из 65: падение случайное и находится только на сборке. Там же
ужимается длинная сторона до 2200 px — полноразмерные снимки давали деку на
97 МБ при той же картинке на экране. Исходник НЕ подменяется: производная
кладётся рядом, в отдельный каталог.
"""
from pathlib import Path

from PIL import Image
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches

from . import metrics, prims, report
from .markup import plain
from .theme import LIGHT, OUTLINE, SLATE, SURFACE, WHITE, SLIDE_H, SLIDE_W

OKFMT = {"JPEG", "PNG", "GIF", "BMP", "TIFF"}
MAXSIDE = 2200
_NORM = {}

MIN_IMG_W = 2.20        # уже этого кадр в колонке не читается
CAP_MAX_ROWS = 2        # подпись под кадром — строка-две, не абзац
CAP_GAP = 0.10          # просвет между нижним краем кадра и подписью
CAP_PT = 12.0        # пол видимого слоя, см. `floor.py` (было 11,5)
RIBBON_AR = 3.0         # шире этого — лента: поперёк кадра под текстом
SPAN_MAX_H = 2.90       # предел высоты полосы кадров
PAD = 0.18

HERO_MIN_AREA = 0.40    # §5.9: обязательный hero — не меньше 40% площади
HERO_PAD = 0.05


def norm_image(path, derived_dir):
    """Производная кадра: формат из белого списка `python-pptx` и длинная
    сторона не больше `MAXSIDE`. Исходник не подменяется."""
    path = Path(path)
    key = (path, Path(derived_dir))
    if key in _NORM:
        return _NORM[key]
    out = path
    if path.exists():
        try:
            with Image.open(path) as im:
                fmt, (iw, ih) = im.format, im.size
                if fmt not in OKFMT or max(iw, ih) > MAXSIDE:
                    jpg = fmt == "JPEG"
                    out = Path(derived_dir) / (path.stem
                                               + (".jpg" if jpg else ".png"))
                    if (not out.exists()
                            or out.stat().st_mtime < path.stat().st_mtime):
                        out.parent.mkdir(parents=True, exist_ok=True)
                        k = min(1.0, MAXSIDE / max(iw, ih))
                        im2 = im
                        if im2.mode not in ("RGB", "L"):
                            im2 = im2.convert("RGB")
                        if k < 1.0:
                            im2 = im2.resize((max(1, int(iw * k)),
                                              max(1, int(ih * k))),
                                             Image.LANCZOS)
                        im2.save(out, "JPEG" if jpg else "PNG",
                                 **({"quality": 86} if jpg else {}))
        except Exception as e:                       # noqa: BLE001
            report.warn(f"[кадр] {path.name}: {e}")
            out = path
    _NORM[key] = out
    return out


def img_ar(path, sid=""):
    """Отношение ШИРИНЫ к высоте. Нужно ДО раскладки: ширину боковой колонки
    назначает не слот, а сам снимок."""
    try:
        with Image.open(path) as im:
            return im.width / max(1, im.height)
    except Exception as e:                           # noqa: BLE001
        report.warn(f"[кадр] {sid}: {Path(path).name} — {e}")
        return None


def img_ratio(path):
    """Отношение высоты к ширине — та же величина наоборот, но так её
    считает код ленты и колонки, перенесённый из lec-06."""
    ar = img_ar(path)
    return 0.75 if not ar else 1.0 / ar


def media_path(meta, root, sid=""):
    """Поле `media` → путь на диске или None. Файла нет — `[кадр]` в отчёт,
    а кадр собирается как кадр без снимка."""
    m = str(meta.get("media") or "").strip()
    if not m or m in ("none", "None"):
        return None
    p = Path(m) if Path(m).is_absolute() else Path(root) / m
    if not p.exists():
        report.warn(f"[кадр] {sid}: файла нет — {m}")
        return None
    return p


def frames(meta, root, sid=""):
    """Кадры кадра, которые ДЕЙСТВИТЕЛЬНО есть на диске: `media` и, если
    задан, `media_2`."""
    first = media_path(meta, root, sid)
    if first is None:
        return []
    out = [first]
    second = str(meta.get("media_2") or "").strip()
    if second:
        p = Path(root) / second
        if p.exists():
            out.append(p)
        else:
            report.MISSING_MEDIA.append(f"{sid}: нет файла {p}")
    return out


def is_ribbon(meta, fr):
    """Лента: заявлено `media_span` ИЛИ сам снимок шире 3:1.

    В боковой колонке 4,45" лента 1920×192 даёт полоску в треть дюйма,
    висящую посреди пустой высокой рамки. Поле `media_span` читается, но на
    него больше ничего не держится: пропорция самого файла решает сама."""
    if not fr:
        return False
    if str(meta.get("media_span") or "").strip():
        return True
    return any(img_ratio(p) <= 1.0 / RIBBON_AR for p in fr)


# ── боковая колонка и лента (lec-06) ────────────────────────────────


def media_box(s, x, y, w, h, image, cap, tag, *, pad=PAD,
              max_cap=CAP_MAX_ROWS):
    """Кадр в рамке-мотиве, который НИКОГДА не вылезает за отведённое место.

    Прежняя версия подгоняла рамку под пропорции снимка сходящимся счётом, и
    на узкой колонке с длинной подписью счёт расходился: на `s14` (два кадра
    в колонке, подпись в две работы) нижняя рамка уехала за край слайда и
    легла поверх подвала. Поэтому здесь жёсткий зажим по высоте, а подпись,
    которой нужно больше двух строк, уходит в ПОДВАЛ целиком вместо того,
    чтобы рвать раскладку.

    Возвращает подпись, которая НЕ поместилась (или None)."""
    cap = plain(cap) if cap else None
    try:
        with Image.open(image) as im:
            iw, ih = im.size
        ratio = ih / iw
    except Exception:                                # noqa: BLE001
        report.MISSING_MEDIA.append(f"{tag}: не читается {Path(image).name}")
        ratio = 0.75
    spill = None
    cap_h = (0.30 + CAP_GAP) if cap else 0.0
    iw_in = w - 2 * pad
    ih_in = 0.0
    for _ in range(4):
        inner = h - 2 * pad - cap_h
        iw_in = max(0.5, min(w - 2 * pad, inner / max(ratio, 0.01)))
        ih_in = iw_in * ratio
        if not cap:
            break
        rows = metrics.n_lines(cap, CAP_PT, iw_in)
        if rows > max_cap:                           # подпись в подвал
            spill, cap, cap_h = cap, None, 0.0
            continue
        new = 0.30 * rows + 0.06 + CAP_GAP
        if abs(new - cap_h) < 0.02:
            break
        cap_h = new
    ph = ih_in + 2 * pad + cap_h
    if ph > h:                                       # зажим: место главнее
        k = max(0.0, h - 2 * pad - cap_h) / max(ih_in, 0.01)
        iw_in, ih_in, ph = iw_in * k, ih_in * k, h
    # Подпись не вправе стоить кадру читаемый размер: сначала кадр, подпись
    # уходит в подвал на всю ширину слайда.
    if cap and iw_in < MIN_IMG_W:
        return media_box(s, x, y, w, h, image, None, tag, pad=pad,
                         max_cap=max_cap) or cap
    if iw_in < MIN_IMG_W:
        report.warn(f"[кадр] {tag}: ширина {iw_in:.2f}\" "
                    f"(< {MIN_IMG_W:.2f}\") — тесное место")
    pw = iw_in + 2 * pad
    px, py = x + (w - pw) / 2, y + (h - ph) / 2
    prims.ocean_box(s, px, py, pw, ph, fill=WHITE)
    prims.add_image(s, image, px + pad, py + pad, iw_in, ih_in, tag=tag)
    if cap:
        prims.text_box(s, px + pad, py + ph - pad - cap_h + 0.02, iw_in,
                       cap_h, cap, size=CAP_PT, italic=True, color=LIGHT,
                       align=PP_ALIGN.CENTER, ls=1.12, tag=f"{tag}/подпись")
    return spill


def media_column(s, fr, caps, x, y, w, h, tag):
    """Колонка кадра: один снимок или два друг под другом.

    Высота двух делится НЕ пополам, а по их собственным пропорциям. При
    ровном делении портрет схлопывается в миниатюру: на `s47a` второй
    снимок (661×935) выходил 1,04" шириной в колонке 4,45"."""
    if len(fr) < 2:
        a = media_box(s, x, y, w, h, fr[0], caps[0] if caps else None,
                      f"{tag}/кадр")
        return ([a] if a else []) + [plain(z) for z in caps[1:]]
    gap = PAD
    avail = h - gap
    nat = [(w - 2 * PAD) * img_ratio(p) + 2 * PAD for p in fr]
    k = min(1.0, avail / max(sum(nat), 0.01))
    hs = [n * k for n in nat]
    slack = (avail - sum(hs)) / 2
    hs = [n + slack for n in hs]
    a = media_box(s, x, y, w, hs[0], fr[0], caps[0] if caps else None,
                  f"{tag}/кадр1")
    c = media_box(s, x, y + hs[0] + gap, w, hs[1], fr[1],
                  caps[1] if len(caps) > 1 else None, f"{tag}/кадр2")
    used = min(2, len(caps))
    return [z for z in (a, c) if z] + [plain(z) for z in caps[used:]]


def band_h(fr, caps, full_w, col_gap):
    """Высота полосы кадров под текстом — по пропорциям самих снимков, но не
    выше `SPAN_MAX_H`."""
    cw = full_w if len(fr) < 2 else (full_w - col_gap) / 2
    cap_h = 0.36 if caps else 0.0
    h = max((cw - 2 * PAD) * img_ratio(p) + 2 * PAD + cap_h for p in fr)
    return min(max(h, 0.80), SPAN_MAX_H)


def media_band(s, fr, caps, x, y, h, full_w, col_gap, tag):
    """Полоса кадров поперёк слайда под текстом: один во всю ширину или два
    рядом по половине. Два стоят в своих половинах и на подпись друг друга
    не наезжают — каждый зажат своей рамкой."""
    if len(fr) < 2:
        a = media_box(s, x, y, full_w, h, fr[0], caps[0] if caps else None,
                      f"{tag}/лента")
        return ([a] if a else []) + [plain(z) for z in caps[1:]]
    cw = (full_w - col_gap) / 2
    a = media_box(s, x, y, cw, h, fr[0], caps[0] if caps else None,
                  f"{tag}/лента1")
    c = media_box(s, x + cw + col_gap, y, cw, h, fr[1],
                  caps[1] if len(caps) > 1 else None, f"{tag}/лента2")
    used = min(2, len(caps))
    return [z for z in (a, c) if z] + [plain(z) for z in caps[used:]]


# ── боковая коробка по СНИМКУ, а не по слоту (lec-07) ───────────────


def media_box_size(ar, w_slot, h_slot, pad=0.3):
    """Коробка боковой иллюстрации по СНИМКУ, а не по слоту.

    Слот — место, которое кадр готов отдать; коробка — то, что снимок
    действительно занимает. Разница видна сразу: вертикальный снимок,
    вписанный в слот 4,30×3,40", берёт от ширины полтора дюйма, а остальные
    два с половиной держит пустыми — и ровно на эту ширину сужается текст
    рядом. Горизонтальный поступает так же с высотой.

    Возвращает (ширина коробки, высота коробки)."""
    if not ar:
        return w_slot, h_slot
    iw, ih = w_slot - pad, h_slot - pad
    if iw / ih > ar:
        iw = ih * ar
    else:
        ih = iw / ar
    return iw + pad, ih + pad


def draw_media(s, sid, p, x, y, w, h, pad=0.3):
    """Боковая иллюстрация в коробке мотива, вписанная по центру."""
    prims.ocean_box(s, x, y, w, h)
    ar = img_ar(p, sid)
    if not ar:
        return
    iw, ih = w - pad, h - pad
    if iw / ih > ar:
        iw = ih * ar
    else:
        ih = iw / ar
    s.shapes.add_picture(str(p), Inches(x + (w - iw) / 2),
                         Inches(y + (h - ih) / 2), Inches(iw), Inches(ih))
    report.PRIM["кадр"] += 1


# ── обязательный кадр-свидетельство, §5.9 (lec-07) ──────────────────


def draw_hero(s, sid, meta, p, x_right, y_top, y_bot, w_max):
    """Крупный кадр-свидетельство с подписью об источнике.

    Отличается от `draw_media` не размером, а назначением: здесь картинка И
    ЕСТЬ утверждение кадра, поэтому правило курса (§5.9) требует ей не
    меньше 40% площади слайда и ВИДИМОЙ подписи об источнике. Подпись тут не
    вежливость: неподписанный снимок записи в базе регулятора неотличим от
    карточки, нарисованной нами самими, — ровно та подмена, которую курс
    запрещает отдельным правилом.

    Саму подпись рисует не эта функция, а подвал деки: место источника одно,
    и поле `media_caption` уезжает туда последним сегментом. Отдельная
    строка под картинкой обошлась бы в 0,27" высоты, а их на кадре с
    40-процентной иллюстрацией нет — замерено.

    Возвращает (левый край коробки, доля площади слайда, низ коробки)."""
    ar = img_ar(p, sid)
    if not ar:
        return x_right, 0.0, y_top
    cap = str(meta.get("media_caption") or "").strip()
    if not cap:
        report.warn(
            f"[кадр] {sid}: обязательный hero без поля `media_caption` — "
            "§5.9 требует видимой подписи об источнике")
    w = w_max
    while w > 1.0 and 2 * HERO_PAD + w / ar > y_bot - y_top:
        w -= 0.02
    h = w / ar
    box_w, box_h = w + 2 * HERO_PAD, h + 2 * HERO_PAD
    x_box = x_right - box_w
    prims.ocean_box(s, x_box, y_top, box_w, box_h)
    s.shapes.add_picture(str(p), Inches(x_box + HERO_PAD),
                         Inches(y_top + HERO_PAD), Inches(w), Inches(h))
    report.PRIM["кадр"] += 1
    return x_box, w * h / (SLIDE_W * SLIDE_H), y_top + box_h


def place_hero(s, sid, meta, p, x_right, y_top, y_bot, w_max):
    """Поставить обязательный hero и проверить порог §5.9 ОДНИМ ходом.

    Отдельной функцией, потому что мест вызова ДВА — содержательный кадр и
    обложка, — а порог и текст предупреждения обязаны быть одни. Пока
    проверка стояла внутри раскладки содержательного кадра, обложка молча
    прошла мимо неё: ровно так `s01` Лекции 7 и оказался без hero при том,
    что снимок лежал в `assets/` с проверенными правами.

    Возвращает (левый край коробки, низ коробки)."""
    x_box, area, bot = draw_hero(s, sid, meta, p, x_right, y_top, y_bot,
                                 w_max)
    report.HERO_AREA[sid] = area
    if area < HERO_MIN_AREA:
        report.warn(
            f"[кадр] {sid}: обязательный hero занял {100 * area:.1f}% "
            f"площади кадра при пороге {100 * HERO_MIN_AREA:.0f}% (§5.9)")
    return x_box, bot
