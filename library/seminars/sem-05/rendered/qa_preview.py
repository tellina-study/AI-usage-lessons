#!/usr/bin/env python3
"""Предпросмотр отрендеренных слайдов и проверка, что текст помещается в отведённые рамки.

Конвертера pptx→pdf в окружении нет, поэтому картинка собирается из реальной геометрии
готового .pptx: позиции, размеры, заливки, шрифты и текст читаются из файла, а не из исходников.
Строки переносятся по той же ширине, что задана фигуре, и переполнение печатается списком.

ПОПРАВКА НА ШРИФТ РАБОТАЕТ В ДВЕ СТОРОНЫ — и перепутать их значит либо врать
картинкой, либо врать списком переполнений. Дека задаёт Arial, в окружении есть
только DejaVu, который для кириллицы шире примерно на 5–10%.

  * РИСУЕМ по мерке DejaVu без поправки (`K_LAYOUT`, ровно 1.0). Картинка
    показывает худший случай переноса — если на ней текст стоит в рамке, в
    настоящем Arial он стоит с запасом. Подкручивать картинку в «как будет в
    Arial» нельзя: смотреть глазами тогда не на что.
  * ОБЪЯВЛЯЕМ переполнение по мерке `K_REPORT` (0.88), то есть считая текст
    УЗКИМ. Иначе каждая вторая рамка деки попадала бы в список как дефект,
    которого в настоящем Arial нет, — и список перестал бы читаться.

Обе величины берутся из `metrics.py`, а не задаются здесь вторым экземпляром:
разойтись сборка и предпросмотр в оценке «влезло или нет» не должны.
"""
import sys
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu

import metrics as M

SC = 120 / 914400          # пикселей на EMU при 120 px/дюйм
FR = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
FB = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
FM = "/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf"
OUT = Path(__file__).parent / "preview"; OUT.mkdir(exist_ok=True)

def font(pt, bold=False, mono=False):
    return ImageFont.truetype(FM if mono else (FB if bold else FR), max(7, int(pt * 120 / 72)))

def wrap(d, text, fo, w, mono=False):
    out = []
    for para in text.split("\n"):
        if mono:
            out.append(para); continue
        line = ""
        for word in para.split():
            probe = (line + " " + word).strip()
            if d.textlength(probe, font=fo) > w and line:
                out.append(line); line = word
            else:
                line = probe
        out.append(line)
    return out

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def grad_stops(sh):
    """Остановки градиентной заливки, если она есть.

    Без этого предпросмотр рисовал дивайдеры БЕЛЫМИ: градиент задаётся узлом
    `gradFill`, а не `solidFill`, проверка `fill.type == 1` его не видела — и
    белый текст на «белом» фоне пропадал вовсе. Смотреть на такую картинку и
    делать вывод о вёрстке нельзя, поэтому градиент разбирается явно."""
    try: sppr = sh._element.spPr
    except Exception: return None
    g = sppr.find(A + "gradFill")
    if g is None: return None
    out = []
    for gs in g.iter(A + "gs"):
        clr = gs.find(A + "srgbClr")
        if clr is None: continue
        v = clr.get("val")
        out.append((int(gs.get("pos", 0)) / 100000.0,
                    (int(v[0:2], 16), int(v[2:4], 16), int(v[4:6], 16))))
    return sorted(out) or None


def paint_gradient(im, x, y, w, h, stops):
    """Линейный градиент по диагонали — тот же угол 45°, что задаёт вёрстка."""
    if w <= 0 or h <= 0: return
    box = Image.new("RGB", (max(w, 1), max(h, 1)))
    px = box.load()
    for j in range(box.size[1]):
        for i in range(0, box.size[0], 4):
            t = (i / max(box.size[0] - 1, 1) + j / max(box.size[1] - 1, 1)) / 2
            lo = stops[0]; hi = stops[-1]
            for k in range(len(stops) - 1):
                if stops[k][0] <= t <= stops[k + 1][0]:
                    lo, hi = stops[k], stops[k + 1]; break
            span = max(hi[0] - lo[0], 1e-6)
            f = min(max((t - lo[0]) / span, 0.0), 1.0)
            col = tuple(int(lo[1][c] + (hi[1][c] - lo[1][c]) * f) for c in range(3))
            for d in range(4):
                if i + d < box.size[0]: px[i + d, j] = col
    im.paste(box, (x, y))


def _is_mono(r):
    nm = (r.font.name or "").lower()
    return nm.startswith("consol") or nm.startswith("dejavu sans mono")


def _margin(tf, attr):
    """Поле текстовой рамки в пикселях предпросмотра. У ячейки таблицы оно
    своё (python-pptx ставит 0,1″ по бокам), у наших фигур — ноль."""
    try:
        v = getattr(tf, attr)
        return int(v * SC) if v is not None else 0
    except Exception:
        return 0


def run_info(p, r=None):
    """Кегль, шрифт, насыщенность и цвет — для абзаца или для ОДНОГО прогона.

    Прогон нужен отдельно, потому что `**жирный**` внутри строки теперь
    что-то значит (целевой ответ в таблице, острая часть кромки на Такте Б), а
    прежде предпросмотр красил весь абзац по нулевому прогону — и жирного на
    картинке не было видно вовсе. Смотреть глазами на разметку, которой не
    видно, нельзя."""
    if r is None:
        r = p.runs[0] if p.runs else None
    sz = (r.font.size if r else None) or p.font.size
    nm = ((r.font.name if r else None) or p.font.name) or "Arial"
    bo = bool((r.font.bold if r else None) or p.font.bold)
    col = None
    for src in (r.font if r else None, p.font):
        try:
            if src and src.color and src.color.rgb: col = tuple(src.color.rgb); break
        except Exception: pass
    return (sz.pt if sz else 18), nm, bo, col or (20, 27, 46)


import re as _re
_WS = _re.compile(r"(\s+)")


def wrap_runs(d, runs, w):
    """Перенос строки, в которой соседствуют прогоны разной насыщенности.

    Возвращает список строк, каждая — список кусков `(текст, шрифт, цвет)`.
    Слово меряется тем шрифтом, которым будет нарисовано, а пробелы берутся из
    самого текста: досочинять их на стыке прогонов нельзя — получается двойной
    пробел перед жирным куском и пробел перед точкой после него.

    Перенос делает `metrics.wrap_tokens` — тот же код, что и при сборке, чтобы
    предпросмотр и вёрстка не расходились в том, сколько вышло строк."""
    tokens, meta = [], []
    for text, fo, col in runs:
        for tok in _WS.split(text):
            if tok:
                tokens.append((tok, len(meta), False)); meta.append((fo, col))
    lines = M.wrap_tokens(tokens, 0, w,
                          measure=lambda t, i, _mo: d.textlength(t, font=meta[i][0]))
    return [[(t, meta[i][0], meta[i][1]) for t, i, _ in ln] for ln in lines] or [[]]

def draw_tf(d, tf, x, y, w, h, report, tag):
    """Нарисовать рамку и сказать, переполнена ли она.

    Две высоты считаются РАЗНЫМИ мерками, и это не дублирование:

      `need_draw`   — по фактической ширине рамки, в DejaVu. Столько строк
                      видно на картинке, и именно это глаз проверяет.
      `need_report` — по ширине, увеличенной на 1/K_REPORT, то есть в
                      предположении более узкого Arial. Только переполнение,
                      которое переживает эту поправку, попадает в список.

    Прежде список считался по `need_draw`: на нём висели рамки, у которых в
    настоящем Arial запас, и настоящие дефекты в нём терялись."""
    # Поля берутся у САМОЙ рамки, а не назначаются здесь константой. Вёрстка
    # ставит `margin_* = 0` (текст идёт от края рамки), и прежние «4 px с
    # каждой стороны» отнимали у каждой рамки 0,067″ ширины и столько же
    # высоты — этого хватало, чтобы ровно уложенный текст переносился на
    # лишнюю строку И попадал в список переполнений. То есть предпросмотр
    # сообщал о дефекте, который создавал сам.
    ml, mr, mt, mb = (_margin(tf, a) for a in
                      ("margin_left", "margin_right", "margin_top", "margin_bottom"))
    cy, need_draw, need_report = y + mt, 0, 0
    rows = []                       # копим строки, рисуем после — см. ниже
    wide = w - ml - mr                      # как есть — по этой ширине рисуем
    narrow = wide / M.K_REPORT              # как в Arial — по этой судим
    room = h - mt - mb
    for p in tf.paragraphs:
        pt, nm, _bo, _col = run_info(p)
        try:
            mult = float(p.line_spacing) if isinstance(p.line_spacing, float) else 1.22
        except Exception:
            mult = 1.22
        lh = int(pt * 120 / 72 * mult)
        runs = []
        for r in (p.runs or []):
            rpt, rnm, rbo, rcol = run_info(p, r)
            low = rnm.lower()
            runs.append((r.text, font(rpt, rbo, low.startswith("consol")
                                      or low.startswith("dejavu sans mono")), rcol))
        if not runs:
            need_draw += lh; need_report += lh; rows.append(([], lh))
            continue
        # «Не переносить» — режим ТЕРМИНАЛЬНОЙ карточки, где моноширинным набран
        # весь абзац. Прежде режим включался по ПЕРВОМУ прогону — и ячейка вида
        # «`---` первой строкой плюс описание» разъезжалась на строку-на-прогон:
        # `---` отдельной строкой, остальное отдельной, поверх разделителя.
        mono = all(_is_mono(r) for r in p.runs)
        if mono:
            # Моноширинный вывод не переносится — но и не разваливается на
            # строку-на-прогон: после того как `` `кусок` `` в листинге разметки
            # стал отдельным (жирным) прогоном, одна строка файла рисовалась
            # тремя, и карточка на картинке «переполнялась», будучи целой.
            lines = [[]]
            for seg, fo, col in runs:
                for k, part in enumerate(seg.split("\n")):
                    if k:
                        lines.append([])
                    if part:
                        lines[-1].append((part, fo, col))
            wide_n = len(lines)
        else:
            lines = wrap_runs(d, runs, wide)
            wide_n = len(lines)
        for ln in lines:
            rows.append((ln, lh))
        need_draw += wide_n * lh
        need_report += (wide_n if mono else len(wrap_runs(d, runs, narrow))) * lh

    # Вертикальное выравнивание рамки. Прежде предпросмотр его ИГНОРИРОВАЛ и
    # всегда рисовал от верха — отчего текст в пилюле-термине и в ячейке
    # таблицы (обе выровнены по центру) вылезал на картинке за верхний край,
    # и каждая такая рамка выглядела сломанной, будучи целой.
    try:
        va = int(tf.vertical_anchor) if tf.vertical_anchor is not None else 0
    except Exception:
        va = 0
    if va == 3 and need_draw < room:            # MIDDLE
        cy = y + mt + (room - need_draw) / 2
    elif va == 4 and need_draw < room:          # BOTTOM
        cy = y + mt + (room - need_draw)
    else:
        cy = y + mt
    for ln, lh in rows:
        cx = x + ml
        for seg, fo, col in ln:
            d.text((cx, cy), seg, font=fo, fill=col)
            cx += d.textlength(seg, font=fo)
        cy += lh
    if need_report > room + 2:
        report.append(f"{tag}: текст выше рамки на {(need_report - room) / 120:.2f}\" "
                      f"(по мерке Arial; в DejaVu на картинке "
                      f"{(need_draw - room) / 120:+.2f}\")")

def deck_order(pptx=None):
    """Порядок слайдов файла. Для блочной сборки `sem-05-n.pptx` берётся из
    тех же файлов, из которых её собрали, а не из `deck.yaml`, где блока ещё
    нет — иначе предпросмотр показывал бы чужой слайд под нужным именем."""
    name = Path(pptx).stem if pptx else ""
    if name.startswith("sem-05-") and len(name) > 7:
        import build_sem05 as B
        return [s["id"] for s in B.deck_from_files(name[7:])]
    import yaml
    deck = yaml.safe_load((Path(__file__).parent.parent / "deck.yaml").read_text())
    return [s["id"] for s in deck["slides"]]


def render(pptx, ids, order=None):
    """`order` — перечень идентификаторов в порядке слайдов файла. По умолчанию
    берётся из `deck.yaml`. Отдельный параметр нужен пробнику приёма
    (`probe_base_edge.py`): у него своя однослайдовая дека, которой в
    `deck.yaml` нет и быть не должно."""
    prs = Presentation(pptx)
    if order is None:
        order = deck_order(pptx)
    problems = []
    for sid in ids:
        sl = prs.slides[order.index(sid) if sid in order else ids.index(sid)]
        W, H = int(prs.slide_width * SC), int(prs.slide_height * SC)
        bgc = (255, 255, 255)
        try: bgc = tuple(sl.background.fill.fore_color.rgb)
        except Exception: pass
        im = Image.new("RGB", (W, H), bgc); d = ImageDraw.Draw(im)
        for sh in sl.shapes:
            x, y, w, h = [int(v * SC) for v in (sh.left, sh.top, sh.width, sh.height)]
            if sh.shape_type is not None and sh.has_chart if False else False: pass
            if sh.shape_type == 13 or sh.__class__.__name__ == "Picture":
                try:
                    im.paste(Image.open(__import__("io").BytesIO(sh.image.blob)).convert("RGB").resize((w, h)), (x, y))
                except Exception as e: problems.append(f"{sid} картинка: {e}")
                continue
            if getattr(sh, "has_table", False):
                tb = sh.table
                cw = [int(c.width * SC) for c in tb.columns]
                ry = y
                for i, row in enumerate(tb.rows):
                    rh = int(row.height * SC); cx = x
                    for j, cell in enumerate(row.cells):
                        fc = (255, 255, 255)
                        try: fc = tuple(cell.fill.fore_color.rgb)
                        except Exception: pass
                        d.rectangle([cx, ry, cx + cw[j], ry + rh], fill=fc, outline=(200, 210, 220))
                        draw_tf(d, cell.text_frame, cx, ry, cw[j], rh, problems, f"{sid} ячейка [{i},{j}]")
                        cx += cw[j]
                    ry += rh
                if ry > H: problems.append(f"{sid}: таблица уходит за нижний край на {(ry - H) / 120:.2f}\"")
                continue
            if sh.has_text_frame and not sh.text_frame.text.strip():
                pass
            if sh.shape_type == 1 or getattr(sh, "fill", None) is not None and not sh.has_text_frame:
                pass
            grad = grad_stops(sh)
            if grad:
                paint_gradient(im, x, y, w, h, grad)
                continue
            try:
                fill = tuple(sh.fill.fore_color.rgb) if sh.fill.type == 1 else None
            except Exception: fill = None
            try:
                line = tuple(sh.line.color.rgb) if sh.line.fill.type == 1 else None
            except Exception: line = None
            if fill or line:
                d.rounded_rectangle([x, y, x + w, y + h], radius=10, fill=fill, outline=line, width=2)
            if sh.has_text_frame and sh.text_frame.text.strip():
                draw_tf(d, sh.text_frame, x, y, w, h, problems, f"{sid} «{sh.text_frame.text[:28]}…»")
            if x < 0 or y < 0 or x + w > W + 2 or y + h > H + 2:
                problems.append(f"{sid}: фигура вне канвы ({x},{y},{w},{h})")
        im.save(OUT / f"{sid}.png")
    return problems

if __name__ == "__main__":
    args = sys.argv[1:]
    pptx = Path(__file__).parent / "sem-05.pptx"
    if args and args[0].endswith(".pptx"):
        pptx = Path(args[0]); args = args[1:]
    ids = args or ["s01","s02","s03","s03","s04","s05","s06","s46","s47","s48"]
    for p in render(pptx, ids):
        print("•", p)
    print("предпросмотр:", ", ".join(ids))
