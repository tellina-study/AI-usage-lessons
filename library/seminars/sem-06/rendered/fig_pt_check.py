"""Пол кегля на схемах — В ПУНКТАХ, по собранной деке, а не на глаз.

Кегль схемы живёт в пикселях холста, а порог деки задан в пунктах, и
переводной коэффициент — не свойство рисовалки: он зависит от того, какой
ШИРИНЫ схема в итоге встала на слайд (`deck_kit.figure` вписывает её в
отведённый прямоугольник, и при нехватке высоты ужимает целиком). Поэтому
считается по трём замерам: минимальный кегль в px (его печатает
`fig_toolkit.report()`), ширина холста в px и ширина картинки на слайде в
дюймах — из самого .pptx, по sha1 блоба.

    python3 fig_pt_check.py sem-06-en.pptx 'mcp-n09-mehanizm-en.png=23' ...
"""
import hashlib, sys
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Emu

FIG = Path(__file__).parent / "figures"
pptx = Path(sys.argv[1])
want = dict(a.split("=") for a in sys.argv[2:])

by_sha = {}
for p in sorted(FIG.glob("*.png")):
    by_sha[hashlib.sha1(p.read_bytes()).hexdigest()] = p

pr = Presentation(pptx)
print(f"{'схема':44} {'стр.':>4} {'холст px':>9} {'на слайде':>10} {'px/дюйм':>8} {'кегль px':>8} {'pt':>6}")
rows = []
for n, sl in enumerate(pr.slides, 1):
    for sh in sl.shapes:
        if sh.shape_type != 13:                      # PICTURE
            continue
        p = by_sha.get(sh.image.sha1)
        if p is None:
            print(f"  ⚠ стр. {n}: картинка не из figures/ (sha1 {sh.image.sha1[:8]})")
            continue
        iw, _ = Image.open(p).size
        inch = Emu(sh.width).inches
        ppi = iw / inch
        px = want.get(p.name)
        pt = (float(px) / ppi * 72) if px else None
        rows.append((p.name, n, iw, inch, ppi, px, pt))
        print(f"{p.name:44} {n:>4} {iw:>9} {inch:>9.2f}″ {ppi:>8.1f} "
              f"{(px or '—'):>8} {(f'{pt:.2f}' if pt else '—'):>6}")
bad = [r for r in rows if r[6] is not None and r[6] < 7.5]
print(f"\nсхем на деке: {len(rows)} · ниже пола 7,5 pt: {len(bad)}"
      + ("".join(f"\n  ✘ {r[0]} — {r[6]:.2f} pt" for r in bad) if bad else ""))
